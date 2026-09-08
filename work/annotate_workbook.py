import copy
import json
import math
import re
import statistics
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


BASE = Path("/Users/yyt/Documents/Codex/2026-08-14/diy")
JSON_PATH = Path("/Users/yyt/Desktop/违法测2json.txt")
TARGET_PATH = Path("/Users/yyt/Downloads/自办违法测2.xlsx")
REF_PATH = Path("/Users/yyt/Downloads/南京-自办违纪模拟测试集.xlsx")
OUT_DIR = BASE / "outputs"
OUT_PATH = OUT_DIR / "自办违法测2_模型标注统计.xlsx"


def norm_text(value):
    if value is None:
        return ""
    s = str(value)
    s = s.replace("（", "(").replace("）", ")")
    s = s.replace("／", "/").replace("：", ":")
    s = re.sub(r"\s+", "", s)
    replacements = {
        "签名时间": "日期",
        "落款时间": "日期",
        "时间": "日期",
        "分管常委(委员)": "分管常委委员",
        "分管常委/委员": "分管常委委员",
        "分管副书记(副主任)": "分管副书记副主任",
        "分管副书记（副主任）": "分管副书记副主任",
        "主要领导": "书记主任",
        "书记(主任)": "书记主任",
        "审批": "意见",
        "批示": "意见",
        "栏位": "",
        "处": "",
        "[": "",
        "]": "",
        "【": "",
        "】": "",
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    return s


def ratio(a, b):
    return SequenceMatcher(None, norm_text(a), norm_text(b)).ratio()


def is_yes(value):
    return str(value).strip() == "是"


def gold_correct(value):
    s = "" if value is None else str(value).strip()
    if not s or s in {"/", "None", "none", "无"}:
        return None
    return "是" in s


def classify_rule(rule_name):
    s = "" if rule_name is None else str(rule_name)
    if "捺印" in s or "按捺" in s:
        return "捺印"
    if "印章" in s or "盖章" in s:
        return "印章"
    if "签名" in s and ("时间" in s or "日期" in s):
        return "日期"
    if "落款时间" in s or "日期" in s:
        return "日期"
    if "签名" in s:
        return "签名"
    return "其他"


def pct(numer, denom):
    if denom == 0:
        return 0
    return numer / denom


def pct_text(value):
    return f"{value:.2%}"


def fill_down_context(ws):
    current = {}
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        for c in range(1, 8):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                current[c] = v
            row[c] = current.get(c)
        for c in range(8, ws.max_column + 1):
            row[c] = ws.cell(r, c).value
        rows.append((r, row))
    return rows


def flatten_json_results(data):
    out = []
    for material in data.get("ruleResults", []):
        material_name = material.get("materialName")
        outline_name = material.get("outlineName")
        for rule in material.get("ruleResults", []):
            out.append(
                {
                    "materialName": material_name,
                    "outlineName": outline_name,
                    "ruleName": rule.get("ruleName"),
                    "isCorrect": rule.get("isCorrect"),
                    "questionDesc": rule.get("questionDesc"),
                    "raw": rule,
                }
            )
    return out


def best_match(row, candidates, used):
    mat_type = row.get(4)
    mat_name = row.get(5)
    rule = row.get(8)
    best = None
    best_score = -1
    best_idx = None
    for idx, cand in enumerate(candidates):
        if idx in used:
            continue
        material_score = max(
            ratio(mat_type, cand["materialName"]),
            ratio(mat_name, cand["outlineName"]),
            ratio(mat_name, cand["materialName"]),
            ratio(mat_type, cand["outlineName"]),
        )
        rule_score = ratio(rule, cand["ruleName"])
        score = 0.36 * material_score + 0.64 * rule_score
        if norm_text(rule) == norm_text(cand["ruleName"]):
            score += 0.08
        if norm_text(mat_type) and norm_text(mat_type) in norm_text(cand["materialName"] + cand["outlineName"]):
            score += 0.04
        if score > best_score:
            best_score = score
            best = cand
            best_idx = idx
    if best_score >= 0.67:
        return best_idx, best, best_score
    return None, None, best_score


def copy_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy.copy(src.alignment)


def set_cell(ws, row, col, value):
    cell = ws.cell(row, col)
    cell.value = value


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    candidates = flatten_json_results(data)

    wb = openpyxl.load_workbook(TARGET_PATH)
    ws = wb.active
    rows = fill_down_context(ws)

    used = set()
    annotated = []
    for r, row in rows:
        idx, match, score = best_match(row, candidates, used)
        if match:
            used.add(idx)
        expected = is_yes(row.get(13))
        model_triggered = bool(match)
        model_correct = match.get("isCorrect") if match else None
        gold = gold_correct(row.get(14))
        if model_triggered and gold is not None:
            aa = "是" if bool(model_correct) == bool(gold) else "否"
        elif model_triggered and not expected:
            aa = "否"
        elif expected and not model_triggered:
            aa = "否"
        else:
            aa = "/"
        category = classify_rule(row.get(8))
        annotated.append(
            {
                "row": r,
                "rule": row.get(8),
                "expected": expected,
                "gold_correct": gold,
                "model_triggered": model_triggered,
                "model_correct": model_correct,
                "aa": aa,
                "category": category,
                "match_score": score,
                "match_rule": match.get("ruleName") if match else "",
                "match_material": (match.get("materialName") if match else ""),
                "desc": (match.get("questionDesc") if match else ""),
            }
        )

    # Per-row category flags and model annotations.
    category_cols = {
        "签名": (15, 16),
        "日期": (17, 18),
        "捺印": (19, 20),
        "印章": (21, 22),
        "其他": (24, 25),
    }
    flag_fill = PatternFill("solid", fgColor="EAF4FF")
    err_fill = PatternFill("solid", fgColor="FCE4D6")
    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    for item in annotated:
        r = item["row"]
        for c in range(15, 26):
            ws.cell(r, c).value = None
        total_col, err_col = category_cols[item["category"]]
        ws.cell(r, total_col).value = 1
        ws.cell(r, total_col).fill = flag_fill
        is_error = item["aa"] == "否"
        ws.cell(r, err_col).value = 1 if is_error else 0
        ws.cell(r, err_col).fill = err_fill if is_error else ok_fill

        ws.cell(r, 23).value = 1 if item["category"] in {"签名", "日期", "捺印", "印章"} else None
        ws.cell(r, 24).value = 1 if item["category"] == "其他" else None
        ws.cell(r, 25).value = 1 if item["category"] == "其他" and is_error else 0 if item["category"] == "其他" else None
        ws.cell(r, 26).value = "是" if item["model_triggered"] else "否"
        ws.cell(r, 27).value = item["aa"]
        ws.cell(r, 28).value = item["desc"] if item["aa"] == "否" and item["desc"] else None
        ws.cell(r, 29).value = f"匹配规则：{item['match_rule']}；匹配分：{item['match_score']:.2f}" if item["model_triggered"] else None

    # Header summary in O:Y, Z:AA.
    def cat_total(cat):
        return sum(1 for x in annotated if x["category"] == cat)

    def cat_errors(cat):
        return sum(1 for x in annotated if x["category"] == cat and x["aa"] == "否")

    sig_total, sig_err = cat_total("签名"), cat_errors("签名")
    date_total, date_err = cat_total("日期"), cat_errors("日期")
    press_total, press_err = cat_total("捺印"), cat_errors("捺印")
    seal_total, seal_err = cat_total("印章"), cat_errors("印章")
    comp_total = sig_total + date_total + press_total + seal_total
    comp_err = sig_err + date_err + press_err + seal_err
    other_total, other_err = cat_total("其他"), cat_errors("其他")

    ws["O1"] = f"签名\n（{sig_total}）\n准确率：{pct_text(1 - pct(sig_err, sig_total))}\n=1-({sig_err}/{sig_total})"
    ws["P1"] = f"签名错误（{sig_err}）"
    ws["Q1"] = f"日期\n（{date_total}）\n准确率：{pct_text(1 - pct(date_err, date_total))}\n=1-({date_err}/{date_total})"
    ws["R1"] = f"日期错误（{date_err}）"
    ws["S1"] = f"捺印\n（{press_total}）\n准确率：{pct_text(1 - pct(press_err, press_total))}\n=1-({press_err}/{press_total})"
    ws["T1"] = f"捺印错误（{press_err}）"
    ws["U1"] = f"印章\n（{seal_total}）\n准确率：{pct_text(1 - pct(seal_err, seal_total))}\n=1-({seal_err}/{seal_total})"
    ws["V1"] = f"印章错误（{seal_err}）"
    ws["W1"] = f"综合\n（签名、日期、捺印和印章）\n（总：{comp_total}，错：{comp_err}）\n准确率：{pct_text(1 - pct(comp_err, comp_total))}\n1-({comp_err}/{comp_total})"
    ws["X1"] = f"其他\n（{other_total}）\n准确率：{pct_text(1 - pct(other_err, other_total))}\n1-({other_err}/{other_total})"
    ws["Y1"] = f"其他错误（{other_err}）"

    expected_total = sum(1 for x in annotated if x["expected"])
    expected_triggered = sum(1 for x in annotated if x["expected"] and x["model_triggered"])
    model_triggered = sum(1 for x in annotated if x["model_triggered"])
    model_correct_yes = sum(1 for x in annotated if x["model_triggered"] and x["aa"] == "是")
    model_error = sum(1 for x in annotated if x["model_triggered"] and x["aa"] == "否")
    correct_recalled = sum(1 for x in annotated if x["expected"] and x["model_triggered"] and x["aa"] == "是")
    recall = pct(correct_recalled, expected_total)
    trigger_rate = pct(expected_triggered, expected_total)
    accuracy = pct(model_correct_yes, model_triggered)
    ws["Z1"] = f"（20260814）\nJSON统计-\n触发率：{pct_text(trigger_rate)}\n模型是否触发\n模型触发且金标准应触发/金标准应触发={expected_triggered}/{expected_total}"
    ws["AA1"] = f"（20260814）\nJSON统计-\n准确率：{pct_text(accuracy)}\n模型是否正确\n模型实际检测：是：{model_triggered}，其中错误：{model_error}，1-({model_error}/{model_triggered})"

    for c in range(15, 30):
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.cell(1, c).font = Font(bold=True)
    ws.row_dimensions[1].height = 92

    # Summary sheet.
    if "统计汇总" in wb.sheetnames:
        del wb["统计汇总"]
    summary = wb.create_sheet("统计汇总", 0)
    headers = ["指标", "分子", "分母", "结果", "口径"]
    summary.append(headers)
    rows_out = [
        ["召回率", correct_recalled, expected_total, pct_text(recall), "金标准应覆盖的检测项中，模型触发且回答正确的比例"],
        ["触发率", expected_triggered, expected_total, pct_text(trigger_rate), "金标准应覆盖的检测项中，JSON 找到对应检测项的比例"],
        ["准确率", model_correct_yes, model_triggered, pct_text(accuracy), "JSON 实际触发的检测项中，AA列为“是”的比例"],
        ["签名准确率", sig_total - sig_err, sig_total, pct_text(1 - pct(sig_err, sig_total)), "按规则名称分类"],
        ["日期准确率", date_total - date_err, date_total, pct_text(1 - pct(date_err, date_total)), "按规则名称分类"],
        ["捺印准确率", press_total - press_err, press_total, pct_text(1 - pct(press_err, press_total)), "按规则名称分类"],
        ["印章准确率", seal_total - seal_err, seal_total, pct_text(1 - pct(seal_err, seal_total)), "按规则名称分类"],
        ["综合准确率", comp_total - comp_err, comp_total, pct_text(1 - pct(comp_err, comp_total)), "签名、日期、捺印、印章合计"],
        ["其他准确率", other_total - other_err, other_total, pct_text(1 - pct(other_err, other_total)), "除签名、日期、捺印、印章外"],
    ]
    for row in rows_out:
        summary.append(row)
    summary.append([])
    summary.append(["匹配检查", "数量", "", "", "说明"])
    summary.append(["JSON规则总数", len(candidates), "", "", "违法测2json.txt 中展开后的规则条数"])
    summary.append(["Excel数据行数", len(annotated), "", "", "自办违法测2.xlsx 数据行"])
    summary.append(["已匹配JSON规则", len(used), "", "", "通过材料+规则名称匹配"])
    summary.append(["未匹配JSON规则", len(candidates) - len(used), "", "", "JSON中未落到目标表的规则"])
    summary.append(["Excel未匹配行", len(annotated) - len(used), "", "", "目标表中没有对应JSON规则的行，Z列记为否"])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=thin)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    for col, width in enumerate([18, 12, 12, 14, 54], 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = "A2"

    # Preserve usability on source sheet.
    ws.freeze_panes = "A2"
    for c in range(15, 30):
        ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 10, 14)
    ws.column_dimensions["AB"].width = 42
    ws.column_dimensions["AC"].width = 42

    wb.save(OUT_PATH)

    print(json.dumps({
        "output": str(OUT_PATH),
        "expected_total": expected_total,
        "expected_triggered": expected_triggered,
        "model_triggered": model_triggered,
        "model_correct": model_correct_yes,
        "model_error": model_error,
        "recall": recall,
        "trigger_rate": trigger_rate,
        "accuracy": accuracy,
        "category": {
            "签名": [sig_total, sig_err, 1 - pct(sig_err, sig_total)],
            "日期": [date_total, date_err, 1 - pct(date_err, date_total)],
            "捺印": [press_total, press_err, 1 - pct(press_err, press_total)],
            "印章": [seal_total, seal_err, 1 - pct(seal_err, seal_total)],
            "综合": [comp_total, comp_err, 1 - pct(comp_err, comp_total)],
            "其他": [other_total, other_err, 1 - pct(other_err, other_total)],
        }
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

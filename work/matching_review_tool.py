import argparse
import base64
import copy
import json
import re
import sqlite3
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from collections import defaultdict
from datetime import datetime
from difflib import SequenceMatcher
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Keep runtime data beside the project so the tool remains portable when moved.
BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "outputs"
UPLOAD_DIR = BASE_DIR / "work" / "uploads"
HISTORY_DIR = BASE_DIR / "work" / "history"
HISTORY_INDEX = HISTORY_DIR / "index.json"
DB_PATH = BASE_DIR / "work" / "review_tool.db"
DEFAULT_OUTPUT = OUTPUT_DIR / "自办违法测2_复核后标注统计.xlsx"
DEFAULT_ERROR_REASON_OUTPUT = OUTPUT_DIR / "错误原因明细.xlsx"
DEFAULT_HISTORY_OUTPUT = OUTPUT_DIR / "比对历史导出.xlsx"
RULE_ANNOTATION_OUTPUT_PREFIX = "规则金标准标注"
ANALYSIS_JOBS = {}
ANALYSIS_JOBS_LOCK = threading.Lock()
ANALYSIS_JOB_TTL_SECONDS = 30 * 60
BATCH_JOBS = {}
BATCH_JOBS_LOCK = threading.Lock()
BATCH_JOB_TTL_SECONDS = 12 * 60 * 60
DEFAULT_MODEL_RESULT_API = "http://120.195.114.115:5025/intelli-judge/cpc/procedure/rule/result"
DEFAULT_CASE_DETAIL_API = "http://120.195.114.115:5025/intelli-judge/case/caseDetail"
CATEGORY_NAMES = ("签名", "日期", "捺印", "印章", "其他")
ERROR_REASON_HEADERS = (
    "错误原因", "来源文件", "材料类型", "对应材料名称",
    "所属阶段", "阶段时间顺序", "规则名称", "特殊条件", "判断规则", "错误场景", "是否触发",
    "是否正确", "模型是否触发", "模型是否正确", "具体问题内容",
)


def norm_text(value):
    if value is None:
        return ""
    s = str(value)
    s = s.replace("（", "(").replace("）", ")")
    s = s.replace("／", "/").replace("：", ":")
    s = re.sub(r"\s+", "", s)
    replacements = {
        "签名时间": "日期",
        "落款时间": "日期",
        "时间": "日期",
        "分管常委(委员)": "分管常委委员",
        "分管常委/委员": "分管常委委员",
        "分管副书记(副主任)": "分管副书记副主任",
        "分管副书记（副主任）": "分管副书记副主任",
        "主要领导": "书记主任",
        "书记(主任)": "书记主任",
        "审批": "意见",
        "批示": "意见",
        "栏位": "",
        "处": "",
        "[": "",
        "]": "",
        "【": "",
        "】": "",
    }
    for old, new in replacements.items():
        s = s.replace(old, new)
    return s


def ratio(left, right):
    return SequenceMatcher(None, norm_text(left), norm_text(right)).ratio()


def is_yes(value):
    return str(value).strip() == "是"


def gold_correct(value):
    s = "" if value is None else str(value).strip()
    if not s or s in {"/", "None", "none", "无"}:
        return None
    return s == "是"


def split_copy_values(value, allow_text_parts=False):
    """Split duplicate-copy notation without splitting ordinary error descriptions."""
    s = "" if value is None else str(value).strip()
    if not s:
        return [""]
    parts = re.split(r"\s*--\s*|\s*-\s*", s)
    if not allow_text_parts and not all(normalize_gold_trigger(part) in {"是", "否"} for part in parts):
        return [s]
    return [part.strip() for part in parts] if len(parts) > 1 else [s]


def normalize_gold_trigger(value):
    s = "" if value is None else str(value).strip()
    if s in {"是", "否"}:
        return s
    return ""


def normalize_result_bool(value):
    """Normalize JSON boolean fields that may be true/false, 1/0, or text."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    s = str(value or "").strip().lower()
    if s in {"true", "1", "是", "yes"}:
        return True
    if s in {"false", "0", "否", "no"}:
        return False
    return None


def db_connect():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_app_db():
    with db_connect() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS api_cache (
                cache_key TEXT PRIMARY KEY,
                base_url TEXT NOT NULL,
                case_id TEXT NOT NULL,
                fetched_at TEXT NOT NULL,
                rule_count INTEGER NOT NULL,
                content TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS batch_jobs (
                id TEXT PRIMARY KEY,
                base_url TEXT NOT NULL,
                excel_file_name TEXT NOT NULL,
                total INTEGER NOT NULL,
                completed INTEGER NOT NULL DEFAULT 0,
                success INTEGER NOT NULL DEFAULT 0,
                failed INTEGER NOT NULL DEFAULT 0,
                status TEXT NOT NULL,
                message TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS batch_cases (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id TEXT NOT NULL,
                case_id TEXT NOT NULL,
                excel_file_name TEXT,
                excel_path TEXT,
                status TEXT NOT NULL,
                message TEXT NOT NULL,
                json_rules INTEGER,
                excel_rows INTEGER,
                review_rows INTEGER,
                missed_rows INTEGER,
                model_triggered INTEGER,
                model_correct INTEGER,
                expected_total INTEGER,
                expected_triggered INTEGER,
                error_total INTEGER,
                error_recalled INTEGER,
                trigger_rate REAL,
                recall REAL,
                accuracy REAL,
                analysis_path TEXT,
                updated_at TEXT NOT NULL
            )
            """
        )
        existing_columns = {
            row["name"]
            for row in conn.execute("PRAGMA table_info(batch_cases)").fetchall()
        }
        if "excel_file_name" not in existing_columns:
            conn.execute("ALTER TABLE batch_cases ADD COLUMN excel_file_name TEXT")
        if "excel_path" not in existing_columns:
            conn.execute("ALTER TABLE batch_cases ADD COLUMN excel_path TEXT")


def now_text():
    return time.strftime("%Y-%m-%d %H:%M:%S")


def model_result_url(base_url, case_id):
    base_url = str(base_url or "").strip()
    case_id = str(case_id or "").strip()
    if not base_url:
        raise ValueError("请填写接口地址。")
    if not case_id:
        raise ValueError("请填写 caseId。")
    parsed = urllib.parse.urlparse(base_url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("接口地址必须是 http 或 https。")
    query = urllib.parse.parse_qs(parsed.query, keep_blank_values=True)
    query["caseId"] = [case_id]
    return urllib.parse.urlunparse(parsed._replace(query=urllib.parse.urlencode(query, doseq=True)))


def extract_model_result_items(body):
    data = body
    if isinstance(body, dict) and "data" in body:
        code = body.get("code")
        if code not in (None, 0, "0"):
            raise ValueError(str(body.get("msg") or body.get("message") or f"接口返回 code={code}"))
        data = body.get("data")
    if not isinstance(data, list):
        raise ValueError("接口返回格式不正确：未找到 data 数组。")
    return [item for item in data if isinstance(item, dict)]


def fetch_model_result_from_api(base_url, case_id, timeout=20):
    url = model_result_url(base_url, case_id)
    req = urllib.request.Request(url, headers={"Accept": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            raw = res.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        raise ValueError(f"接口请求失败：HTTP {exc.code}") from exc
    except urllib.error.URLError as exc:
        raise ValueError(f"接口请求失败：{exc.reason}") from exc
    try:
        body = json.loads(raw)
    except ValueError as exc:
        raise ValueError("接口返回不是合法 JSON。") from exc
    items = extract_model_result_items(body)
    cache_key = f"{base_url.strip()}::{case_id.strip()}"
    with db_connect() as conn:
        conn.execute(
            """
            INSERT OR REPLACE INTO api_cache(cache_key, base_url, case_id, fetched_at, rule_count, content)
            VALUES(?, ?, ?, ?, ?, ?)
            """,
            (cache_key, base_url.strip(), case_id.strip(), now_text(), len(items), json.dumps(items, ensure_ascii=False)),
        )
    return {"url": url, "caseId": case_id.strip(), "count": len(items), "content": items}


def case_detail_url(base_url, case_id):
    base_url = str(base_url or "").strip().rstrip("/")
    case_id = str(case_id or "").strip()
    if not base_url:
        raise ValueError("请填写案件详情接口地址。")
    if not case_id:
        raise ValueError("请填写 caseId。")
    parsed = urllib.parse.urlparse(base_url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("案件详情接口地址必须是 http 或 https。")
    return f"{base_url}/{urllib.parse.quote(case_id, safe='')}"


def fetch_case_detail_from_api(base_url, case_id, timeout=20):
    url = case_detail_url(base_url, case_id)
    req = urllib.request.Request(url, headers={"Accept": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            raw = res.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        raise ValueError(f"案件接口请求失败：HTTP {exc.code}") from exc
    except urllib.error.URLError as exc:
        raise ValueError(f"案件接口请求失败：{exc.reason}") from exc
    try:
        body = json.loads(raw)
    except ValueError as exc:
        raise ValueError("案件接口返回不是合法 JSON。") from exc
    if isinstance(body, dict):
        code = body.get("code")
        if code not in (None, 0, "0"):
            raise ValueError(str(body.get("msg") or body.get("message") or f"案件接口返回 code={code}"))
        data = body.get("data", body)
    else:
        data = body
    if not isinstance(data, dict):
        raise ValueError("案件接口返回格式不正确：未找到案件详情对象。")
    return {"url": url, "caseId": str(case_id).strip(), "data": data}


def gold_row_pending(row):
    """A duplicate copy stays outside statistics until both gold-standard fields are confirmed."""
    return bool(row.get("materialCopyPending")) and not (
        row.get("goldTrigger") in {"是", "否"} and row.get("goldCorrect") in {"是", "否"}
    )


def display_material_type(row):
    value = str(row.get("materialType") or "")
    if row.get("materialCopyPending") and not gold_row_pending(row):
        return value.replace("·待确认", "")
    return value


def pct(numer, denom):
    return 0 if denom == 0 else numer / denom


def pct_text(value):
    return f"{value:.2%}"


def classify_rule(rule_name):
    s = "" if rule_name is None else str(rule_name)
    # These mixed rules follow the classification used in the reference workbook.
    if "落款组织名称" in s or "签名或印章" in s or "签名或盖章" in s:
        return "签名"
    if "捺印" in s or "按捺" in s:
        return "捺印"
    if "材料签批顺序" in s or "材料时间顺序" in s:
        return "日期"
    if "印章" in s or "盖章" in s or "公章" in s or "委印" in s:
        return "印章"
    if "时间" in s or "日期" in s:
        return "日期"
    if "签名" in s:
        return "签名"
    return "其他"


def load_json_candidates(json_path):
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("data"), list):
        data = data["data"]
    if isinstance(data, list):
        # Newer exports are already flat: each array item is one model detection result.
        # They do not expose a material-copy boundary, so keep one material group per
        # normalized material name and let the Excel M/N columns define duplicate copies.
        candidates = []
        material_index_by_key = {}
        for item in data:
            if not isinstance(item, dict):
                continue
            material_name = item.get("materialName") or ""
            outline_name = item.get("outlineName") or material_name
            material_key = f"{norm_text(material_name)}||{norm_text(item.get('materialStage') or outline_name)}"
            if not material_name and not outline_name:
                material_key = f"__material_{len(material_index_by_key) + 1}"
            if material_key not in material_index_by_key:
                material_index_by_key[material_key] = len(material_index_by_key) + 1
            is_correct = normalize_result_bool(item.get("isCorrect"))
            candidates.append(
                {
                    "id": len(candidates),
                    "materialName": material_name,
                    "outlineName": outline_name,
                    "ruleName": item.get("ruleName") or "",
                    "isCorrect": is_correct,
                    "modelTriggered": True,
                    "sourceField": "isCorrect",
                    "sourceFile": item.get("fileName") or item.get("fileId") or "",
                    "materialIndex": material_index_by_key[material_key],
                    "materialKey": material_key,
                    "materialOccurrence": 1,
                    "materialOccurrenceTotal": 1,
                    "questionDesc": item.get("questionDesc") or "",
                }
            )
        return candidates
    if not isinstance(data, dict):
        raise ValueError("JSON 格式不正确：顶层必须是数组或包含 ruleResults 的对象。")
    materials = data.get("ruleResults", [])
    total_by_material_key = defaultdict(int)
    seen_by_material_key = defaultdict(int)
    material_keys = []
    for material_index, material in enumerate(materials, start=1):
        material_name = material.get("materialName") or ""
        outline_name = material.get("outlineName") or ""
        material_key = f"{norm_text(material_name)}||{norm_text(outline_name)}"
        if not material_name and not outline_name:
            material_key = f"__material_{material_index}"
        material_keys.append(material_key)
        total_by_material_key[material_key] += 1
    candidates = []
    for material_index, material in enumerate(materials, start=1):
        material_name = material.get("materialName") or ""
        outline_name = material.get("outlineName") or ""
        source_file = material.get("fileName") or ""
        material_key = material_keys[material_index - 1]
        seen_by_material_key[material_key] += 1
        material_occurrence = seen_by_material_key[material_key]
        material_occurrence_total = total_by_material_key[material_key]
        if "isExit" in material:
            is_exit = bool(material.get("isExit"))
            cid = len(candidates)
            candidates.append(
                {
                    "id": cid,
                    "materialName": material_name,
                    "outlineName": outline_name,
                    "ruleName": "材料类型检测",
                    "isCorrect": is_exit,
                    "modelTriggered": is_exit,
                    "sourceField": "isExit",
                    "sourceFile": source_file,
                    "materialIndex": material_index,
                    "materialKey": material_key,
                    "materialOccurrence": material_occurrence,
                    "materialOccurrenceTotal": material_occurrence_total,
                    "questionDesc": f"{material_name or outline_name}：isExit={str(is_exit).lower()}，用于判断材料类型检测。"
                }
            )
        for rule in material.get("ruleResults", []):
            cid = len(candidates)
            candidates.append(
                {
                    "id": cid,
                    "materialName": material_name,
                    "outlineName": outline_name,
                    "ruleName": rule.get("ruleName") or "",
                    "isCorrect": rule.get("isCorrect"),
                    "modelTriggered": True,
                    "sourceField": "isCorrect",
                    "sourceFile": source_file,
                    "materialIndex": material_index,
                    "materialKey": material_key,
                    "materialOccurrence": material_occurrence,
                    "materialOccurrenceTotal": material_occurrence_total,
                    "questionDesc": rule.get("questionDesc") or "",
                }
            )
    return candidates


def safe_filename(name, fallback):
    name = Path(name or fallback).name
    name = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", name)
    return name or fallback


def review_output_path(analysis):
    excel_name = analysis.get("excelFileName") or Path(analysis.get("excelPath") or "").name or "复核结果.xlsx"
    safe_name = safe_filename(excel_name, "复核结果.xlsx")
    stem = Path(safe_name).stem or "复核结果"
    stamp = time.strftime("%Y%m%d_%H%M%S")
    return OUTPUT_DIR / f"{stem}_复核后标注统计_{stamp}.xlsx"


def history_id(value):
    value = str(value or "")
    return value if re.fullmatch(r"[0-9A-Za-z_-]{8,80}", value) else ""


def read_history_index():
    if not HISTORY_INDEX.exists():
        return []
    try:
        data = json.loads(HISTORY_INDEX.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (OSError, ValueError, TypeError):
        return []


def write_history_index(items):
    HISTORY_DIR.mkdir(parents=True, exist_ok=True)
    HISTORY_INDEX.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


def history_metadata(record):
    summary = record.get("summary") or {}
    return {
        "id": record["id"],
        "title": record.get("title") or "未命名比对",
        "createdAt": record.get("createdAt") or "",
        "jsonFileName": record.get("jsonFileName") or Path(record.get("jsonPath") or "").name,
        "excelFileName": record.get("excelFileName") or Path(record.get("excelPath") or "").name,
        "jsonPath": record.get("jsonPath") or "",
        "excelPath": record.get("excelPath") or "",
        "excelRows": summary.get("excelRows", 0),
        "jsonRules": summary.get("jsonRules", 0),
        "reviewRows": summary.get("reviewRows", 0),
        "recall": summary.get("recall", 0),
        "triggerRate": summary.get("triggerRate", 0),
        "accuracy": summary.get("accuracy", 0),
        "missedRows": summary.get("missedRows", 0),
        "modelReviewStartAt": record.get("modelReviewStartAt") or "",
        "modelReviewEndAt": record.get("modelReviewEndAt") or "",
        "modelReviewDurationSeconds": record.get("modelReviewDurationSeconds"),
        "modelReviewDurationText": record.get("modelReviewDurationText") or "",
    }


def parse_review_datetime(value):
    value = str(value or "").strip()
    if not value:
        return None
    for parser in (datetime.fromisoformat,):
        try:
            return parser(value)
        except (TypeError, ValueError):
            continue
    raise ValueError("模型阅卷时间格式不正确。")


def normalize_review_times(start_value, end_value):
    start = parse_review_datetime(start_value)
    end = parse_review_datetime(end_value)
    if not start and not end:
        return "", "", None, ""
    if not start or not end:
        raise ValueError("请同时填写模型阅卷开始时间和结束时间。")
    seconds = int((end - start).total_seconds())
    if seconds < 0:
        raise ValueError("模型阅卷结束时间不能早于开始时间。")
    duration_text = format_duration(seconds)
    return start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S"), seconds, duration_text


def format_duration(seconds):
    if seconds is None:
        return ""
    seconds = max(0, int(seconds))
    days, seconds = divmod(seconds, 86400)
    hours, seconds = divmod(seconds, 3600)
    minutes, seconds = divmod(seconds, 60)
    parts = []
    if days:
        parts.append(f"{days}天")
    if hours or days:
        parts.append(f"{hours}小时")
    if minutes or hours or days:
        parts.append(f"{minutes}分")
    parts.append(f"{seconds}秒")
    return "".join(parts)


def save_history_record(analysis, assignments, verdict_overrides, category_overrides, error_reason_overrides, title="", completed_materials=None, confirmed_rows=None, summary_snapshot=None, model_review_start="", model_review_end=""):
    model_review_start, model_review_end, model_review_duration_seconds, model_review_duration_text = normalize_review_times(
        model_review_start,
        model_review_end,
    )
    reviewed = apply_review_assignments(
        analysis,
        assignments,
        verdict_overrides,
        category_overrides,
        error_reason_overrides,
    )
    saved_summary = merge_summary_snapshot(reviewed["summary"], summary_snapshot)
    record_id = f"{time.strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}"
    created_at = time.strftime("%Y-%m-%d %H:%M:%S")
    record = {
        "id": record_id,
        "title": str(title or "").strip() or f"{analysis.get('jsonFileName') or Path(analysis.get('jsonPath', '')).name} × {analysis.get('excelFileName') or Path(analysis.get('excelPath', '')).name}",
        "createdAt": created_at,
        "jsonFileName": analysis.get("jsonFileName") or Path(analysis.get("jsonPath") or "").name,
        "excelFileName": analysis.get("excelFileName") or Path(analysis.get("excelPath") or "").name,
        "jsonPath": analysis.get("jsonPath") or "",
        "excelPath": analysis.get("excelPath") or "",
        "analysis": analysis,
        "assignments": assignments or {},
        "verdictOverrides": verdict_overrides or {},
        "categoryOverrides": category_overrides or {},
        "errorReasonOverrides": error_reason_overrides or {},
        "completedMaterials": completed_materials or {},
        "confirmedRows": confirmed_rows or {},
        "errorReasonDetails": collect_error_reason_details(reviewed),
        "summarySnapshot": copy.deepcopy(summary_snapshot) if isinstance(summary_snapshot, dict) else {},
        "reviewed": reviewed,
        "summary": saved_summary,
        "modelReviewStartAt": model_review_start,
        "modelReviewEndAt": model_review_end,
        "modelReviewDurationSeconds": model_review_duration_seconds,
        "modelReviewDurationText": model_review_duration_text,
    }
    HISTORY_DIR.mkdir(parents=True, exist_ok=True)
    (HISTORY_DIR / f"{record_id}.json").write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    items = [history_metadata(record)] + [item for item in read_history_index() if item.get("id") != record_id]
    write_history_index(items)
    return history_metadata(record)


def load_history_record(record_id):
    record_id = history_id(record_id)
    if not record_id:
        return None
    path = HISTORY_DIR / f"{record_id}.json"
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict) or data.get("id") != record_id:
            return None
        saved_summary = data.get("summarySnapshot") or data.get("summary")
        if isinstance(data.get("analysis"), dict):
            reviewed = apply_review_assignments(
                data["analysis"],
                data.get("assignments", {}),
                data.get("verdictOverrides", {}),
                data.get("categoryOverrides", {}),
                data.get("errorReasonOverrides", {}),
            )
            data["reviewed"] = reviewed
            data["summary"] = merge_summary_snapshot(reviewed["summary"], saved_summary)
            data["errorReasonDetails"] = collect_error_reason_details(reviewed)
        elif "errorReasonDetails" not in data and isinstance(data.get("reviewed"), dict):
            reviewed = copy.deepcopy(data["reviewed"])
            reviewed["excelFileName"] = data.get("excelFileName") or reviewed.get("excelFileName")
            data["errorReasonDetails"] = collect_error_reason_details(reviewed)
        return data
    except (OSError, ValueError, TypeError):
        return None


def merge_summary_snapshot(base_summary, summary_snapshot):
    summary = copy.deepcopy(base_summary or {})
    if isinstance(summary_snapshot, dict):
        summary.update(copy.deepcopy(summary_snapshot))
    return summary


def refreshed_history_index():
    current = read_history_index()
    refreshed = []
    for item in current:
        record = load_history_record(item.get("id", ""))
        refreshed.append(history_metadata(record) if record else item)
    if refreshed != current:
        write_history_index(refreshed)
    return refreshed


def delete_history_record(record_id):
    record_id = history_id(record_id)
    if not record_id:
        return False
    path = HISTORY_DIR / f"{record_id}.json"
    existed = path.exists()
    if existed:
        path.unlink()
    write_history_index([item for item in read_history_index() if item.get("id") != record_id])
    return existed


def save_uploaded_inputs(payload):
    json_upload = payload.get("jsonUpload") or {}
    excel_upload = payload.get("excelUpload") or {}
    if json_upload.get("content") is None:
        raise ValueError("请先上传 JSON / TXT 文件。")
    if not excel_upload.get("base64"):
        raise ValueError("请先上传 Excel 文件。")

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    json_filename = safe_filename(json_upload.get("name"), "model_result.json")
    json_path = UPLOAD_DIR / f"{stamp}-{json_filename}"
    Path(json_path).write_text(json_upload["content"], encoding="utf-8")

    excel_filename = safe_filename(excel_upload.get("name"), "gold_standard.xlsx")
    excel_path = UPLOAD_DIR / f"{stamp}-{excel_filename}"
    Path(excel_path).write_bytes(base64.b64decode(excel_upload["base64"]))

    return json_path, excel_path


def json_safe_cell(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def load_rule_annotation_workbook(excel_path):
    """Return the workbook rows used by the standalone rule annotation module."""
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    sheets = []
    for ws in workbook.worksheets:
        if ws.max_row < 1 or ws.max_column < 1:
            continue
        headers = [json_safe_cell(ws.cell(1, c).value) or f"第{c}列" for c in range(1, max(ws.max_column, 14) + 1)]
        rows = []
        carried_values = {}
        for row_number in range(2, ws.max_row + 1):
            values = [json_safe_cell(ws.cell(row_number, c).value) for c in range(1, len(headers) + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            for index in range(min(11, len(values))):
                if values[index] not in (None, ""):
                    carried_values[index] = values[index]
                elif index in carried_values:
                    values[index] = carried_values[index]
            while len(values) < 14:
                values.append(None)
            rows.append({
                "rowNumber": row_number,
                "values": values,
                "m": str(values[12] or "").strip(),
                "n": str(values[13] or "").strip(),
            })
        sheets.append({"name": ws.title, "headers": headers, "rows": rows})
    return {"fileName": Path(excel_path).name, "excelPath": str(excel_path), "sheets": sheets}


def safe_uploaded_excel_path(path_value):
    path = Path(str(path_value or "")).resolve()
    upload_root = UPLOAD_DIR.resolve()
    if upload_root not in path.parents or path.suffix.lower() != ".xlsx" or not path.exists():
        raise ValueError("规则 Excel 文件不存在或路径无效，请重新上传。")
    return path


def export_rule_annotation_workbook(excel_path, sheet_name, edits, new_rules, m_column=13, n_column=14):
    m_column = int(m_column)
    n_column = int(n_column)
    if m_column < 1 or n_column < 1 or m_column == n_column:
        raise ValueError("M列和N列必须选择不同的列。")
    workbook = openpyxl.load_workbook(excel_path)
    ws = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
    for item in edits or []:
        row_number = int(item.get("rowNumber", 0))
        if row_number < 2 or row_number > ws.max_row:
            continue
        ws.cell(row_number, m_column).value = str(item.get("m") or "").strip()
        ws.cell(row_number, n_column).value = str(item.get("n") or "").strip()
    for item in new_rules or []:
        target_sheet = str(item.get("sheetName") or ws.title)
        target = workbook[target_sheet] if target_sheet in workbook.sheetnames else ws
        values = list(item.get("values") or [])
        if not any(str(value or "").strip() for value in values):
            continue
        while len(values) < max(14, m_column, n_column):
            values.append("")
        requested_row = int(item.get("insertRow") or 0)
        row_number = requested_row if 2 <= requested_row <= target.max_row + 1 else target.max_row + 1
        if row_number <= target.max_row:
            target.insert_rows(row_number, 1)
        style_source_row = row_number - 1 if row_number > 2 else row_number + 1
        if 1 <= style_source_row <= target.max_row:
            for column in range(1, min(target.max_column, max(14, m_column, n_column)) + 1):
                source_cell = target.cell(style_source_row, column)
                target_cell = target.cell(row_number, column)
                if source_cell.has_style:
                    target_cell._style = copy.copy(source_cell._style)
                if source_cell.number_format:
                    target_cell.number_format = source_cell.number_format
                if source_cell.alignment:
                    target_cell.alignment = copy.copy(source_cell.alignment)
        for column, value in enumerate(values, start=1):
            target.cell(row_number, column).value = value
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / f"{RULE_ANNOTATION_OUTPUT_PREFIX}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    workbook.save(output_path)
    return output_path


def fill_down_rows(ws):
    # The workbook uses merged cells for rule metadata. Carry A:K forward so
    # each error-scene row remains a complete, independently matchable item;
    # L:N stay row-specific because they contain the error scene and gold values.
    current = {}
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {"excelRow": r}
        for c in range(1, min(11, ws.max_column) + 1):
            value = ws.cell(r, c).value
            if value not in (None, ""):
                current[c] = value
            row[c] = current.get(c)
        for c in range(12, ws.max_column + 1):
            row[c] = ws.cell(r, c).value
        rows.append(row)
    return rows


def material_score(row, candidate):
    return max(
        ratio(row.get(4), candidate["materialName"]),
        ratio(row.get(5), candidate["outlineName"]),
        ratio(row.get(5), candidate["materialName"]),
        ratio(row.get(4), candidate["outlineName"]),
    )


def score_pair(row, candidate):
    material_score_value = material_score(row, candidate)
    rule_score = ratio(row.get(8), candidate["ruleName"])
    score = 0.36 * material_score_value + 0.64 * rule_score
    if norm_text(row.get(8)) == norm_text(candidate["ruleName"]):
        score += 0.08
    joined_material = candidate["materialName"] + candidate["outlineName"]
    if norm_text(row.get(4)) and norm_text(row.get(4)) in norm_text(joined_material):
        score += 0.04
    return min(score, 1)


def top_options(row, candidates, limit=20):
    scored = []
    for candidate in candidates:
        scored.append((score_pair(row, candidate), candidate["id"]))
    scored.sort(key=lambda item: (-item[0], item[1]))
    return scored[:limit]


def best_material_group(row, candidates):
    grouped = {}
    for candidate in candidates:
        key = candidate.get("materialKey") or ""
        score = material_score(row, candidate)
        current = grouped.get(key)
        if current is None or score > current["score"]:
            grouped[key] = {
                "key": key,
                "score": score,
                "count": int(candidate.get("materialOccurrenceTotal") or 1),
            }
    return max(grouped.values(), key=lambda item: (item["score"], item["count"]), default={"key": "", "score": 0, "count": 1})


def expand_gold_rows(source_rows, candidates):
    """Collapse alternative error-scenario rows into one detection item per rule."""
    groups = []
    for source_row in source_rows:
        group_key = tuple(norm_text(source_row.get(column)) for column in range(1, 12))
        if (
            groups
            and groups[-1]["key"] == group_key
            and source_row["excelRow"] == groups[-1]["rows"][-1]["excelRow"] + 1
        ):
            groups[-1]["rows"].append(source_row)
        else:
            groups.append({"key": group_key, "rows": [source_row]})

    expanded = []
    for detection_group in groups:
        scenario_rows = detection_group["rows"]
        source_row = scenario_rows[0]
        material_group = best_material_group(source_row, candidates)
        copy_count = max(
            max(
                len(split_copy_values(row.get(13)))
                for row in scenario_rows
            ),
            material_group.get("count", 1),
            1,
        )
        start_row = scenario_rows[0]["excelRow"]
        end_row = scenario_rows[-1]["excelRow"]
        rule_excel_row = start_row
        scenario_metadata = [
            {
                "sourceExcelRow": row["excelRow"],
                "errorScene": row.get(12) or "",
                "goldTrigger": row.get(13) or "",
                "goldRaw": row.get(14) or "",
            }
            for row in scenario_rows
        ]
        for copy_index in range(1, copy_count + 1):
            parts = []
            for scenario_row in scenario_rows:
                trigger_values = split_copy_values(scenario_row.get(13))
                correct_values = split_copy_values(scenario_row.get(14), allow_text_parts=len(trigger_values) > 1)
                trigger = trigger_values[copy_index - 1] if copy_index <= len(trigger_values) else ""
                correct = correct_values[copy_index - 1] if copy_index <= len(correct_values) else ""
                if trigger or correct:
                    parts.append({"row": scenario_row, "trigger": trigger, "correct": correct})
            selected = next((part for part in parts if normalize_gold_trigger(part["trigger"]) == "是"), None)
            selected = selected or next((part for part in parts if normalize_gold_trigger(part["trigger"]) == "否"), None)
            selected = selected or (parts[0] if parts else {"row": source_row, "trigger": "", "correct": ""})
            row = copy.deepcopy(selected["row"])
            has_trigger_value = any(part["trigger"] for part in parts)
            has_correct_value = any(part["correct"] for part in parts)
            row["sourceExcelRow"] = selected["row"]["excelRow"]
            row["sourceExcelRows"] = [item["sourceExcelRow"] for item in scenario_metadata]
            row["sourceRowStart"] = start_row
            row["sourceRowEnd"] = end_row
            row["sourceRowLabel"] = str(start_row) if start_row == end_row else f"{start_row}-{end_row}"
            row["ruleExcelRow"] = rule_excel_row
            row["scenarioRows"] = scenario_metadata
            row["copyIndex"] = copy_index if copy_count > 1 else 0
            row["copyCount"] = copy_count
            row["expectedMaterialKey"] = material_group.get("key", "")
            row["copyPending"] = copy_count > 1 and (not has_trigger_value or not has_correct_value)
            row[13] = selected["trigger"] if has_trigger_value else ""
            row[14] = selected["correct"] if has_correct_value else ""
            if copy_count > 1:
                row["rowKey"] = f"{start_row}-{end_row}#{copy_index}"
            else:
                row["rowKey"] = str(start_row) if start_row == end_row else f"{start_row}-{end_row}"
            expanded.append(row)
    return expanded


def build_analysis(json_path, excel_path, match_threshold=0.67, review_threshold=0.82, progress=None):
    def report(percent, message, current=None, total=None):
        if progress:
            progress(percent, message, current, total)

    report(8, "正在读取 JSON 规则...")
    candidates = load_json_candidates(json_path)
    report(18, f"已读取 {len(candidates)} 条 JSON 规则，正在读取 Excel...")
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    ws = workbook.active
    report(28, "正在展开 Excel 金标准行...")
    source_rows = expand_gold_rows(fill_down_rows(ws), candidates)
    total_rows = max(len(source_rows), 1)
    report(36, f"已展开 {len(source_rows)} 条检测项，正在计算匹配候选...", 0, len(source_rows))

    options_by_row = {}
    for index, row in enumerate(source_rows, start=1):
        options = top_options(row, candidates)
        if row.get("copyCount", 1) > 1 and row.get("expectedMaterialKey"):
            copy_candidates = [
                candidate for candidate in candidates
                if candidate.get("materialKey") == row.get("expectedMaterialKey")
                and int(candidate.get("materialOccurrence") or 1) == int(row.get("copyIndex") or 1)
            ]
            if copy_candidates:
                options = top_options(row, copy_candidates)
        options_by_row[row["rowKey"]] = options
        if index == len(source_rows) or index == 1 or index % 10 == 0:
            percent = 36 + int(34 * index / total_rows)
            report(percent, f"正在计算匹配候选：{index} / {len(source_rows)}", index, len(source_rows))

    used = set()
    rows = []
    report(72, "正在分配最佳匹配...", 0, len(source_rows))
    for index, row in enumerate(source_rows, start=1):
        row_options = options_by_row[row["rowKey"]]
        best_score, best_id = row_options[0] if row_options else (0, None)
        assigned_id = None
        assigned_score = best_score
        for option_score, option_id in row_options:
            if option_score >= match_threshold and option_id not in used:
                assigned_id = option_id
                assigned_score = option_score
                used.add(option_id)
                break
        candidate = candidates[assigned_id] if assigned_id is not None else None
        status = "auto"
        if row.get("copyPending"):
            status = "review"
        elif assigned_id is None:
            status = "unmatched"
        elif assigned_score < review_threshold:
            status = "review"
        expected_value = normalize_gold_trigger(row.get(13))
        expected = expected_value == "是"
        gold = gold_correct(row.get(14))
        aa = compute_aa(expected, gold, candidate)
        model_triggered = candidate_triggered(candidate)
        raw_material_type = row.get(4) or ""
        display_material_type = raw_material_type
        if row.get("copyCount", 1) > 1:
            pending_suffix = "·待确认" if row.get("copyPending") else ""
            display_material_type = f"{raw_material_type}（第{row.get('copyIndex')}份{pending_suffix}）" if raw_material_type else f"第{row.get('copyIndex')}份{pending_suffix}"
        rows.append(
            {
                "excelRow": row["rowKey"],
                "sourceExcelRow": row.get("sourceExcelRow", row["excelRow"]),
                "sourceExcelRows": row.get("sourceExcelRows", [row.get("sourceExcelRow", row["excelRow"])]),
                "sourceRowStart": row.get("sourceRowStart", row.get("sourceExcelRow", row["excelRow"])),
                "sourceRowEnd": row.get("sourceRowEnd", row.get("sourceExcelRow", row["excelRow"])),
                "sourceRowLabel": row.get("sourceRowLabel", str(row.get("sourceExcelRow", row["excelRow"]))),
                "ruleExcelRow": row.get("ruleExcelRow", row.get("sourceRowStart", row.get("sourceExcelRow", row["excelRow"]))),
                "scenarioRows": row.get("scenarioRows", []),
                "rowKey": row["rowKey"],
                "materialSerial": row.get(1) or "",
                "volume": row.get(2) or "",
                "issueType": row.get(3) or "",
                "materialType": display_material_type,
                "materialTypeRaw": raw_material_type,
                "materialName": row.get(5) or "",
                "materialCopyIndex": row.get("copyIndex", 0),
                "materialCopyCount": row.get("copyCount", 1),
                "materialCopyPending": bool(row.get("copyPending")),
                "stage": row.get(6) or "",
                "stageTimeOrder": row.get(7) or "",
                "ruleName": row.get(8) or "",
                "specialCondition": row.get(9) or "",
                "judgementRule": row.get(10) or "",
                "errorScene": row.get(12) or "",
                "goldTrigger": expected_value,
                "goldCorrect": "" if gold is None else ("是" if gold else "否"),
                "goldRaw": row.get(14) or "",
                "goldTriggerOriginal": expected_value,
                "goldCorrectOriginal": "" if gold is None else ("是" if gold else "否"),
                "goldRawOriginal": row.get(14) or "",
                "category": classify_rule(row.get(8)),
                "assignedId": assigned_id,
                "modelTriggered": model_triggered,
                "score": round(assigned_score, 4),
                "status": status,
                "aa": aa,
                "errorReason": row.get(28) or "",
                "options": [
                    {"candidateId": cid, "score": round(score, 4)}
                    for score, cid in row_options
                ],
            }
        )
        if index == len(source_rows) or index == 1 or index % 20 == 0:
            percent = 72 + int(18 * index / total_rows)
            report(percent, f"正在分配最佳匹配：{index} / {len(source_rows)}", index, len(source_rows))

    report(94, "正在汇总统计结果...")
    unused_candidates = [c for c in candidates if c["id"] not in used]
    summary = summarize(rows, candidates)
    report(100, "分析完成")
    return {
        "jsonPath": str(json_path),
        "excelPath": str(excel_path),
        "matchThreshold": match_threshold,
        "reviewThreshold": review_threshold,
        "summary": summary,
        "rows": rows,
        "candidates": candidates,
        "unusedCandidates": unused_candidates,
    }


def compute_aa(expected, gold, candidate):
    if gold is None:
        return "/"
    if candidate is not None and gold is not None:
        return "是" if bool(candidate["isCorrect"]) == bool(gold) else "否"
    if candidate is not None and not expected:
        return "否"
    if expected and candidate is None:
        return "否"
    return "/"


def candidate_triggered(candidate):
    return bool(candidate) and bool(candidate.get("modelTriggered", True))


def dimension_stats(rows, candidates):
    stats = {}
    for category in CATEGORY_NAMES:
        gold_rows = [
            row
            for row in rows
            if row["category"] == category
            and row["goldTrigger"] == "是"
            and not gold_row_pending(row)
        ]
        model_triggered = sum(1 for row in gold_rows if row.get("modelTriggered"))
        correct = sum(1 for row in gold_rows if row.get("modelTriggered") and row.get("aa") == "是")
        stats[category] = {
            "goldTriggered": len(gold_rows),
            "modelTriggered": model_triggered,
            "correct": correct,
            "recall": pct(model_triggered, len(gold_rows)),
            "accuracy": pct(correct, model_triggered),
            "hitRate": pct(correct, len(gold_rows)),
        }
    combined_gold = sum(stats[category]["goldTriggered"] for category in ("签名", "日期", "捺印", "印章"))
    combined_model = sum(stats[category]["modelTriggered"] for category in ("签名", "日期", "捺印", "印章"))
    combined_correct = sum(stats[category]["correct"] for category in ("签名", "日期", "捺印", "印章"))
    stats["综合"] = {
        "goldTriggered": combined_gold,
        "modelTriggered": combined_model,
        "correct": combined_correct,
        "recall": pct(combined_model, combined_gold),
        "accuracy": pct(combined_correct, combined_model),
        "hitRate": pct(combined_correct, combined_gold),
    }
    return stats


def error_reason_stats(rows, candidates):
    candidate_by_id = {candidate["id"]: candidate for candidate in candidates}
    grouped = {}
    for row in rows:
        if gold_row_pending(row):
            continue
        candidate = candidate_by_id.get(row.get("assignedId"))
        is_unmatched_error = not candidate and row.get("status") in {"unmatched", "confirmed_none"}
        is_model_error = candidate and candidate.get("isCorrect") is False
        if not is_unmatched_error and not is_model_error:
            continue
        reason = str(row.get("errorReason") or "").strip()
        if not reason:
            continue
        item = grouped.setdefault(reason, {"reason": reason, "count": 0, "ruleNames": set(), "materials": set()})
        item["count"] += 1
        if row.get("ruleName"):
            item["ruleNames"].add(str(row["ruleName"]))
        material_label = " / ".join(value for value in [display_material_type(row), row.get("materialName")] if value)
        if material_label:
            item["materials"].add(material_label)
    return [
        {
            "reason": item["reason"],
            "count": item["count"],
            "ruleNames": sorted(item["ruleNames"]),
            "materials": sorted(item["materials"]),
        }
        for item in sorted(grouped.values(), key=lambda value: (-value["count"], value["reason"]))
    ]


def collect_error_reason_details(reviewed):
    candidate_by_id = {candidate["id"]: candidate for candidate in reviewed.get("candidates", [])}
    analysis_source = reviewed.get("excelFileName") or Path(reviewed.get("excelPath") or "").name
    details = []
    for row in reviewed.get("rows", []):
        if gold_row_pending(row):
            continue
        candidate = candidate_by_id.get(row.get("assignedId"))
        is_unmatched_error = not candidate and row.get("status") in {"unmatched", "confirmed_none"}
        is_model_error = candidate and candidate.get("isCorrect") is False
        reason = str(row.get("errorReason") or "").strip()
        if (not is_unmatched_error and not is_model_error) or not reason:
            continue
        details.append({
            "错误原因": reason,
            "来源文件": analysis_source,
            "材料类型": display_material_type(row),
            "对应材料名称": row.get("materialName", ""),
            "所属阶段": row.get("stage", ""),
            "阶段时间顺序": row.get("stageTimeOrder", ""),
            "规则名称": row.get("ruleName", ""),
            "特殊条件": row.get("specialCondition", ""),
            "判断规则": row.get("judgementRule", ""),
            "错误场景": row.get("errorScene", ""),
            "是否触发": row.get("goldTrigger", ""),
            "是否正确": row.get("goldRaw", ""),
            "模型是否触发": "是" if row.get("modelTriggered") else "否",
            "模型是否正确": row.get("aa", ""),
            "具体问题内容": (candidate or {}).get("questionDesc") or "未匹配到对应的 JSON 检测项",
        })
    grouped = {}
    order = []
    for item in details:
        reason = item["错误原因"]
        if reason not in grouped:
            grouped[reason] = []
            order.append(reason)
        grouped[reason].append(item)
    return [item for reason in order for item in grouped[reason]]


def write_error_reason_sheet(ws, reviewed, exported_at=None):
    exported_at = exported_at or time.strftime("%Y-%m-%d %H:%M:%S")
    ws.append([f"错误原因明细（导出时间：{exported_at}）"] + [""] * (len(ERROR_REASON_HEADERS) - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(ERROR_REASON_HEADERS))
    ws.append(list(ERROR_REASON_HEADERS))
    details = collect_error_reason_details(reviewed)
    for item in details:
        ws.append([item.get(header, "") for header in ERROR_REASON_HEADERS])
    header_fill = PatternFill("solid", fgColor="C55A11")
    header_font = Font(color="FFFFFF", bold=True)
    title_fill = PatternFill("solid", fgColor="1F4E78")
    title_font = Font(color="FFFFFF", bold=True, size=14)
    thin = Side(style="thin", color="D9E2F3")
    widths = [32, 30, 22, 38, 18, 12, 38, 18, 58, 45, 12, 20, 14, 14, 72]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
    start = 3
    while start <= ws.max_row:
        reason = ws.cell(start, 1).value
        end = start
        while end < ws.max_row and ws.cell(end + 1, 1).value == reason:
            end += 1
        if end > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        start = end + 1
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:O{max(2, ws.max_row)}"
    ws.sheet_view.showGridLines = False


def write_error_reason_workbook(reviewed, output_path):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "错误原因明细"
    write_error_reason_sheet(ws, reviewed)
    workbook.save(output_path)


def source_excel_header(reviewed, column, fallback):
    excel_path = reviewed.get("excelPath") or ""
    if excel_path and Path(excel_path).exists():
        try:
            workbook = openpyxl.load_workbook(excel_path, data_only=True, read_only=True)
            value = workbook.active.cell(1, column).value
            workbook.close()
            if value not in (None, ""):
                return str(value).strip()
        except (OSError, ValueError, TypeError):
            pass
    return fallback


def write_history_review_sheet(ws, reviewed, summary, exported_at=None):
    """Write the detailed review sheet requested for each history record."""
    exported_at = exported_at or time.strftime("%Y-%m-%d %H:%M:%S")
    summary = summary or reviewed.get("summary") or {}
    dimension_summary = summary.get("dimensionStats") or {}
    trigger_header = source_excel_header(
        reviewed,
        13,
        f"{reviewed.get('excelFileName') or 'Excel'}-是否触发",
    )
    correct_header = source_excel_header(
        reviewed,
        14,
        f"{reviewed.get('excelFileName') or 'Excel'}-是否正确",
    )
    recall_text = pct_text(summary.get("recall", 0))
    if "召回率：%" in correct_header:
        correct_header = correct_header.replace("召回率：%", f"召回率：{recall_text}")
    elif "召回率" not in correct_header:
        correct_header = f"{correct_header}\n召回率：{recall_text}"

    signature = dimension_item(dimension_summary, "签名")
    date = dimension_item(dimension_summary, "日期")
    print_mark = dimension_item(dimension_summary, "捺印")
    seal = dimension_item(dimension_summary, "印章")
    combined = dimension_item(dimension_summary, "综合")
    other = dimension_item(dimension_summary, "其他")
    headers = [
        "材料序号", "所在卷宗", "问题类型", "材料类型", "对应材料名称", "所属阶段", "阶段时间顺序",
        "规则名称", "特殊条件", "判断规则", "错误场景", trigger_header, correct_header,
        dimension_header("签名", signature), dimension_issue_header("签名", signature),
        dimension_header("日期", date), dimension_issue_header("日期", date),
        dimension_header("捺印", print_mark), dimension_issue_header("捺印", print_mark),
        dimension_header("印章", seal), dimension_issue_header("印章", seal),
        dimension_header("综合", combined),
        dimension_header("其他", other),
        dimension_issue_header("其他", other),
        f"复核统计\n触发率：{pct_text(summary.get('triggerRate', 0))}\n模型触发且金标准应触发/金标准应触发={summary.get('expectedTriggered', 0)}/{summary.get('expectedTotal', 0)}",
        f"复核统计\n准确率：{pct_text(summary.get('accuracy', 0))}\n模型实际检测：是：{summary.get('modelTriggered', 0)}，其中错误：{summary.get('modelError', 0)}",
        "错误原因",
    ]
    title = f"复核记录（导出时间：{exported_at}）"
    ws.append([title] + [""] * (len(headers) - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.append(headers)

    category_columns = {"签名": (14, 15), "日期": (16, 17), "捺印": (18, 19), "印章": (20, 21)}
    for row in reviewed.get("rows", []):
        values = [
            row.get("materialSerial", ""), row.get("volume", ""), row.get("issueType", ""),
            display_material_type(row), row.get("materialName", ""), row.get("stage", ""),
            row.get("stageTimeOrder", ""), row.get("ruleName", ""), row.get("specialCondition", ""),
            row.get("judgementRule", ""), row.get("errorScene", ""), row.get("goldTrigger", ""),
            row.get("goldRaw", ""),
        ]
        category = row.get("category")
        for name in ("签名", "日期", "捺印", "印章"):
            values.extend([
                1 if category == name else None,
                (1 if row.get("aa") == "否" else 0) if category == name else None,
            ])
        values.extend([
            1 if category in {"签名", "日期", "捺印", "印章"} else None,
            1 if category == "其他" else None,
            1 if category == "其他" and row.get("aa") == "否" else 0 if category == "其他" else None,
            "是" if row.get("modelTriggered") else "否",
            row.get("aa", ""),
            row.get("errorReason", ""),
        ])
        ws.append(values)

    title_fill = PatternFill("solid", fgColor="17365D")
    title_font = Font(color="FFFFFF", bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    widths = [12, 16, 18, 18, 32, 18, 16, 38, 30, 52, 52, 42, 34, 22, 16, 22, 16, 22, 16, 22, 16, 30, 22, 16, 34, 34, 32]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for cell in ws[1]:
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 96
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:AA{max(2, ws.max_row)}"
    ws.sheet_view.showGridLines = False


def write_history_export_workbook(records, output_path, exported_at=None):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "比对历史"
    exported_at = exported_at or time.strftime("%Y-%m-%d %H:%M:%S")
    dimension_headers = []
    for category in ["签名", "日期", "捺印", "印章", "综合", "其他"]:
        dimension_headers.extend([f"{category}召回率", f"{category}准确率", f"{category}综合命中率"])
    headers = [
        "比对名称", "Excel 文件名", "保存时间", "阅卷时长",
        *dimension_headers, "触发率", "召回率", "准确率",
    ]
    ws.append([f"比对历史（导出时间：{exported_at}）"] + [""] * (len(headers) - 1))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.append(headers)
    for record in records:
        summary = record.get("summary") or {}
        dimension_stats = summary.get("dimensionStats") or {}
        dimension_values = []
        for category in ["签名", "日期", "捺印", "印章", "综合", "其他"]:
            item = dimension_stats.get(category) or {}
            dimension_values.extend([
                pct_text(item.get("recall", 0)),
                pct_text(item.get("accuracy", 0)),
                pct_text(item.get("hitRate", 0)),
            ])
        ws.append([
            record.get("title") or "未命名比对",
            record.get("excelFileName") or Path(record.get("excelPath") or "").name,
            record.get("createdAt") or "",
            record.get("modelReviewDurationText") or "未填写",
            *dimension_values,
            pct_text(summary.get("triggerRate", 0)),
            pct_text(summary.get("recall", 0)),
            pct_text(summary.get("accuracy", 0)),
        ])
    title_fill = PatternFill("solid", fgColor="17365D")
    title_font = Font(color="FFFFFF", bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    widths = [30, 32, 22, 18] + [14] * len(dimension_headers) + [14, 14, 14]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{max(2, ws.max_row)}"
    ws.sheet_view.showGridLines = False
    for index, record in enumerate(records, start=1):
        prefix = f"历史{index:02d}"
        reviewed = record.get("reviewed") or {}
        summary_sheet = workbook.create_sheet(f"{prefix}_统计")
        write_summary_sheet(summary_sheet, record.get("summary") or reviewed.get("summary") or {}, exported_at)
        review_sheet = workbook.create_sheet(f"{prefix}_复核")
        write_history_review_sheet(review_sheet, reviewed, record.get("summary") or reviewed.get("summary") or {}, exported_at)
        error_sheet = workbook.create_sheet(f"{prefix}_错误原因")
        write_error_reason_sheet(error_sheet, reviewed, exported_at)
    workbook.save(output_path)


def summarize(rows, candidates):
    eligible_rows = [row for row in rows if not gold_row_pending(row)]
    expected_total = sum(1 for row in eligible_rows if row["goldTrigger"] == "是")
    expected_triggered = sum(1 for row in eligible_rows if row["goldTrigger"] == "是" and row.get("modelTriggered"))
    model_triggered = sum(1 for row in eligible_rows if row.get("modelTriggered"))
    model_correct = sum(1 for row in eligible_rows if row.get("modelTriggered") and row["aa"] == "是")
    model_error = sum(1 for row in eligible_rows if row.get("modelTriggered") and row["aa"] == "否")
    correct_recalled = sum(1 for row in eligible_rows if row["goldTrigger"] == "是" and row.get("modelTriggered") and row["aa"] == "是")
    error_total = sum(1 for row in eligible_rows if row["goldTrigger"] == "是" and row["goldCorrect"] != "是")
    error_recalled = sum(
        1
        for row in eligible_rows
        if row["goldTrigger"] == "是" and row["goldCorrect"] != "是" and row.get("modelTriggered")
    )
    missed_rows = sum(1 for row in eligible_rows if row["goldTrigger"] == "是" and not row.get("modelTriggered"))
    by_category = {}
    for category in ["签名", "日期", "捺印", "印章", "其他"]:
        total = sum(1 for row in eligible_rows if row["category"] == category)
        errors = sum(1 for row in eligible_rows if row["category"] == category and row["aa"] == "否")
        by_category[category] = {
            "total": total,
            "errors": errors,
            "accuracy": 1 - pct(errors, total),
        }
    comp_total = sum(by_category[c]["total"] for c in ["签名", "日期", "捺印", "印章"])
    comp_err = sum(by_category[c]["errors"] for c in ["签名", "日期", "捺印", "印章"])
    by_category["综合"] = {"total": comp_total, "errors": comp_err, "accuracy": 1 - pct(comp_err, comp_total)}
    return {
        "excelRows": len(rows),
        "jsonRules": len(candidates),
        "matchedRows": model_triggered,
        "reviewRows": sum(1 for row in rows if row["status"] in {"review", "unmatched"}),
        "expectedTotal": expected_total,
        "expectedTriggered": expected_triggered,
        "modelTriggered": model_triggered,
        "modelCorrect": model_correct,
        "modelError": model_error,
        "correctRecalled": correct_recalled,
        "errorTotal": error_total,
        "errorRecalled": error_recalled,
        "missedRows": missed_rows,
        "recall": pct(error_recalled, error_total),
        "triggerRate": pct(expected_triggered, expected_total),
        "accuracy": pct(model_correct, model_triggered),
        "byCategory": by_category,
        "dimensionStats": dimension_stats(rows, candidates),
        "errorReasons": error_reason_stats(rows, candidates),
    }


def review_row_key(row):
    return str(row.get("rowKey") or row.get("excelRow"))


def apply_review_assignments(analysis, assignments, verdict_overrides=None, category_overrides=None, error_reason_overrides=None):
    verdict_overrides = verdict_overrides or {}
    category_overrides = category_overrides or {}
    error_reason_overrides = error_reason_overrides or {}
    candidate_by_id = {candidate["id"]: candidate for candidate in analysis["candidates"]}
    rows = copy.deepcopy(analysis["rows"])
    for row in rows:
        row["goldTrigger"] = normalize_gold_trigger(row.get("goldTrigger"))
        existing_gold_correct = normalize_gold_trigger(row.get("goldCorrect"))
        if existing_gold_correct:
            row["goldCorrect"] = existing_gold_correct
        else:
            raw_gold_correct = gold_correct(row.get("goldRaw"))
            row["goldCorrect"] = "" if raw_gold_correct is None else ("是" if raw_gold_correct else "否")
        key = review_row_key(row)
        assigned = assignments.get(key, row["assignedId"])
        if assigned in ("", None, "none", "null"):
            row["assignedId"] = None
            if key in assignments:
                row["status"] = "unmatched"
            continue
        assigned = int(assigned)
        # Keep every manual assignment. The review page counts each Excel row,
        # even when several rows intentionally point to the same JSON rule.
        row["assignedId"] = assigned
    for row in rows:
        key = review_row_key(row)
        override_category = category_overrides.get(key)
        if override_category in CATEGORY_NAMES:
            row["category"] = override_category
            row["categorySource"] = "人工"
        else:
            row["categorySource"] = "自动"
        if key in error_reason_overrides:
            row["errorReason"] = str(error_reason_overrides[key] or "").strip()
    for row in rows:
        candidate = candidate_by_id.get(row["assignedId"])
        row["modelTriggered"] = candidate_triggered(candidate)
        expected = row["goldTrigger"] == "是"
        gold = True if row["goldCorrect"] == "是" else False if row["goldCorrect"] == "否" else None
        row["aa"] = compute_aa(expected, gold, candidate)
        override = verdict_overrides.get(review_row_key(row))
        if override in {"是", "否"}:
            row["aa"] = override
            row["aaSource"] = "人工"
            row["status"] = "confirmed"
        else:
            row["aaSource"] = "自动"
    reviewed = copy.deepcopy(analysis)
    reviewed["rows"] = rows
    reviewed["summary"] = summarize(rows, analysis["candidates"])
    return reviewed


def write_reviewed_workbook(excel_path, output_path, reviewed):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(excel_path)
    ws = workbook.active
    exported_at = time.strftime("%Y-%m-%d %H:%M:%S")
    candidates = {candidate["id"]: candidate for candidate in reviewed["candidates"]}

    category_cols = {
        "签名": (15, 16),
        "日期": (17, 18),
        "捺印": (19, 20),
        "印章": (21, 22),
        "其他": (24, 25),
    }
    flag_fill = PatternFill("solid", fgColor="EAF4FF")
    err_fill = PatternFill("solid", fgColor="FCE4D6")
    ok_fill = PatternFill("solid", fgColor="E2F0D9")
    rows_by_source = defaultdict(list)
    for row in reviewed["rows"]:
        write_row = row.get("ruleExcelRow", row.get("sourceRowStart", row.get("sourceExcelRow", row.get("excelRow"))))
        rows_by_source[int(write_row)].append(row)
    for r, source_rows in rows_by_source.items():
        source_rows.sort(key=lambda item: int(item.get("materialCopyIndex") or 1))
        delimiter = "--" if len(source_rows) > 1 else ""
        scenario_rows = set()
        for row in source_rows:
            for source_row in row.get("sourceExcelRows") or []:
                try:
                    scenario_rows.add(int(source_row))
                except (TypeError, ValueError):
                    continue
        trigger_values = [row.get("goldTrigger") or "待确认" for row in source_rows]
        correct_values = [row.get("goldRaw") or "待确认" for row in source_rows]
        ws.cell(r, 13).value = delimiter.join(trigger_values) if delimiter else trigger_values[0]
        ws.cell(r, 14).value = delimiter.join(correct_values) if delimiter else correct_values[0]
        for source_row in scenario_rows:
            if source_row != r:
                ws.cell(source_row, 27).value = "/"
        for c in range(15, 26):
            ws.cell(r, c).value = None
        category_counts = defaultdict(lambda: {"total": 0, "errors": 0})
        descriptions = []
        reasons = []
        model_triggers = []
        model_answers = []
        for row in source_rows:
            category = row.get("category") or "其他"
            category_counts[category]["total"] += 1
            if row.get("aa") == "否":
                category_counts[category]["errors"] += 1
            candidate = candidates.get(row.get("assignedId"))
            has_error_reason = (candidate and candidate.get("isCorrect") is False) or (not candidate and row.get("status") in {"unmatched", "confirmed_none"})
            if has_error_reason and row.get("errorReason"):
                reasons.append(str(row["errorReason"]))
            if candidate:
                descriptions.append(f"匹配规则：{candidate['ruleName']}；匹配分：{row['score']:.2f}；状态：{row['status']}")
            else:
                descriptions.append(f"未匹配；状态：{row['status']}")
            model_triggers.append("是" if row.get("modelTriggered") else "否")
            model_answers.append(row.get("aa") or "待确认")
        for category, (total_col, err_col) in category_cols.items():
            counts = category_counts.get(category, {"total": 0, "errors": 0})
            if counts["total"]:
                ws.cell(r, total_col).value = counts["total"]
                ws.cell(r, total_col).fill = flag_fill
                ws.cell(r, err_col).value = counts["errors"]
                ws.cell(r, err_col).fill = err_fill if counts["errors"] else ok_fill
        combined_total = sum(category_counts[c]["total"] for c in ("签名", "日期", "捺印", "印章"))
        other_total = category_counts["其他"]["total"]
        other_errors = category_counts["其他"]["errors"]
        ws.cell(r, 23).value = combined_total or None
        ws.cell(r, 24).value = other_total or None
        ws.cell(r, 25).value = other_errors if other_total else None
        ws.cell(r, 26).value = delimiter.join(model_triggers) if delimiter else model_triggers[0]
        ws.cell(r, 27).value = delimiter.join(model_answers) if delimiter else model_answers[0]
        ws.cell(r, 28).value = delimiter.join(reasons) if reasons else None
        ws.cell(r, 29).value = delimiter.join(descriptions) if delimiter else descriptions[0]

    summary = reviewed["summary"]
    dimension_summary = summary.get("dimensionStats") or {}
    signature = dimension_item(dimension_summary, "签名")
    date = dimension_item(dimension_summary, "日期")
    print_mark = dimension_item(dimension_summary, "捺印")
    seal = dimension_item(dimension_summary, "印章")
    comp = dimension_item(dimension_summary, "综合")
    other = dimension_item(dimension_summary, "其他")
    ws["O1"] = dimension_header("签名", signature)
    ws["P1"] = dimension_issue_header("签名", signature)
    ws["Q1"] = dimension_header("日期", date)
    ws["R1"] = dimension_issue_header("日期", date)
    ws["S1"] = dimension_header("捺印", print_mark)
    ws["T1"] = dimension_issue_header("捺印", print_mark)
    ws["U1"] = dimension_header("印章", seal)
    ws["V1"] = dimension_issue_header("印章", seal)
    ws["W1"] = dimension_header("综合", comp)
    ws["X1"] = dimension_header("其他", other)
    ws["Y1"] = dimension_issue_header("其他", other)
    ws["Z1"] = f"复核统计\n触发率：{pct_text(summary['triggerRate'])}\n模型触发且金标准应触发/金标准应触发={summary['expectedTriggered']}/{summary['expectedTotal']}"
    ws["AA1"] = f"复核统计\n准确率：{pct_text(summary['accuracy'])}\n模型实际检测：是：{summary['modelTriggered']}，其中错误：{summary['modelError']}"
    for c in range(15, 30):
        ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.cell(1, c).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 10, 14)
    ws.row_dimensions[1].height = 92
    ws.freeze_panes = "A2"

    if "统计汇总" in workbook.sheetnames:
        del workbook["统计汇总"]
    sheet = workbook.create_sheet("统计汇总", 0)
    write_summary_sheet(sheet, summary, exported_at)

    if "复核记录" in workbook.sheetnames:
        del workbook["复核记录"]
    review_sheet = workbook.create_sheet("复核记录")
    write_review_sheet(review_sheet, reviewed, candidates, exported_at)
    workbook.save(output_path)


def dimension_item(dimension_summary, category):
    data = (dimension_summary or {}).get(category) or {}
    return {
        "goldTriggered": int(data.get("goldTriggered") or 0),
        "modelTriggered": int(data.get("modelTriggered") or 0),
        "correct": int(data.get("correct") or 0),
        "recall": data.get("recall") or 0,
        "accuracy": data.get("accuracy") or 0,
        "hitRate": data.get("hitRate") or 0,
    }


def dimension_header(name, data):
    return (
        f"{name}\n"
        f"金标准触发数：{data['goldTriggered']}\n"
        f"模型触发数：{data['modelTriggered']}\n"
        f"判断正确数：{data['correct']}\n"
        f"召回率：{pct_text(data['recall'])}\n"
        f"准确率：{pct_text(data['accuracy'])}\n"
        f"综合命中率：{pct_text(data['hitRate'])}"
    )


def dimension_issue_header(name, data):
    return f"{name}未命中/误判（{max(0, data['goldTriggered'] - data['correct'])}）"


def write_summary_sheet(ws, summary, exported_at=None):
    rows = [
        ["指标", "分子", "分母", "结果", "口径"],
        ["召回率", summary["errorRecalled"], summary["errorTotal"], pct_text(summary["recall"]), "错误场景中模型触发的检测项 / 错误场景总数"],
        ["触发率", summary["expectedTriggered"], summary["expectedTotal"], pct_text(summary["triggerRate"]), "模型触发的金标准检测项 / 金标准要求触发的检测项"],
        ["准确率", summary["modelCorrect"], summary["modelTriggered"], pct_text(summary["accuracy"]), "模型判断正确的已触发检测项 / 模型已触发的检测项"],
    ]
    rows.append(["", "", "", "", "", "", "", ""])
    dimension_header_row = len(rows) + 1
    rows.append(["统计维度", "金标准触发数", "模型触发数", "判断正确数", "召回率", "准确率", "综合命中率", "计算公式"])
    dimension_stats = summary["dimensionStats"]
    for category in ["签名", "日期", "捺印", "印章", "综合", "其他"]:
        data = dimension_stats[category]
        rows.append([
            category,
            data["goldTriggered"],
            data["modelTriggered"],
            data["correct"],
            pct_text(data["recall"]),
            pct_text(data["accuracy"]),
            pct_text(data["hitRate"]),
            "召回率=模型触发数/金标准触发数；准确率=判断正确数/模型触发数；综合命中率=判断正确数/金标准触发数",
        ])
    rows.append(["", "", "", "", "", "", "", ""])
    error_reason_header_row = len(rows) + 1
    rows.append(["错误原因", "出现次数", "检测项", "材料", "说明"])
    for item in summary.get("errorReasons", []):
        rows.append([
            item["reason"],
            item["count"],
            "；".join(item["ruleNames"]),
            "；".join(item["materials"]),
            "来源：模型 isCorrect=false 行的人工填写",
        ])
    title = f"统计汇总（导出时间：{exported_at or time.strftime('%Y-%m-%d %H:%M:%S')}）"
    ws.append([title, "", "", "", "", "", "", ""])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="FFFFFF", bold=True, size=14)
    title_fill = PatternFill("solid", fgColor="17365D")
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = title_fill
        cell.font = title_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = header_font
    for cell in ws[dimension_header_row + 1]:
        cell.fill = PatternFill("solid", fgColor="70AD47")
        cell.font = header_font
    for cell in ws[error_reason_header_row + 1]:
        cell.fill = PatternFill("solid", fgColor="C55A11")
        cell.font = header_font
    for col, width in enumerate([22, 18, 16, 16, 14, 14, 14, 86], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"


def write_review_sheet(ws, reviewed, candidates, exported_at=None):
    title = f"复核记录（导出时间：{exported_at or time.strftime('%Y-%m-%d %H:%M:%S')}）"
    ws.append([title] + [""] * 15)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=16)
    ws.append(["Excel行", "状态", "匹配分", "材料类型", "对应材料名称", "Excel规则", "统计维度", "金标准触发", "金标准正确", "AA判定", "判定来源", "JSON材料", "JSON规则", "JSON是否正确", "错误原因", "模型描述"])
    for row in reviewed["rows"]:
        candidate = candidates.get(row["assignedId"])
        has_error_reason = (candidate and candidate.get("isCorrect") is False) or (not candidate and row.get("status") in {"unmatched", "confirmed_none"})
        ws.append([
            row.get("sourceRowLabel", row.get("sourceExcelRow", row["excelRow"])),
            row["status"],
            row["score"],
            display_material_type(row),
            row["materialName"],
            row["ruleName"],
            row["category"],
            row["goldTrigger"],
            row["goldRaw"],
            row["aa"],
            row.get("aaSource", "自动"),
            candidate["materialName"] if candidate else "",
            candidate["ruleName"] if candidate else "",
            "是" if candidate and candidate["isCorrect"] else "否" if candidate else "",
            row.get("errorReason", "") if has_error_reason else "",
            candidate["questionDesc"] if candidate else "",
        ])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True, size=14)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(color="FFFFFF", bold=True)
    widths = [10, 14, 10, 18, 24, 36, 12, 12, 18, 10, 12, 22, 36, 14, 28, 58]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"


HTML = r"""
<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>JSON 表格匹配复核工具</title>
  <style>
    :root { --bg:#ffffff; --surface:#ffffff; --soft:#f3f4f6; --ink:#111827; --muted:#4b5563; --line:#e5e7eb; --blue:#3b82f6; --blue-dark:#2563eb; --blue-soft:#eff6ff; --red:#b91c1c; --red-soft:#fef2f2; --amber:#b45309; --amber-soft:#fffbeb; --green:#047857; --green-soft:#ecfdf5; --shadow:none; }
    * { box-sizing: border-box; }
    body { margin:0; font-family:'Outfit', -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; color:var(--ink); background:var(--bg); }
    body::before { content:""; display:block; height:3px; background:linear-gradient(90deg, var(--blue), #60a5fa 48%, #93c5fd); opacity:.9; }
    body[data-theme="compact"] { font-size:13px; }
    header { padding:18px 22px; background:var(--surface); border-bottom:1px solid var(--line); position:relative; z-index:4; }
    .header-topline { display:flex; align-items:center; justify-content:space-between; gap:18px; margin-bottom:12px; }
    h1 { margin:0 0 12px; font-size:20px; letter-spacing:0; }
    .header-topline h1 { margin:0; }
    .controls { display:grid; grid-template-columns: minmax(300px, 1.2fr) minmax(260px, 1fr) 96px 96px 100px; gap:10px; align-items:end; margin-bottom:14px; }
    .controls > div { min-width:0; }
    .source-panel { display:grid; gap:10px; padding:12px; margin-bottom:12px; border:1px solid var(--line); border-radius:8px; background:var(--surface); }
    .source-tabs { display:flex; gap:6px; flex-wrap:wrap; }
    .source-tabs button { width:auto; background:var(--soft); color:var(--muted); border-color:var(--line); }
    .source-tabs button.active { background:var(--blue); color:#fff; border-color:var(--blue); }
    .api-source-grid { display:grid; grid-template-columns:minmax(260px, 1fr) minmax(170px, 240px) auto; gap:10px; align-items:end; }
    .batch-grid { display:grid; grid-template-columns:minmax(260px, 1fr) minmax(260px, .8fr) auto auto; gap:10px; align-items:end; }
    .batch-grid textarea { min-height:82px; resize:vertical; }
    .batch-result { border:1px solid var(--line); border-radius:8px; overflow:hidden; background:var(--surface); margin:10px 0 14px; }
    .batch-result-head { display:flex; align-items:center; justify-content:space-between; gap:10px; padding:10px 12px; border-bottom:1px solid var(--line); }
    .batch-result-title { font-weight:800; }
    .batch-result-actions { display:flex; gap:8px; align-items:center; flex-wrap:wrap; }
    .batch-result-actions button { width:auto; }
    .batch-summary { display:flex; flex-wrap:wrap; gap:8px; padding:10px 12px; border-bottom:1px solid var(--line); color:var(--muted); font-size:12px; }
    .batch-summary b { color:var(--ink); }
    .batch-table-wrap { max-height:300px; overflow:auto; }
    .batch-table th, .batch-table td { font-size:12px; padding:8px; }
    .dataset-actions { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .dataset-actions button { width:auto; }
    .dataset-current { display:flex; justify-content:space-between; gap:12px; align-items:center; padding:11px 12px; margin-bottom:12px; border:1px solid var(--line); border-radius:8px; background:var(--blue-soft); }
    .dataset-current-title { font-weight:800; }
    .dataset-current-meta { color:var(--muted); font-size:12px; margin-top:4px; word-break:break-all; }
    .json-preview { border:1px solid var(--line); border-radius:8px; background:var(--surface); overflow:hidden; }
    .json-preview-head { display:flex; justify-content:space-between; gap:10px; align-items:center; padding:10px 12px; border-bottom:1px solid var(--line); }
    .json-preview-actions { display:flex; gap:8px; flex-wrap:wrap; }
    .json-preview-actions button { width:auto; }
    .json-preview-meta { padding:8px 12px; color:var(--muted); font-size:12px; border-bottom:1px solid var(--line); word-break:break-all; }
    .json-preview pre { margin:0; padding:12px; height:calc(100vh - 390px); min-height:520px; max-height:none; overflow:auto; background:var(--soft); font-family:ui-monospace, SFMono-Regular, Menlo, Consolas, monospace; font-size:12px; line-height:1.5; white-space:pre-wrap; word-break:break-word; }
    label { display:block; font-size:12px; color:var(--muted); margin-bottom:4px; }
    input, select, button { width:100%; border:1px solid var(--line); border-radius:6px; padding:9px 10px; background:var(--surface); color:var(--ink); font-size:14px; }
    button { cursor:pointer; background:var(--blue); color:#fff; border-color:var(--blue); font-weight:600; }
    button.secondary { background:#fff; color:var(--blue); }
    button.mini { width:auto; padding:5px 8px; font-size:12px; }
    main { padding:18px 22px 40px; }
    .cards { display:grid; grid-template-columns: repeat(6, minmax(120px, 1fr)); gap:10px; margin-bottom:14px; }
    .card { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px; }
    .card .k { color:var(--muted); font-size:12px; }
    .card .v { font-size:22px; font-weight:700; margin-top:4px; }
    .card .formula { color:var(--muted); font-size:12px; line-height:1.45; margin-top:7px; }
    .card .formula code { font-size:12px; }
    .toolbar { display:flex; gap:10px; align-items:center; margin:12px 0; flex-wrap:wrap; }
    .toolbar button, .toolbar select, .toolbar input { width:auto; }
    .toolbar-actions { display:flex; align-items:center; gap:6px; flex-wrap:wrap; }
    .status { font-size:13px; color:var(--muted); }
    .analysis-progress { display:flex; align-items:center; gap:8px; min-width:min(420px, 100%); color:var(--muted); font-size:12px; }
    .analysis-progress.hide { display:none; }
    .analysis-progress-track { width:150px; height:8px; border-radius:999px; background:var(--soft); overflow:hidden; border:1px solid var(--line); }
    .analysis-progress-bar { width:0%; height:100%; border-radius:999px; background:var(--blue); transition:width .2s ease; }
    .analysis-progress-text { min-width:0; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
    .app-layout { display:grid; grid-template-columns:176px minmax(0,1fr); gap:16px; align-items:start; width:auto; min-width:0; }
    .app-layout.module-menu-collapsed { grid-template-columns:54px minmax(0,1fr); }
    .module-menu { position:sticky; top:18px; display:flex; flex-direction:column; gap:6px; padding:10px; background:var(--soft); border-radius:8px; }
    .module-menu-head { display:flex; align-items:center; justify-content:flex-end; gap:8px; padding:4px 4px 8px; color:var(--muted); font-size:12px; font-weight:800; }
    .module-menu-title { letter-spacing:.04em; }
    .module-tab { width:100%; display:flex; align-items:center; gap:7px; text-align:left; background:transparent; color:var(--muted); border:0; border-radius:6px; padding:9px 9px; white-space:nowrap; }
    .module-icon { width:19px; flex:0 0 19px; text-align:center; font-size:15px; }
    .module-tab:hover, .module-tab.active { color:var(--blue); background:var(--surface); }
    .module-tab.active { background:var(--blue); color:#fff; }
    .module-menu-toggle { width:40px; min-width:40px; min-height:34px; padding:0; background:transparent; color:var(--muted); font-size:22px; line-height:1; }
    .module-menu-toggle:hover { background:var(--surface); color:var(--ink); }
    .module-menu.collapsed .module-menu-title,
    .module-menu.collapsed .module-label,
    .module-menu.collapsed .module-menu-toggle span { display:none; }
    .module-menu.collapsed .module-menu-head { justify-content:center; }
    .module-menu.collapsed .module-tab { justify-content:center; padding-left:6px; padding-right:6px; }
    .module-menu.collapsed .module-menu-toggle { width:100%; padding-left:2px; padding-right:2px; }
    .volume-block { display:grid; gap:5px; }
    .volume-switch { width:100%; display:flex; align-items:center; justify-content:space-between; gap:8px; text-align:left; background:var(--surface); color:var(--ink); border:1px solid var(--line); border-radius:6px; padding:9px; font-weight:800; }
    .volume-switch.active { background:var(--blue); border-color:var(--blue); color:#fff; }
    .volume-switch .volume-arrow { color:inherit; font-size:14px; }
    .module-submenu { display:grid; gap:3px; padding:0 0 5px 8px; }
    .module-menu.collapsed .volume-switch { justify-content:center; padding-left:6px; padding-right:6px; }
    .module-menu.collapsed .volume-name,
    .module-menu.collapsed .volume-arrow,
    .module-menu.collapsed .module-submenu { display:none; }
    .module-content { min-width:0; width:100%; }
    .module-panel { min-width:0; animation:module-enter .28s ease both; }
    @keyframes module-enter { from { opacity:0; transform:translateY(8px); } to { opacity:1; transform:translateY(0); } }
    .analysis-empty { background:var(--blue-soft); border-left:5px solid var(--blue); border-radius:8px; padding:30px 32px; margin-bottom:14px; }
    .analysis-empty-kicker { color:var(--blue-dark); font-size:12px; font-weight:800; letter-spacing:.08em; text-transform:uppercase; }
    .analysis-empty h2 { margin:8px 0 6px; font-size:24px; }
    .analysis-empty p { margin:0; color:var(--muted); line-height:1.6; }
    .module-heading { display:flex; justify-content:space-between; gap:14px; align-items:flex-start; margin-bottom:12px; }
    .module-heading h2 { margin:0; font-size:18px; }
    .module-heading p { margin:5px 0 0; color:var(--muted); font-size:12px; line-height:1.45; }
    .case-panel { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; margin-bottom:14px; }
    .case-panel-head { display:flex; justify-content:space-between; gap:12px; align-items:flex-start; margin-bottom:10px; }
    .case-panel-title { font-weight:800; }
    .case-panel-subtitle { margin-top:3px; color:var(--muted); font-size:12px; line-height:1.45; }
    .case-fetch-grid { display:grid; grid-template-columns:minmax(260px, 1fr) minmax(280px, 1.2fr) auto; gap:10px; align-items:end; }
    .case-fetch-grid button { width:auto; white-space:nowrap; }
    .case-info-grid { display:grid; grid-template-columns:repeat(4, minmax(150px, 1fr)); gap:8px; margin-top:12px; }
    .case-info-item { padding:9px 10px; border:1px solid var(--line); border-radius:6px; background:var(--soft); min-width:0; }
    .case-info-label { color:var(--muted); font-size:11px; margin-bottom:3px; }
    .case-info-value { font-weight:800; line-height:1.4; word-break:break-word; }
    .case-template { margin-top:10px; padding:10px 12px; border-left:5px solid var(--blue); background:var(--blue-soft); border-radius:6px; line-height:1.55; }
    .case-template b { color:var(--blue-dark); }
    .identity-placeholder { padding:34px; border:1px dashed var(--line); border-radius:8px; color:var(--muted); text-align:center; background:var(--surface); }
    .annotation-heading { position:relative; overflow:hidden; padding:18px 20px; border:1px solid #dbe7f7; border-radius:8px; background:linear-gradient(110deg, #f8fbff, #eef5ff); }
    .annotation-heading::after { content:""; position:absolute; right:28px; top:-36px; width:110px; height:110px; border:1px solid rgba(59,130,246,.18); border-radius:50%; box-shadow:0 0 0 14px rgba(59,130,246,.04), 0 0 0 28px rgba(59,130,246,.025); }
    .annotation-heading .module-heading { position:relative; z-index:1; margin:0; }
    .annotation-kicker { margin-bottom:7px; color:var(--blue-dark); font-size:11px; font-weight:800; letter-spacing:.12em; }
    .annotation-controls, .annotation-toolbar, .annotation-table-pdf-layout, .annotation-layout { animation:content-rise .34s ease both; }
    @keyframes content-rise { from { opacity:0; transform:translateY(5px); } to { opacity:1; transform:translateY(0); } }
    .history-save { display:flex; justify-content:flex-end; align-items:center; min-width:min(760px, 100%); }
    .history-save button { width:auto; white-space:nowrap; min-height:38px; }
    .history-toolbar { display:flex; gap:8px; align-items:center; margin-bottom:10px; flex-wrap:wrap; }
    .history-toolbar button { width:auto; }
    .history-list { display:grid; gap:10px; }
    .history-item { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; }
    .history-item-head { display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    .history-item-title { font-weight:800; line-height:1.4; }
    .history-item-time { color:var(--muted); font-size:12px; white-space:nowrap; }
    .history-files { margin-top:6px; color:var(--muted); font-size:12px; line-height:1.5; word-break:break-all; }
    .history-review-time { margin-top:6px; color:var(--muted); font-size:12px; line-height:1.5; }
    .history-select { width:18px; height:18px; flex:0 0 auto; accent-color:var(--blue); }
    .history-title-wrap { display:flex; align-items:flex-start; gap:9px; min-width:0; }
    .history-metrics { display:flex; flex-wrap:wrap; gap:8px 16px; margin-top:9px; color:var(--muted); font-size:12px; }
    .history-metrics b { color:var(--ink); }
    .history-actions { display:flex; gap:7px; flex-wrap:wrap; margin-top:10px; }
    .history-actions button { width:auto; padding:6px 9px; font-size:12px; }
    .history-empty { background:var(--surface); border:1px dashed var(--line); border-radius:8px; padding:28px; text-align:center; color:var(--muted); }
    .history-pagination { display:flex; justify-content:center; align-items:center; gap:10px; margin-top:12px; }
    .history-pagination button { width:auto; min-width:78px; }
    .history-page-status { min-width:130px; text-align:center; color:var(--muted); font-size:13px; }
    .history-detail { margin-top:14px; }
    .history-detail-title { font-weight:800; margin-bottom:8px; }
    .history-detail .dimension-panel, .history-detail .error-reason-panel { margin-bottom:10px; }
    .annotation-controls { display:grid; grid-template-columns:minmax(280px,1fr) minmax(280px,1fr) auto auto; gap:10px; align-items:end; margin-bottom:12px; }
    .annotation-toolbar { display:flex; align-items:center; gap:8px; flex-wrap:wrap; margin:10px 0 12px; }
    .annotation-toolbar button { width:auto; }
    .annotation-controls button { width:auto; white-space:nowrap; }
    .annotation-layout { display:grid; grid-template-columns:minmax(230px,270px) minmax(300px,.72fr) minmax(430px,1.35fr); gap:12px; min-height:calc(100vh - 250px); }
    .annotation-pane { min-width:0; background:var(--surface); border:1px solid var(--line); border-radius:8px; overflow:hidden; }
    .annotation-pane-head { display:flex; align-items:center; justify-content:space-between; gap:8px; padding:11px 12px; border-bottom:1px solid var(--line); font-weight:800; }
    .annotation-pane-body { padding:10px; }
    .annotation-list-stats { display:grid; gap:6px; margin-bottom:9px; padding:9px 10px; border:1px solid var(--line); border-radius:6px; background:var(--soft); font-size:12px; line-height:1.45; }
    .annotation-list-stats b { color:var(--ink); }
    .annotation-filter { display:flex; gap:5px; flex-wrap:wrap; margin-bottom:9px; }
    .annotation-filter button { width:auto; flex:1 1 auto; min-width:48px; padding:6px 7px; background:var(--surface); color:var(--muted); font-size:12px; }
    .annotation-filter button.active { background:var(--blue); border-color:var(--blue); color:#fff; }
    .annotation-file-nav { display:flex; align-items:center; gap:7px; flex-wrap:wrap; margin:10px 0 12px; padding:9px 10px; border:1px solid var(--line); border-radius:7px; background:var(--surface); }
    .annotation-file-nav button { width:auto; min-width:82px; padding:7px 9px; }
    .annotation-file-nav select { width:min(360px, 42vw); min-width:180px; }
    .annotation-file-status { display:inline-flex; align-items:center; gap:5px; padding:5px 8px; border-radius:5px; background:var(--soft); color:var(--muted); font-size:12px; font-weight:700; }
    .annotation-file-status.completed { background:var(--green-soft); color:var(--green); }
    .annotation-file-status.progress { background:var(--amber-soft); color:var(--amber); }
    .annotation-file-status.issue { background:var(--red-soft); color:var(--red); }
    .annotation-file-status.pending { background:var(--soft); color:var(--muted); }
    .annotation-rule-list { max-height:calc(100vh - 315px); overflow:auto; }
    .annotation-rule-item { display:block; width:100%; text-align:left; padding:9px 10px; border:0; border-bottom:1px solid var(--line); border-radius:0; background:transparent; color:var(--ink); font-weight:500; }
    .annotation-rule-item:hover, .annotation-rule-item.active { background:var(--blue-soft); color:var(--blue-dark); }
    .annotation-rule-item small { display:block; margin-top:4px; color:var(--muted); font-weight:400; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
    .annotation-group-title { padding:9px 10px 6px; background:var(--soft); color:var(--muted); font-size:12px; font-weight:800; border-bottom:1px solid var(--line); }
    .annotation-material-group { border-bottom:1px solid var(--line); }
    .annotation-material-toggle { display:flex; align-items:center; justify-content:space-between; gap:8px; width:100%; padding:11px 10px; border:0; border-radius:0; background:var(--soft); color:var(--ink); text-align:left; font-weight:800; }
    .annotation-material-toggle:hover, .annotation-material-toggle.expanded { background:var(--blue-soft); color:var(--blue-dark); }
    .annotation-material-toggle .material-toggle-label { min-width:0; overflow-wrap:anywhere; }
    .annotation-material-toggle .material-toggle-meta { flex:0 0 auto; color:var(--muted); font-size:12px; font-weight:700; white-space:nowrap; }
    .annotation-material-toggle .material-toggle-meta.completed { color:var(--green); }
    .annotation-material-toggle .material-toggle-meta.progress { color:var(--amber); }
    .annotation-material-items { background:var(--surface); }
    .annotation-pdf { height:calc(100vh - 315px); min-height:420px; background:#f2f4f7; }
    .annotation-pdf iframe { width:100%; height:100%; border:0; background:#fff; }
    .annotation-empty { display:flex; align-items:center; justify-content:center; height:100%; min-height:420px; padding:24px; text-align:center; color:var(--muted); }
    .annotation-meta { display:grid; gap:8px; max-height:calc(100vh - 520px); overflow:auto; }
    .annotation-meta-item { padding-bottom:8px; border-bottom:1px solid var(--line); }
    .annotation-meta-label { color:var(--muted); font-size:11px; margin-bottom:3px; }
    .annotation-meta-value { white-space:pre-wrap; word-break:break-word; line-height:1.45; }
    .annotation-decision { margin-top:12px; padding-top:12px; border-top:1px solid var(--line); }
    .annotation-choice { display:flex; gap:8px; margin:6px 0 10px; }
    .annotation-choice button { width:auto; min-width:72px; background:var(--surface); color:var(--ink); }
    .annotation-choice button.active { background:var(--blue); color:#fff; }
    .annotation-decision textarea { min-height:94px; resize:vertical; }
    .annotation-new-rule { margin-top:12px; padding:12px; border:1px solid var(--line); border-radius:8px; background:var(--soft); }
    .annotation-new-grid { display:grid; grid-template-columns:repeat(2, minmax(0,1fr)); gap:8px; }
    .annotation-new-grid textarea { min-height:70px; resize:vertical; }
    .annotation-new-actions { display:flex; justify-content:flex-end; gap:8px; margin-top:10px; }
    .annotation-new-actions button { width:auto; }
    .annotation-table-pdf-layout { display:grid; grid-template-columns:minmax(300px,.72fr) minmax(0,1.35fr); gap:12px; min-height:calc(100vh - 315px); }
    .annotation-table-pdf-layout > .annotation-pdf { order:-1; }
    .annotation-table-nav { display:flex; align-items:center; gap:5px; min-width:0; }
    .annotation-table-nav button { width:auto; min-width:84px; padding:5px 8px; }
    .annotation-table-nav select { width:min(220px, 32vw); padding:6px 7px; font-size:12px; }
    .annotation-excel-table-wrap { height:calc(100vh - 315px); min-height:420px; overflow-y:scroll; overflow-x:scroll; scrollbar-gutter:stable both-edges; border:1px solid var(--line); background:var(--surface); }
    .annotation-excel-table-wrap::-webkit-scrollbar { width:12px; height:12px; }
    .annotation-excel-table-wrap::-webkit-scrollbar-track { background:#eef1f5; }
    .annotation-excel-table-wrap::-webkit-scrollbar-thumb { background:#9aa8b8; border:3px solid #eef1f5; border-radius:8px; }
    .annotation-excel-table { border-collapse:collapse; min-width:1900px; width:max-content; table-layout:fixed; }
    .annotation-excel-table th, .annotation-excel-table td { border:1px solid #aeb8c4; padding:7px 8px; vertical-align:top; white-space:pre-wrap; word-break:break-word; line-height:1.35; font-size:12px; }
    .annotation-excel-table th { position:sticky; top:0; z-index:2; background:#b4c7e7; color:#111827; text-align:center; font-weight:800; }
    .annotation-excel-table td:nth-child(-n+7) { background:#f7f8fa; }
    .annotation-excel-table td:nth-child(13), .annotation-excel-table td:nth-child(14) { background:#fff2cc; text-align:center; vertical-align:middle; }
    .annotation-excel-table tr.selected td { outline:2px solid var(--blue); outline-offset:-2px; }
    .annotation-excel-table .table-row-button { display:block; width:100%; min-height:44px; padding:0; border:0; background:transparent; color:inherit; text-align:left; font-weight:500; }
    .annotation-excel-table .table-row-button:hover { color:var(--blue-dark); }
    .annotation-excel-table .table-choice { display:flex; gap:4px; justify-content:center; }
    .annotation-excel-table .table-choice button { width:auto; min-width:34px; padding:5px 7px; font-size:12px; background:#fff; color:var(--ink); }
    .annotation-excel-table .table-choice button.active { background:var(--blue); color:#fff; }
    .annotation-excel-table td:nth-child(13), .annotation-excel-table td:nth-child(14) { min-width:120px; }
    .annotation-excel-table td:nth-child(14) .table-choice { margin-bottom:6px; }
    .annotation-excel-table .table-n-input { width:100%; min-width:150px; min-height:42px; border:1px solid #c8d0da; padding:6px; resize:vertical; background:#fffdf2; }
    .workspace { display:grid; grid-template-columns:minmax(196px, clamp(196px, 18vw, 260px)) minmax(0,1fr); gap:14px; align-items:start; width:100%; min-width:0; }
    .workspace.materials-collapsed { grid-template-columns:54px minmax(0,1fr); }
    .theme-switcher { display:flex; gap:6px; align-items:center; }
    .theme-switcher button { width:auto; padding:7px 10px; background:var(--surface); color:var(--blue); }
    .theme-switcher button.active { background:var(--blue); color:#fff; }
    header .theme-switcher { gap:4px; padding:4px; background:rgba(255,255,255,.14); border:1px solid rgba(255,255,255,.28); border-radius:8px; }
    header .theme-switcher button { min-height:34px; padding:6px 10px; background:transparent; color:#fff; }
    header .theme-switcher button:hover { background:rgba(255,255,255,.18); }
    header .theme-switcher button.active { background:#fff; color:var(--blue-dark); }
    .info-popover { position:relative; }
    .info-icon { width:36px; min-height:36px; padding:0; border-radius:50%; font-family:ui-monospace, SFMono-Regular, Menlo, Consolas, monospace; font-size:16px; }
    .info-popover-content { position:absolute; left:0; top:calc(100% + 8px); z-index:12; width:min(420px, calc(100vw - 48px)); padding:14px 16px; background:var(--ink); color:#fff; border-radius:8px; line-height:1.55; font-size:12px; visibility:hidden; opacity:0; transform:translateY(-4px); transition:opacity .16s ease, transform .16s ease, visibility .16s ease; pointer-events:none; }
    .info-popover-content::before { content:""; position:absolute; left:12px; top:-6px; width:12px; height:12px; background:var(--ink); transform:rotate(45deg); }
    .info-popover:hover .info-popover-content,
    .info-popover:focus-within .info-popover-content { visibility:visible; opacity:1; transform:translateY(0); }
    .info-popover-content .formula-title { color:#fff; margin-bottom:6px; }
    .info-popover-content .formula-line { color:rgba(255,255,255,.86); }
    .info-popover-content .formula-line b { color:rgba(255,255,255,.94); font-weight:700; }
    .info-popover-content code { color:var(--ink); background:#fff; }
    .sidebar { background:var(--surface); border:1px solid var(--line); border-radius:8px; overflow:hidden; position:sticky; top:96px; height:calc(100vh - 120px); min-height:420px; display:flex; flex-direction:column; }
    .sidebar-head { padding:12px; border-bottom:1px solid var(--line); }
    .sidebar-title-row { display:flex; align-items:center; justify-content:space-between; gap:8px; }
    .sidebar-title { font-weight:800; font-size:15px; }
    .sidebar-toggle { width:40px; min-width:40px; min-height:32px; padding:0; white-space:nowrap; font-size:22px; line-height:1; }
    .sidebar-tools { display:grid; grid-template-columns:minmax(120px, .85fr) minmax(150px, 1.15fr); gap:8px; margin-top:10px; }
    .sidebar-filter { min-width:0; }
    .sidebar-filter label { margin-bottom:5px; font-weight:600; }
    .sidebar-filter select { width:100%; }
    .sidebar-filter input { width:100%; }
    .material-summary { margin-top:10px; padding:8px 10px; border-radius:6px; background:var(--blue-soft); color:var(--blue-dark); font-size:12px; font-weight:700; line-height:1.45; }
    .status-legend { display:flex; flex-wrap:wrap; gap:5px; margin-top:8px; }
    .status-legend-item, .material-status-chip { display:inline-flex; align-items:center; min-height:22px; padding:2px 7px; border-radius:5px; font-size:11px; font-weight:800; white-space:nowrap; }
    .status-legend-item.pending, .material-status-chip.pending { color:var(--amber); background:var(--amber-soft); }
    .status-legend-item.missed, .material-status-chip.missed { color:var(--red); background:var(--red-soft); }
    .status-legend-item.incomplete, .material-status-chip.incomplete { color:var(--muted); background:var(--soft); }
    .status-legend-item.completed, .material-status-chip.completed { color:var(--green); background:var(--green-soft); }
    .status-legend-item[data-tip] { position:relative; cursor:help; }
    .status-legend-item[data-tip]::after { content:attr(data-tip); position:absolute; left:0; top:calc(100% + 8px); z-index:8; width:230px; padding:8px 10px; border:1px solid var(--line); border-radius:6px; background:var(--ink); color:var(--surface); font-size:11px; font-weight:500; line-height:1.45; white-space:normal; opacity:0; visibility:hidden; transform:translateY(-3px); transition:opacity .16s ease, transform .16s ease, visibility .16s ease; pointer-events:none; }
    .status-legend-item[data-tip]::before { content:""; position:absolute; left:12px; top:calc(100% + 3px); z-index:9; width:9px; height:9px; background:var(--ink); transform:rotate(45deg); opacity:0; visibility:hidden; transition:opacity .16s ease, visibility .16s ease; }
    .status-legend-item[data-tip]:hover::after, .status-legend-item[data-tip]:focus-visible::after,
    .status-legend-item[data-tip]:hover::before, .status-legend-item[data-tip]:focus-visible::before { opacity:1; visibility:visible; transform:translateY(0); }
    .status-legend-item[data-tip]:focus-visible { outline:2px solid var(--blue); outline-offset:2px; }
    .status-legend-note { margin-top:7px; color:var(--muted); font-size:11px; line-height:1.45; }
    .toolbar-status-legend { margin-top:0; align-items:center; }
    .material-list { flex:1 1 auto; min-height:0; overflow:auto; }
    .workspace.materials-collapsed .sidebar-title,
    .workspace.materials-collapsed .sidebar-tools,
    .workspace.materials-collapsed .material-list { display:none; }
    .workspace.materials-collapsed .sidebar-head { padding:12px 10px; }
    .workspace.materials-collapsed .sidebar-title-row { justify-content:center; }
    .workspace.materials-collapsed .sidebar-toggle { width:40px; padding:0; }
    .material-item { width:100%; display:block; border:0; border-bottom:1px solid var(--line); border-radius:0; background:var(--surface); color:var(--ink); text-align:left; padding:10px 12px; font-weight:500; }
    .material-item:hover, .material-item.active { background:var(--soft); }
    .material-type { font-size:13px; font-weight:700; line-height:1.35; }
    .material-doc { margin-top:4px; font-size:12px; color:var(--muted); line-height:1.35; }
    .material-meta { margin-top:7px; color:var(--muted); font-size:12px; display:flex; gap:8px; flex-wrap:wrap; }
    .material-meta strong { color:var(--red); }
    .material-item-select { display:block; width:100%; min-height:0; padding:10px 12px 7px; border:0; border-radius:0; background:transparent; color:var(--ink); text-align:left; font-weight:500; }
    .material-item-select:hover { background:transparent; color:var(--ink); transform:none; }
    .material-item-footer { display:flex; align-items:center; justify-content:space-between; gap:8px; padding:0 12px 10px; }
    .material-file-status { display:flex; align-items:center; gap:6px; color:var(--muted); font-size:12px; font-weight:700; }
    .material-file-status.completed { color:var(--green); }
    .material-file-status.missed { color:var(--red); }
    .material-file-status.pending { color:var(--amber); }
    .material-submit { width:auto; min-height:30px; padding:5px 9px; font-size:12px; }
    .material-pagination { flex:0 0 auto; display:flex; align-items:center; justify-content:center; gap:8px; flex-wrap:wrap; padding:9px 10px; border-top:1px solid var(--line); background:var(--surface); }
    .material-pagination button { width:auto; min-height:32px; padding:5px 9px; }
    .material-page-status { min-width:130px; color:var(--muted); font-size:12px; text-align:center; }
    .review-panel { min-width:0; width:100%; }
    .review-table-wrap { width:100%; min-width:0; height:calc(100vh - 238px); min-height:320px; overflow-y:auto; overflow-x:scroll; overscroll-behavior:contain; scrollbar-gutter:stable both-edges; scrollbar-width:auto; }
    .review-table-wrap::-webkit-scrollbar, .error-reason-table-wrap::-webkit-scrollbar { width:11px; height:11px; }
    .review-table-wrap::-webkit-scrollbar-track, .error-reason-table-wrap::-webkit-scrollbar-track { background:var(--soft); }
    .review-table-wrap::-webkit-scrollbar-thumb, .error-reason-table-wrap::-webkit-scrollbar-thumb { background:var(--muted); border:3px solid var(--soft); border-radius:8px; }
    .review-table-wrap > table { min-width:1420px; margin-bottom:220px; }
    .gold-edit-status { margin:-3px 0 8px; color:var(--muted); font-size:12px; text-align:right; }
    .gold-editor { display:grid; gap:7px; min-width:170px; }
    .gold-editor select, .gold-editor textarea { width:100%; min-height:34px; padding:6px 8px; font-size:12px; }
    .gold-editor textarea { min-height:48px; resize:vertical; line-height:1.4; }
    .gold-editor-note { color:var(--muted); font-size:11px; line-height:1.35; }
    .gold-editor-reset { width:auto; min-height:30px; padding:4px 7px; font-size:11px; }
    .review-title-group { display:flex; align-items:center; gap:9px; min-width:0; }
    .file-status { display:inline-flex; align-items:center; min-height:26px; padding:3px 8px; border-radius:999px; background:var(--amber-soft); color:var(--amber); font-size:12px; font-weight:700; white-space:nowrap; }
    .file-status.completed { background:var(--green-soft); color:var(--green); }
    .review-head { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; margin-bottom:10px; display:flex; gap:10px; justify-content:space-between; align-items:center; }
    .review-head h2 { margin:0; font-size:17px; letter-spacing:0; }
    .review-actions { display:flex; gap:8px; flex-wrap:wrap; justify-content:flex-end; }
    .review-actions button { width:auto; }
    table { width:100%; border-collapse:collapse; background:var(--surface); border:1px solid var(--line); }
    th, td { border-bottom:1px solid var(--line); padding:8px; vertical-align:top; font-size:13px; }
    th { position:sticky; top:0; background:#eef2f8; z-index:2; text-align:left; }
    tr.review { background:#fff8ed; }
    tr.unmatched { background:#fff3f0; }
    tr.confirmed { background:#f0fdf4; }
    .pill { display:inline-block; padding:2px 7px; border-radius:999px; font-size:12px; font-weight:700; }
    .pill.auto { color:var(--green); background:#ecfdf3; }
    .pill.review { color:var(--amber); background:#fffaeb; }
    .pill.unmatched, .pill.duplicate_dropped { color:var(--red); background:#fef3f2; }
    .pill.confirmed, .pill.confirmed_none { color:#175cd3; background:#eff4ff; }
    .rule { min-width:260px; }
    .dimension { min-width:150px; }
    .dimension-select { width:100%; min-width:120px; padding:7px 8px; font-weight:700; color:var(--ink); background:var(--surface); }
    .dimension-select.manual { border-color:var(--blue); }
    .dimension-note { margin-top:5px; color:var(--muted); font-size:12px; line-height:1.4; }
    .desc { max-width:460px; }
    .candidate { min-width:390px; }
    .candidate label { margin-top:8px; }
    .choice-card { border:1px solid var(--line); background:var(--surface); border-radius:8px; padding:9px; margin-bottom:8px; }
    .choice-head { display:flex; justify-content:space-between; gap:8px; align-items:flex-start; }
    .choice-title { font-weight:700; line-height:1.35; }
    .choice-meta { color:var(--muted); font-size:12px; margin-top:4px; line-height:1.45; }
    .choice-action { width:auto; white-space:nowrap; padding:6px 9px; font-size:12px; background:var(--surface); color:var(--blue); }
    .actions { display:flex; flex-wrap:wrap; gap:6px; margin-top:8px; }
    .manual { display:flex; gap:6px; margin-top:8px; align-items:center; }
    .manual input { width:120px; padding:6px 8px; font-size:12px; }
    .manual button { width:auto; padding:6px 9px; font-size:12px; }
    .recommend { margin-top:8px; padding:7px; background:var(--soft); border:1px solid var(--line); border-radius:6px; }
    .recommend-list { display:flex; flex-wrap:wrap; gap:6px; margin-top:5px; }
    .rec-chip { width:auto; border:1px solid var(--line); background:var(--surface); color:var(--ink); padding:5px 7px; font-size:12px; max-width:100%; text-align:left; }
    .detail-line { margin-top:7px; color:var(--muted); font-size:12px; line-height:1.45; }
    .candidate-summary { background:#fffaf0; border:1px dashed #f59e0b; border-radius:7px; padding:8px 9px; }
    .candidate-summary .summary-title { color:#92400e; font-weight:800; margin-bottom:3px; }
    .model-desc-box { background:#f8fafc; border-left:4px solid var(--blue); padding:9px 10px; line-height:1.55; color:var(--ink); }
    .model-desc-box .desc-title { font-weight:800; color:var(--blue); margin-bottom:5px; }
    .model-desc-empty { color:var(--muted); }
    .error-reason-box { margin-top:8px; padding:8px 9px; background:#fff8ed; border:1px solid #f3c27b; border-radius:7px; }
    .error-reason-box label { color:var(--amber); font-weight:800; margin-bottom:5px; }
    .error-reason-box textarea { width:100%; min-height:58px; resize:vertical; border:1px solid #f3c27b; border-radius:6px; padding:7px 8px; font:inherit; font-size:12px; line-height:1.45; color:var(--ink); background:#fffdf8; }
    .verdict { min-width:175px; }
    .verdict-value { font-weight:700; margin-bottom:7px; }
    .verdict-value.good { color:var(--green); }
    .verdict-value.bad { color:var(--red); }
    .verdict-actions { display:flex; flex-wrap:wrap; gap:6px; }
    .verdict-actions button { width:auto; padding:5px 8px; font-size:12px; }
    .verdict-actions button.active { background:var(--blue); border-color:var(--blue); color:#fff; }
    .verdict-actions button.active.secondary { background:var(--blue); color:#fff; }
    .field-stack { display:grid; gap:6px; min-width:92px; }
    .field-row { display:grid; grid-template-columns: 64px minmax(30px, auto); gap:6px; align-items:center; }
    .field-label { color:var(--muted); font-size:12px; }
    .field-value { font-weight:800; color:var(--ink); }
    .field-value.yes, .field-value.true { color:var(--green); }
    .field-value.no, .field-value.false { color:var(--red); }
    .field-value.good { color:var(--green); }
    .field-value.bad { color:var(--red); }
    .field-note { color:var(--muted); font-size:12px; line-height:1.45; margin-top:6px; }
    code { font-family: ui-monospace, SFMono-Regular, Menlo, Consolas, monospace; background:var(--soft); border:1px solid var(--line); border-radius:4px; padding:1px 4px; }
    .formula-panel { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; margin-bottom:14px; display:grid; grid-template-columns: 1fr 1fr; gap:10px 18px; }
    .formula-title { font-weight:700; margin-bottom:4px; }
    .formula-line { font-size:13px; color:var(--muted); line-height:1.5; }
    .formula-line b { color:var(--ink); }
    .dimension-panel { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; margin-bottom:14px; }
    .dimension-panel-head { display:flex; align-items:baseline; gap:10px; margin-bottom:8px; }
    .dimension-panel-title { font-weight:800; }
    .dimension-panel-note { color:var(--muted); font-size:12px; line-height:1.6; }
    .dimension-table { width:100%; border-collapse:collapse; table-layout:fixed; }
    .dimension-table th, .dimension-table td { border:1px solid var(--line); padding:7px 8px; font-size:12px; text-align:left; }
    .dimension-table th { position:static; background:var(--soft); color:var(--ink); }
    .dimension-table th.num, .dimension-table td.num { text-align:center; }
    .dimension-table td.num { text-align:center; font-weight:700; }
    .dimension-table td.rate { color:var(--blue); font-weight:800; }
    .dimension-table td.formula { white-space:normal; line-height:1.55; }
    .error-reason-panel { background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:12px 14px; margin-bottom:14px; }
    .error-reason-panel-title { display:flex; align-items:center; justify-content:space-between; gap:10px; font-weight:800; margin-bottom:8px; }
    .error-reason-panel-title button { width:auto; min-height:34px; padding:6px 10px; }
    .error-reason-table { width:100%; min-width:1900px; border-collapse:collapse; }
    .error-reason-table-wrap { width:100%; overflow:auto; }
    .error-reason-table th, .error-reason-table td { border:1px solid var(--line); padding:7px 8px; font-size:12px; text-align:left; vertical-align:top; }
    .error-reason-table th { position:static; background:#fff2e8; }
    .small { font-size:12px; color:var(--muted); }
    .hide { display:none; }
    .toast-container { position:fixed; top:18px; right:20px; z-index:60; width:min(380px, calc(100vw - 32px)); display:grid; gap:10px; pointer-events:none; }
    .toast { display:flex; align-items:flex-start; gap:10px; padding:12px 14px; border:1px solid var(--line); border-left:4px solid var(--blue); border-radius:8px; background:var(--surface); color:var(--ink); font-size:13px; line-height:1.45; font-weight:700; opacity:0; transform:translateY(-8px); animation:toast-in .18s ease forwards; }
    .toast.success { border-left-color:var(--green); }
    .toast.error { border-left-color:var(--red); }
    .toast.info { border-left-color:var(--blue); }
    .toast-mark { flex:0 0 auto; min-width:20px; font-size:16px; line-height:1.2; }
    .toast.success .toast-mark { color:var(--green); }
    .toast.error .toast-mark { color:var(--red); }
    .toast.info .toast-mark { color:var(--blue); }
    .toast.is-leaving { animation:toast-out .18s ease forwards; }
    @keyframes toast-in { to { opacity:1; transform:translateY(0); } }
    @keyframes toast-out { to { opacity:0; transform:translateY(-8px); } }
    body[data-theme="compact"] .toast { background:#111c2d; color:#e7f4ff; border-color:#2b4059; }
    .picker-overlay { position:fixed; inset:0; z-index:20; background:rgba(15,23,42,.38); display:flex; align-items:flex-start; justify-content:center; padding:78px 24px 24px; }
    .picker-overlay.hide { display:none; }
    .picker-dialog { width:min(980px, calc(100vw - 48px)); max-height:calc(100vh - 90px); background:var(--surface); border:1px solid var(--line); border-radius:10px; box-shadow:none; display:flex; flex-direction:column; overflow:hidden; }
    .picker-head { flex:0 0 auto; position:relative; z-index:2; padding:13px 14px; border-bottom:1px solid var(--line); display:grid; grid-template-columns: 1fr auto; gap:10px; align-items:center; }
    .picker-title { font-weight:800; }
    .picker-subtitle { color:var(--muted); font-size:12px; margin-top:3px; }
    .picker-close { width:auto; background:var(--surface); color:var(--blue); }
    .picker-search { flex:0 0 auto; padding:10px 14px; border-bottom:1px solid var(--line); }
    .picker-list { flex:1 1 auto; min-height:0; overflow:auto; padding:10px 14px 14px; display:flex; flex-direction:column; align-items:stretch; gap:8px; }
    .picker-item { display:block; flex:0 0 auto; width:100%; height:auto; min-height:0; overflow:visible; cursor:pointer; text-align:left; white-space:normal; background:var(--surface); color:var(--ink); border:1px solid var(--line); border-radius:8px; padding:10px 11px; }
    .picker-item:hover { border-color:var(--blue); transform:none; }
    .picker-item.active { border-color:var(--blue); background:var(--soft); transform:none; }
    .picker-item:focus-visible { outline:2px solid var(--blue); outline-offset:2px; }
    .picker-item-top { display:flex; flex-wrap:wrap; gap:8px; align-items:center; margin-bottom:5px; }
    .id-badge, .score-badge, .bool-badge { display:inline-flex; align-items:center; border-radius:999px; padding:2px 7px; font-size:12px; font-weight:700; }
    .id-badge { background:var(--soft); border:1px solid var(--line); color:var(--ink); }
    .score-badge { background:#eff4ff; color:#175cd3; }
    .bool-badge.true { background:#ecfdf3; color:var(--green); }
    .bool-badge.false { background:#fef3f2; color:var(--red); }
    .picker-rule { max-width:100%; font-weight:700; line-height:1.4; overflow-wrap:anywhere; word-break:break-word; }
    .picker-desc { max-width:100%; color:var(--muted); font-size:12px; line-height:1.45; margin-top:4px; white-space:normal; overflow-wrap:anywhere; word-break:break-word; }
    .picker-pagination { flex:0 0 auto; position:relative; z-index:2; display:flex; align-items:center; justify-content:center; gap:8px; flex-wrap:wrap; padding:10px 14px; border-top:1px solid var(--line); background:var(--surface); }
    .picker-pagination button { width:auto; min-height:34px; padding:6px 10px; }
    .picker-page-status { color:var(--muted); font-size:12px; min-width:150px; text-align:center; }
    .history-save-overlay { position:fixed; inset:0; z-index:40; display:flex; align-items:center; justify-content:center; padding:24px; background:rgba(15,23,42,.42); }
    .history-save-overlay.hide { display:none; }
    .history-save-dialog { width:min(720px, calc(100vw - 32px)); background:var(--surface); border:1px solid var(--line); border-radius:10px; overflow:hidden; }
    .history-save-head { display:flex; align-items:flex-start; justify-content:space-between; gap:16px; padding:16px 18px; background:var(--soft); border-bottom:1px solid var(--line); }
    .history-save-title { font-size:17px; font-weight:800; }
    .history-save-subtitle { margin-top:4px; color:var(--muted); font-size:12px; line-height:1.45; }
    .history-save-close { width:auto; min-width:34px; padding:6px 10px; background:transparent; color:var(--muted); border-color:transparent; font-size:20px; line-height:1; }
    .history-save-close:hover { background:var(--surface); color:var(--ink); }
    .history-save-body { padding:18px; display:grid; gap:13px; }
    .history-save-body input { width:100%; }
    .history-save-time-grid { display:grid; grid-template-columns:repeat(2, minmax(280px, 1fr)); gap:12px; }
    .history-save-time-grid input { min-width:0; }
    .history-save-duration { padding:11px 12px; border:1px solid var(--line); border-radius:6px; background:var(--soft); color:var(--muted); font-size:13px; }
    .history-save-actions { display:flex; justify-content:flex-end; gap:8px; padding:0 18px 18px; }
    .history-save-actions button { width:auto; min-width:92px; }

    /* Flat Design system: hierarchy comes from color blocks and type, never depth. */
    * { box-shadow:none !important; }
    h1, h2, .formula-title, .dimension-panel-title, .history-item-title { font-weight:800; letter-spacing:0; }
    header { background:var(--blue); color:#fff; border:0; padding:12px 24px; }
    header .header-topline { min-height:40px; margin-bottom:0; }
    h1 { font-size:20px; line-height:1.25; margin-bottom:0; }
    header label { color:rgba(255,255,255,.88); font-weight:600; }
    input, select, textarea { border:0; border-radius:6px; background:var(--soft); color:var(--ink); min-height:44px; }
    input[type=file] { background:#fff; }
    .controls input[type=file] { display:block; width:100%; height:44px; min-height:44px; line-height:24px; overflow:hidden; }
    input:focus, select:focus, textarea:focus { outline:2px solid var(--blue); outline-offset:2px; background:#fff; }
    button { min-height:44px; border:0; border-radius:6px; background:var(--blue); color:#fff; font-weight:700; transition:background .2s ease; }
    button:hover { background:var(--blue-dark); transform:none; }
    button:disabled { cursor:not-allowed; opacity:.45; transform:none !important; }
    button:disabled:hover { background:var(--blue); }
    button:focus-visible { outline:2px solid var(--ink); outline-offset:3px; }
    button.secondary { background:var(--soft); color:var(--ink); }
    button.secondary:hover { background:#e5e7eb; }
    button.mini { min-height:34px; }
    main { padding:22px 24px 48px; }
    .app-layout { margin-left:-24px; margin-right:-24px; }
    .cards { gap:12px; margin-bottom:18px; }
    .card { border:0; border-radius:8px; padding:16px; min-height:126px; }
    .card:nth-child(1), .card:nth-child(2) { background:var(--soft); }
    .card:nth-child(3) { background:var(--red-soft); }
    .card:nth-child(4) { background:var(--blue-soft); }
    .card:nth-child(5) { background:var(--green-soft); }
    .card:nth-child(6) { background:var(--amber-soft); }
    .card .k { color:var(--muted); font-weight:600; }
    .card .v { font-size:28px; line-height:1; margin-top:10px; }
    .card:nth-child(1) .v, .card:nth-child(2) .v { color:var(--ink); }
    .card:nth-child(3) .v { color:var(--red); }
    .card:nth-child(4) .v { color:var(--blue-dark); }
    .card:nth-child(5) .v { color:var(--green); }
    .card:nth-child(6) .v { color:var(--amber); }
    .card .formula { margin-top:12px; }
    .toolbar { gap:8px; margin:14px 0; }
    .theme-switcher { gap:4px; background:var(--soft); padding:4px; border-radius:8px; }
    .theme-switcher button { min-height:36px; padding:6px 10px; background:transparent; color:var(--muted); }
    .theme-switcher button:hover { background:#e5e7eb; }
    .theme-switcher button.active { background:var(--ink); color:#fff; }
    header .theme-switcher { background:rgba(255,255,255,.14); border-color:rgba(255,255,255,.28); }
    header .theme-switcher button { background:transparent; color:#fff; }
    header .theme-switcher button:hover { background:rgba(255,255,255,.18); }
    header .theme-switcher button.active { background:#fff; color:var(--blue-dark); }
    .module-menu { gap:6px; padding:10px; background:var(--soft); border-radius:8px; }
    .module-menu-head { padding:4px 4px 8px; background:transparent; color:var(--muted); }
    .module-tab { min-height:44px; padding:9px 12px; background:transparent; color:var(--muted); border:0; border-radius:6px; }
    .module-tab:hover { background:#e5e7eb; color:var(--ink); }
    .module-tab.active { background:var(--blue); color:#fff; }
    .module-menu-toggle { min-height:34px; background:transparent; color:var(--muted); }
    .module-menu-toggle:hover { background:#e5e7eb; color:var(--ink); }
    .analysis-empty { border-radius:8px; padding:34px 36px; }
    .module-heading { margin-bottom:16px; }
    .module-heading h2 { font-size:22px; }
    .module-heading p { max-width:720px; font-size:13px; }
    .history-save { min-width:min(760px, 100%); }
    .history-save button { min-height:34px; padding:7px 14px; font-size:13px; }
    .history-toolbar { margin-bottom:12px; }
    .history-item { border:0; border-radius:8px; padding:16px; background:var(--soft); }
    .history-item:nth-child(3n+1) { background:var(--blue-soft); }
    .history-item:nth-child(3n+2) { background:var(--green-soft); }
    .history-item:nth-child(3n+3) { background:var(--amber-soft); }
    .history-item-title { font-size:16px; }
    .history-metrics b { color:var(--ink); font-size:14px; }
    .history-actions { margin-top:12px; }
    .history-empty { border:0; border-radius:8px; background:var(--soft); padding:34px; }
    .workspace { gap:16px; }
    .sidebar { border:0; border-radius:8px; background:var(--soft); top:18px; }
    .sidebar-head { padding:16px; border:0; background:var(--ink); color:#fff; }
    .sidebar-title { font-size:16px; }
    .sidebar-filter label { color:rgba(255,255,255,.78); }
    .sidebar-filter select, .sidebar-filter input { background:#fff; }
    .material-item { border-bottom:2px solid var(--bg); padding:13px 16px; }
    .material-item:hover { background:#e5e7eb; transform:none; }
    .material-item.active { background:var(--blue-soft); border-left:6px solid var(--blue); padding-left:10px; }
    .material-type { color:var(--ink); }
    .material-meta strong { color:var(--red); }
    .review-head { border:0; border-radius:8px; padding:16px; margin-bottom:12px; background:var(--blue-soft); }
    .review-head h2 { font-size:20px; }
    table { border:0; background:#fff; }
    th, td { border-bottom:2px solid #fff; padding:10px; }
    th { position:sticky; top:0; background:var(--ink); color:#fff; }
    tr.review { background:var(--amber-soft); }
    tr.unmatched { background:var(--red-soft); }
    tr.confirmed { background:var(--green-soft); }
    .pill { border-radius:6px; padding:4px 8px; }
    .pill.auto { background:var(--green-soft); color:var(--green); }
    .pill.review { background:var(--amber-soft); color:var(--amber); }
    .pill.unmatched, .pill.duplicate_dropped { background:var(--red-soft); color:var(--red); }
    .pill.confirmed, .pill.confirmed_none { background:var(--blue-soft); color:var(--blue-dark); }
    .dimension-select { min-height:40px; background:var(--soft); }
    .dimension-select.manual { border:2px solid var(--blue); }
    .choice-card { border:0; border-radius:6px; background:#fff; padding:12px; }
    .choice-action { min-height:36px; background:var(--soft); color:var(--ink); }
    .choice-action:hover, .picker-close:hover { background:#e5e7eb; color:var(--ink); }
    .recommend { background:var(--blue-soft); border:0; border-radius:6px; padding:10px; }
    .rec-chip { min-height:34px; border:0; background:#fff; color:var(--ink); }
    .rec-chip:hover { background:var(--blue); color:#fff; }
    .candidate-summary { background:var(--amber-soft); border:0; border-left:5px solid var(--amber); border-radius:0 6px 6px 0; padding:10px; }
    .model-desc-box { background:var(--soft); border-left:5px solid var(--blue); padding:11px 12px; }
    .error-reason-box { background:var(--amber-soft); border:0; border-left:5px solid var(--amber); border-radius:0 6px 6px 0; padding:10px; }
    .error-reason-box textarea { min-height:68px; border:0; background:#fff; }
    .field-value.yes, .field-value.true { color:var(--green); }
    .field-value.no, .field-value.false { color:var(--red); }
    .formula-panel { border:0; border-radius:8px; background:var(--soft); padding:16px; }
    .formula-line { font-size:13px; }
    code { border:0; border-radius:4px; background:#fff; }
    .dimension-panel, .error-reason-panel { border:0; border-radius:8px; background:var(--soft); padding:16px; }
    .dimension-table, .error-reason-table { background:#fff; }
    .dimension-table th, .error-reason-table th { background:var(--ink); color:#fff; }
    .dimension-table th, .dimension-table td, .error-reason-table th, .error-reason-table td { border-bottom:2px solid var(--soft); border-left:0; border-right:0; }
    .dimension-table td.rate { color:var(--blue-dark); }
    .error-reason-table th { background:var(--amber); }
    .picker-overlay { background:rgba(17,24,39,.55); }
    .picker-dialog { border:0; border-radius:8px; }
    .picker-head { background:var(--ink); color:#fff; border:0; }
    .picker-subtitle { color:rgba(255,255,255,.72); }
    .picker-close { min-height:36px; background:#fff; color:var(--ink); }
    .picker-search { background:var(--soft); border:0; }
    .picker-item { display:block; height:auto; min-height:0; border:2px solid var(--line); border-radius:6px; }
    .picker-item:hover { border-color:var(--blue); background:var(--blue-soft); transform:none; }
    .picker-item.active { border-color:var(--blue); background:var(--blue-soft); transform:none; }
    .id-badge, .score-badge, .bool-badge { border-radius:4px; }
    .id-badge { background:var(--soft); border:0; }
    .score-badge { background:var(--blue-soft); color:var(--blue-dark); }
    .bool-badge.true { background:var(--green-soft); color:var(--green); }
    .bool-badge.false { background:var(--red-soft); color:var(--red); }

    /* Theme 2: 黑曜荧光。适合展示和夜间使用，只有少量青色作为状态信号。 */
    body[data-theme="compact"] {
      --bg:#0b1220; --surface:#111c2d; --soft:#17263a; --ink:#e7f4ff; --muted:#9db2c9;
      --line:#2b4059; --blue:#22d3ee; --blue-dark:#06b6d4; --blue-soft:#123949;
      --red:#fb7185; --red-soft:#3a1d2a; --amber:#fbbf24; --amber-soft:#3a2d12;
      --green:#34d399; --green-soft:#123b36;
    }
    body[data-theme="compact"] { background:var(--bg); color:var(--ink); }
    body[data-theme="compact"] header { background:#08101d; border-bottom:1px solid var(--blue); }
    body[data-theme="compact"] header label { color:#a9c8dc; }
    body[data-theme="compact"] header .theme-switcher { background:var(--surface); border-color:var(--line); }
    body[data-theme="compact"] header .theme-switcher button { color:var(--muted); }
    body[data-theme="compact"] header .theme-switcher button.active { background:var(--blue); color:#06131f; }
    body[data-theme="compact"] input,
    body[data-theme="compact"] select,
    body[data-theme="compact"] textarea { background:var(--soft); color:var(--ink); }
    body[data-theme="compact"] input[type=file] { background:var(--surface); }
    body[data-theme="compact"] input:focus,
    body[data-theme="compact"] select:focus,
    body[data-theme="compact"] textarea:focus { background:var(--surface); }
    body[data-theme="compact"] button { color:#06131f; }
    body[data-theme="compact"] button.secondary,
    body[data-theme="compact"] .choice-action,
    body[data-theme="compact"] .picker-close { background:var(--soft); color:var(--ink); }
    body[data-theme="compact"] button.secondary:hover,
    body[data-theme="compact"] .choice-action:hover,
    body[data-theme="compact"] .picker-close:hover { background:#233b55; }
    body[data-theme="compact"] .theme-switcher { background:var(--surface); border:1px solid var(--line); }
    body[data-theme="compact"] .module-menu { background:var(--surface); border:1px solid var(--line); }
    body[data-theme="compact"] .module-menu-head { background:transparent; color:var(--muted); }
    body[data-theme="compact"] .theme-switcher button { color:var(--muted); }
    body[data-theme="compact"] .theme-switcher button.active { background:var(--blue); color:#06131f; }
    body[data-theme="compact"] .module-menu-head { color:var(--muted); }
    body[data-theme="compact"] .module-tab { color:var(--muted); }
    body[data-theme="compact"] .module-tab:hover { background:var(--soft); color:var(--ink); }
    body[data-theme="compact"] .module-tab.active { background:var(--blue); color:#06131f; }
    body[data-theme="compact"] .module-menu-toggle { color:var(--muted); }
    body[data-theme="compact"] .card { border:1px solid var(--line); border-radius:4px; }
    body[data-theme="compact"] .card:nth-child(1),
    body[data-theme="compact"] .card:nth-child(2) { background:var(--surface); }
    body[data-theme="compact"] .card:nth-child(3) { background:var(--red-soft); }
    body[data-theme="compact"] .card:nth-child(4) { background:#123949; }
    body[data-theme="compact"] .card:nth-child(5) { background:#123b36; }
    body[data-theme="compact"] .card:nth-child(6) { background:#3a2d12; }
    body[data-theme="compact"] .card:nth-child(1) .v,
    body[data-theme="compact"] .card:nth-child(2) .v { color:var(--ink); }
    body[data-theme="compact"] .card:nth-child(3) .v { color:var(--red); }
    body[data-theme="compact"] .card:nth-child(4) .v { color:var(--blue); }
    body[data-theme="compact"] .card:nth-child(5) .v { color:var(--green); }
    body[data-theme="compact"] .card:nth-child(6) .v { color:var(--amber); }
    body[data-theme="compact"] .info-popover-content { background:#07111f; color:#e7f4ff; border:1px solid #2b4059; }
    body[data-theme="compact"] .info-popover-content::before { background:#07111f; border-left:1px solid #2b4059; border-top:1px solid #2b4059; }
    body[data-theme="compact"] .info-popover-content .formula-title { color:var(--blue); }
    body[data-theme="compact"] .info-popover-content .formula-line { color:#d8e6f5; }
    body[data-theme="compact"] .info-popover-content .formula-line b { color:#ffffff; }
    body[data-theme="compact"] .info-popover-content code { color:#e7f4ff; background:#17263a; border:1px solid #2b4059; }
    body[data-theme="compact"] .sidebar { background:var(--surface); border:1px solid var(--line); border-radius:4px; }
    body[data-theme="compact"] .sidebar-head { background:#08101d; }
    body[data-theme="compact"] .material-item { background:var(--surface); border-bottom:1px solid var(--line); }
    body[data-theme="compact"] .material-item:hover { background:var(--soft); }
    body[data-theme="compact"] .material-item.active { background:#123949; border-left-color:var(--blue); }
    body[data-theme="compact"] .review-head { background:#123949; border-left:5px solid var(--blue); }
    body[data-theme="compact"] table { background:var(--surface); }
    body[data-theme="compact"] th { background:#08101d; }
    body[data-theme="compact"] tr.review { background:#3a2d12; }
    body[data-theme="compact"] tr.unmatched { background:#3a1d2a; }
    body[data-theme="compact"] tr.confirmed { background:#123b36; }
    body[data-theme="compact"] .formula-panel,
    body[data-theme="compact"] .dimension-panel,
    body[data-theme="compact"] .error-reason-panel,
    body[data-theme="compact"] .history-item { background:var(--surface); border:1px solid var(--line); border-radius:4px; }
    body[data-theme="compact"] .dimension-table,
    body[data-theme="compact"] .error-reason-table { background:var(--surface); }
    body[data-theme="compact"] .dimension-table th,
    body[data-theme="compact"] .error-reason-table th { background:#08101d; }
    body[data-theme="compact"] .choice-card,
    body[data-theme="compact"] .rec-chip,
    body[data-theme="compact"] .picker-item { background:var(--surface); border-color:var(--line); color:var(--ink); }
    body[data-theme="compact"] .recommend { background:#123949; }
    body[data-theme="compact"] .candidate-summary { background:var(--amber-soft); }
    body[data-theme="compact"] .model-desc-box { background:var(--soft); }
    body[data-theme="compact"] code { background:var(--soft); color:var(--ink); }
    body[data-theme="compact"] .picker-overlay { background:rgba(0,0,0,.72); }
    body[data-theme="compact"] .picker-dialog { background:var(--surface); border:1px solid var(--line); }
    body[data-theme="compact"] .picker-head { background:#08101d; }
    body[data-theme="compact"] .picker-search { background:var(--soft); }
    body[data-theme="compact"] .material-item-select { background:transparent; color:var(--ink); }
    body[data-theme="compact"] .material-item-select:hover { background:transparent; color:var(--ink); }
    body[data-theme="compact"] .material-submit { color:#06131f; }
    body[data-theme="compact"] .toast { background:#111c2d; color:#e7f4ff; border-color:#2b4059; }

    .material-item { padding:0; }
    .material-item:hover { background:var(--surface); transform:none; }
    .material-item.active { background:var(--blue-soft); border-left:6px solid var(--blue); padding-left:0; }
    .material-item.active .material-item-select { background:var(--blue-soft); }
    .material-item.active .material-item-select:hover { background:var(--blue-soft); }
    .material-item.completed { background:var(--green-soft); }
    .material-item.completed .material-item-select { background:var(--green-soft); }
    .material-item.completed.active { border-left-color:var(--green); }
    .material-item.has-missed:not(.active) { border-left:4px solid var(--red); }
    .material-item.pending:not(.active) { border-left:4px solid var(--amber); }
    .material-item.incomplete:not(.active) { border-left:4px solid var(--muted); }
    .material-item.ready:not(.active) { border-left:4px solid var(--blue); }
    .material-item.has-missed .material-item-select { background:var(--red-soft); }
    .material-item.pending .material-item-select { background:var(--amber-soft); }
    .material-item.completed .material-item-select { background:var(--green-soft); }
    .material-item.incomplete .material-item-select { background:var(--soft); }
    .material-item.active .material-item-select { background:var(--blue-soft); }
    .row-focus { outline:3px solid var(--blue); outline-offset:-3px; animation:row-focus-pulse 1.5s ease-out; }
    @keyframes row-focus-pulse { 0% { outline-color:var(--blue); } 100% { outline-color:transparent; } }
    .review-actions button, .actions button, .verdict-actions button, .history-toolbar button, .toolbar button, .material-submit, .material-pagination button, .history-pagination button { min-height:34px; padding:7px 10px; font-size:13px; }
    .review-actions, .actions, .verdict-actions, .history-toolbar { gap:6px; }

    /* 黑曜荧光主题：正文和辅助文字使用更高对比度，避免深色背景上发灰。 */
    body[data-theme="compact"] .small,
    body[data-theme="compact"] .choice-meta,
    body[data-theme="compact"] .dimension-note,
    body[data-theme="compact"] .field-label,
    body[data-theme="compact"] .field-note,
    body[data-theme="compact"] .history-files,
    body[data-theme="compact"] .history-review-time,
    body[data-theme="compact"] .history-page-status,
    body[data-theme="compact"] .material-doc,
    body[data-theme="compact"] .material-meta,
    body[data-theme="compact"] .card .formula,
    body[data-theme="compact"] .picker-desc { color:#c4d7e8; }
    body[data-theme="compact"] .material-summary { color:#72e8ff; background:#123949; }
    body[data-theme="compact"] .material-item.has-missed .material-item-select { background:#3a1d2a; }
    body[data-theme="compact"] .material-item.pending .material-item-select { background:#3a2d12; }
    body[data-theme="compact"] .material-item.incomplete .material-item-select { background:#17263a; }
    body[data-theme="compact"] .material-item.completed .material-item-select { background:#123b36; }
    body[data-theme="compact"] .material-item.active .material-item-select { background:#123949; }
    body[data-theme="compact"] .status-legend-item.incomplete, body[data-theme="compact"] .material-status-chip.incomplete { color:#c4d7e8; background:#17263a; }
    body[data-theme="compact"] .status-legend-item.pending, body[data-theme="compact"] .material-status-chip.pending { color:#ffd166; background:#3a2d12; }
    body[data-theme="compact"] .status-legend-item.missed, body[data-theme="compact"] .material-status-chip.missed { color:#ff9aaa; background:#3a1d2a; }
    body[data-theme="compact"] .status-legend-item.completed, body[data-theme="compact"] .material-status-chip.completed { color:#6ee7b7; background:#123b36; }
    body[data-theme="compact"] .review-table-wrap::-webkit-scrollbar-track, body[data-theme="compact"] .error-reason-table-wrap::-webkit-scrollbar-track { background:#17263a; }
    body[data-theme="compact"] .review-table-wrap::-webkit-scrollbar-thumb, body[data-theme="compact"] .error-reason-table-wrap::-webkit-scrollbar-thumb { background:#6f8da8; border-color:#17263a; }
    .material-submit { background:var(--blue); color:#fff; }
    .material-submit:hover { background:var(--blue-dark); color:#fff; transform:none; }
    .workspace.materials-collapsed .material-pagination { display:none; }

    @media (max-width: 1280px) {
      header { padding:10px 20px; }
      main { padding:18px 20px 40px; }
      .app-layout { margin-left:-20px; margin-right:-20px; }
      .app-layout { grid-template-columns:160px minmax(0,1fr); gap:14px; }
      .app-layout.module-menu-collapsed { grid-template-columns:50px minmax(0,1fr); }
      .workspace { grid-template-columns:minmax(196px, 20vw) minmax(0,1fr); }
      th, td { padding:8px; }
    }
    @media (max-width: 1100px) {
      .controls { grid-template-columns: 1fr; }
      .cards { grid-template-columns: repeat(2, 1fr); }
      .app-layout { grid-template-columns:1fr; margin-left:-20px; margin-right:-20px; }
      .module-menu { position:relative; top:0; min-height:0; padding:0; flex-direction:row; align-items:center; overflow:auto; }
      .app-layout.module-menu-collapsed { grid-template-columns:1fr; }
      .module-menu-head { flex:0 0 auto; padding:4px 6px; }
      .module-menu-title { display:none; }
      .module-menu-toggle { flex:0 0 40px; }
      .module-tab { width:auto; flex:0 0 auto; white-space:nowrap; }
      .volume-block { display:flex; flex-direction:row; gap:6px; }
      .volume-switch { width:auto; white-space:nowrap; }
      .module-submenu { display:flex; flex-direction:row; padding:0; }
      .case-fetch-grid, .case-info-grid { grid-template-columns:1fr; }
      .workspace { grid-template-columns: 1fr; }
      .sidebar { position:relative; top:0; height:auto; min-height:0; max-height:300px; }
      .annotation-controls { grid-template-columns:1fr; }
      .annotation-controls button { width:100%; }
      .annotation-layout { grid-template-columns:1fr; min-height:0; }
      .annotation-table-pdf-layout { grid-template-columns:1fr; min-height:0; }
      .annotation-excel-table-wrap { height:520px; }
      .annotation-rule-list, .annotation-meta { max-height:300px; }
      .annotation-pdf { height:520px; }
      th { top:0; }
      .module-heading { flex-direction:column; }
      .history-save { width:100%; min-width:0; }
      .history-save > button { width:auto; }
      .history-save-time-grid { grid-template-columns:1fr; }
    }
  </style>
</head>
<body>
  <header>
    <div class="header-topline">
      <h1>JSON 表格匹配复核工具</h1>
      <div class="theme-switcher" aria-label="切换界面主题">
        <button id="theme-review" title="冷白蓝线" onclick="setTheme('review')">冷白蓝线</button>
        <button id="theme-compact" title="黑曜荧光" onclick="setTheme('compact')">黑曜荧光</button>
      </div>
    </div>
  </header>
  <main>
    <div class="app-layout" id="appLayout">
      <nav class="module-menu" id="moduleMenu" aria-label="主菜单">
        <div class="module-menu-head">
          <button id="toggleModuleMenu" class="module-menu-toggle" title="收起菜单" aria-label="收起菜单" aria-expanded="true" onclick="toggleModuleMenu()">‹</button>
        </div>
        <div class="volume-block">
          <button id="volume-procedure" class="volume-switch active" onclick="setVolume('procedure')"><span class="module-icon">▣</span><span class="volume-name">程序卷</span><span class="volume-arrow">▾</span></button>
          <div id="procedureSubmenu" class="module-submenu">
            <button id="module-annotation" class="module-tab active" onclick="setModule('annotation')"><span class="module-icon">✎</span><span class="module-label">规则标注</span></button>
            <button id="module-data" class="module-tab" onclick="setModule('data')"><span class="module-icon">⇣</span><span class="module-label">数据获取</span></button>
            <button id="module-workbench" class="module-tab" onclick="setModule('workbench')"><span class="module-icon">▦</span><span class="module-label">复核工作台</span></button>
            <button id="module-stats" class="module-tab" onclick="setModule('stats')"><span class="module-icon">▤</span><span class="module-label">专项统计</span></button>
            <button id="module-history" class="module-tab" onclick="setModule('history')"><span class="module-icon">↻</span><span class="module-label">比对历史</span></button>
          </div>
        </div>
        <div class="volume-block">
          <button id="volume-identity" class="volume-switch" onclick="setVolume('identity')"><span class="module-icon">◇</span><span class="volume-name">主体身份卷</span><span class="volume-arrow">▾</span></button>
        </div>
      </nav>
      <div class="module-content">
    <section class="case-panel" id="caseInfoPanel">
      <div class="case-panel-head">
        <div>
          <div class="case-panel-title" id="casePanelTitle">程序卷 / 案件信息</div>
          <div class="case-panel-subtitle">先根据 caseId 获取案件基础信息，再按案件性质和案件类型推荐模板。</div>
        </div>
        <span class="status" id="caseInfoStatus">未获取案件信息</span>
      </div>
      <div class="case-fetch-grid">
        <div><label>caseId</label><input id="caseInfoCaseId" placeholder="输入 caseId"></div>
        <div><label>案件详情接口</label><input id="caseInfoBaseUrl" placeholder="案件详情接口地址"></div>
        <div><button onclick="fetchCaseInfo()">获取案件信息</button></div>
      </div>
      <div id="caseInfoContent" class="case-info-grid hide"></div>
      <div id="caseTemplateResult" class="case-template hide"></div>
    </section>
    <section id="dataModule" class="module-panel hide">
      <div class="module-heading">
        <div>
          <h2>数据获取</h2>
          <p>统一准备 JSON/TXT 和金标准 Excel；可以手工上传、接口拉取，也可以批量测试。</p>
        </div>
      </div>
      <section class="source-panel">
        <div class="source-tabs" aria-label="模型结果来源">
          <button id="source-upload" class="active" onclick="setModelSource('upload')">上传 JSON/TXT</button>
          <button id="source-api" onclick="setModelSource('api')">接口拉取</button>
          <button id="source-batch" onclick="setModelSource('batch')">批量测试</button>
        </div>
        <div id="uploadSourcePanel">
          <label>上传 JSON / TXT</label>
          <input id="jsonFile" type="file" accept=".json,.txt,application/json,text/plain" onchange="previewUploadedJson()">
          <div class="dataset-actions" style="margin-top:10px"><button onclick="useCurrentDataset()">使用这组数据复核</button></div>
        </div>
        <div id="apiSourcePanel" class="api-source-grid hide">
          <div><label>接口地址</label><input id="apiBaseUrl" placeholder="模型结果接口地址"></div>
          <div><label>caseId</label><input id="apiCaseId" placeholder="输入 caseId" oninput="updateFileReadyState()"></div>
          <div><button class="secondary" onclick="fetchModelResult()">拉取模型结果</button></div>
          <div class="dataset-actions"><button onclick="useCurrentDataset()">使用接口 JSON + Excel 复核</button></div>
        </div>
        <div id="batchSourcePanel" class="batch-grid hide">
          <div><label>批量模式</label><select id="batchExcelMode" onchange="updateBatchModeUI()"><option value="shared">共用一个 Excel</option><option value="mapped">每个 case 对应不同 Excel</option></select></div>
          <div><label>批量接口地址</label><input id="batchBaseUrl" placeholder="模型结果接口地址"></div>
          <div><label>caseId 列表</label><textarea id="batchCaseIds" placeholder="一行一个 caseId"></textarea></div>
          <div><label>caseId 文件</label><input id="batchCaseFile" type="file" accept=".txt,.csv,text/plain,text/csv"></div>
          <div id="batchMappingWrap" class="hide"><label>case-Excel 映射 CSV</label><input id="batchMappingFile" type="file" accept=".csv,.txt,text/csv,text/plain"></div>
          <div id="batchExcelFilesWrap" class="hide"><label>批量 Excel 文件</label><input id="batchExcelFiles" type="file" multiple accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"></div>
          <div><button onclick="startBatchTest()">开始批量测试</button></div>
        </div>
        <div id="singleExcelWrap">
          <label>金标准 Excel</label>
          <input id="excelFile" type="file" accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" onchange="updateFileReadyState()">
        </div>
      </section>
      <section id="jsonPreview" class="json-preview hide">
        <div class="json-preview-head">
          <div class="batch-result-title">JSON 结果预览</div>
          <div class="json-preview-actions">
            <button class="secondary mini" onclick="copyPreviewJson()">复制 JSON</button>
            <button class="secondary mini" onclick="downloadPreviewJson()">下载 JSON</button>
            <button class="secondary mini" onclick="useCurrentDataset()">用此 JSON 去复核</button>
          </div>
        </div>
        <div class="json-preview-meta" id="jsonPreviewMeta">未拉取结果</div>
        <pre id="jsonPreviewBody"></pre>
      </section>
      <section id="batchResult" class="batch-result hide">
        <div class="batch-result-head">
          <div>
            <div class="batch-result-title">批量测试结果</div>
            <div class="status" id="batchStatus">等待开始</div>
          </div>
          <div class="batch-result-actions">
            <button class="secondary" id="batchExportButton" disabled onclick="exportBatchResult()">导出批量结果</button>
          </div>
        </div>
        <div class="batch-summary" id="batchSummary"></div>
        <div class="batch-table-wrap">
          <table class="batch-table">
            <thead><tr><th>caseId</th><th>金标准Excel</th><th>状态</th><th>JSON规则</th><th>Excel检测项</th><th>漏检</th><th>触发率</th><th>召回率</th><th>准确率</th><th>操作</th></tr></thead>
            <tbody id="batchRows"></tbody>
          </table>
        </div>
      </section>
    </section>
    <section id="workbenchModule" class="module-panel hide">
      <div class="dataset-current">
        <div>
          <div class="dataset-current-title">当前复核数据</div>
          <div class="dataset-current-meta" id="currentDatasetInfo">请先到“数据获取”准备 JSON/TXT 和 Excel。</div>
        </div>
        <button class="secondary" onclick="setModule('data')">返回数据获取</button>
      </div>
      <div class="controls">
        <div><label>匹配阈值</label><input id="matchThreshold" type="number" min="0" max="1" step="0.01" value="0.67"></div>
        <div><label>复核阈值</label><input id="reviewThreshold" type="number" min="0" max="1" step="0.01" value="0.82"></div>
        <div><button id="analyzeButton" disabled onclick="analyze()">分析</button></div>
      </div>
      <section class="cards hide" id="cards"></section>
      <div class="toolbar">
        <div class="toolbar-actions">
          <div class="info-popover">
          <button class="info-icon" type="button" aria-label="统计口径说明">Σ</button>
          <div class="info-popover-content" id="statsInfo" role="tooltip">
            <div class="formula-title">统计口径</div>
            <div class="formula-line">召回率 = 错误场景中模型触发的检测项 / 错误场景总数。</div>
            <div class="formula-line">触发率 = 模型触发的金标准检测项 / 金标准要求触发的检测项。</div>
            <div class="formula-line">准确率 = 模型判断正确的已触发检测项 / 模型已触发的检测项。</div>
            <div class="formula-line">漏检项 = 金标准要求触发但模型未触发的检测项。</div>
            <div class="formula-line">错误场景指金标准要求触发、且金标准判定不正确的检测项；材料类型检测按外层 <code>isExit</code> 判断。</div>
          </div>
          </div>
          <div class="info-popover">
            <button class="info-icon" type="button" aria-label="匹配分数说明">#</button>
            <div class="info-popover-content" id="scoreInfo" role="tooltip">
              <div class="formula-title">匹配分数</div>
              <div class="formula-line">匹配分数只表示 Excel 行和 JSON 规则的文本相似度，不代表模型回答对错。</div>
            </div>
          </div>
          <div class="status-legend toolbar-status-legend" aria-label="材料状态说明">
            <span class="status-legend-item pending" tabindex="0" data-tip="还有检测项没有完成匹配确认，通常包括待复核、未匹配或重复候选。">待复核</span>
            <span class="status-legend-item missed" tabindex="0" data-tip="金标准要求触发，但模型没有触发的检测项。">漏检</span>
            <span class="status-legend-item incomplete" tabindex="0" data-tip="材料尚未完成全部确认，或还没有点击“提交文件”。">未完成</span>
            <span class="status-legend-item completed" tabindex="0" data-tip="检测项已完成确认，并且已经点击“提交文件”。">已完成</span>
          </div>
          <button id="saveHistoryButton" class="secondary" disabled onclick="openHistorySaveDialog()">保存当前比对</button>
          <button id="exportButton" class="secondary" disabled onclick="exportWorkbook()">导出复核结果</button>
        </div>
        <div class="analysis-progress hide" id="analysisProgress" aria-live="polite">
          <div class="analysis-progress-track" aria-hidden="true"><div class="analysis-progress-bar" id="analysisProgressBar"></div></div>
          <span class="analysis-progress-text" id="analysisProgressText">0%</span>
        </div>
        <span class="status" id="status">等待分析</span>
      </div>
      <section id="analysisEmpty" class="analysis-empty">
        <div class="analysis-empty-kicker">准备开始</div>
        <h2>请先准备复核数据</h2>
        <p>到“数据获取”上传 JSON/TXT 与 Excel，或拉取接口 JSON 后，再回到这里点击“分析”。</p>
      </section>
      <div class="workspace hide" id="reviewWorkspace">
        <aside class="sidebar">
          <div class="sidebar-head">
            <div class="sidebar-title-row">
              <div class="sidebar-title">材料列表</div>
              <button id="toggleMaterials" class="sidebar-toggle" title="收起材料列表" aria-label="收起材料列表" aria-expanded="true" onclick="toggleMaterials()">‹</button>
            </div>
            <div class="sidebar-tools">
              <div class="sidebar-filter">
                <label for="filter">状态筛选</label>
                <select id="filter" onchange="resetMaterialListPage(); renderRows()">
                  <option value="review">只看待复核</option>
                  <option value="all">全部</option>
                  <option value="missed">漏检项</option>
                  <option value="unmatched">未匹配</option>
                  <option value="auto">自动通过</option>
                  <option value="confirmed">已确认</option>
                  <option value="completed">已完成文件</option>
                  <option value="incomplete">未完成文件</option>
                </select>
              </div>
              <div class="sidebar-filter">
                <label for="search">搜索材料或规则</label>
                <input id="search" placeholder="输入关键词" oninput="resetMaterialListPage(); renderRows()">
              </div>
            </div>
            <div class="material-summary" id="materialSummary">筛选文件：0 个 · 已完成文件：0 / 0</div>
          </div>
          <div class="material-list" id="materialList"></div>
          <div class="material-pagination hide" id="materialPagination"></div>
        </aside>
        <section class="review-panel">
          <div class="review-head">
            <div class="review-title-group">
              <h2 id="currentMaterial">全部材料</h2>
              <span id="currentMaterialStatus" class="file-status">未完成</span>
            </div>
            <div class="review-actions">
              <button id="saveGoldButton" class="secondary" disabled onclick="saveGoldStandards()">保存金标准修改</button>
              <button onclick="submitCurrentMaterial()">提交当前文件</button>
              <button class="secondary" onclick="jumpNextPendingMaterial()">下一组待复核</button>
            </div>
          </div>
          <div class="gold-edit-status" id="goldEditStatus">金标准未修改</div>
          <div class="review-table-wrap">
            <table>
              <thead>
                <tr>
                  <th>行号</th>
                  <th>状态</th>
                  <th>Excel 内容</th>
                  <th>金标准</th>
                  <th>统计维度</th>
                  <th>JSON 匹配候选</th>
                  <th>模型判断正确</th>
                  <th>模型描述</th>
                </tr>
              </thead>
              <tbody id="rows"></tbody>
            </table>
          </div>
        </section>
      </div>
    </section>
    <section id="annotationModule" class="module-panel">
      <div class="annotation-heading">
      <div class="module-heading">
        <div>
          <div class="annotation-kicker">RULE STANDARD</div>
          <h2>规则金标准标注</h2>
          <p>对照 PDF 查看规则，只填写 Excel 的 M、N 列；PDF 不会写入标注结果。</p>
        </div>
      </div>
      </div>
      <div class="annotation-controls">
        <div><label>规则 Excel</label><input id="annotationExcelFile" type="file" accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"></div>
        <div><label>对照 PDF（仅查看）</label><input id="annotationPdfFile" type="file" accept=".pdf,application/pdf" onchange="loadAnnotationPdf(this)"></div>
        <div><label>工作表</label><select id="annotationSheet" onchange="switchAnnotationSheet()" disabled></select></div>
        <div><label>M列：是否触发</label><select id="annotationMColumn" onchange="applyAnnotationColumnMapping()" disabled></select></div>
        <div><label>N列：金标准结果</label><select id="annotationNColumn" onchange="applyAnnotationColumnMapping()" disabled></select></div>
        <button onclick="loadAnnotationWorkbook()">打开标注</button>
      </div>
      <div class="annotation-toolbar">
        <span class="status" id="annotationStatus">请上传规则 Excel；PDF 可以单独上传。</span>
        <button class="secondary" onclick="toggleNewRuleForm()">新增规则</button>
        <button class="secondary" id="annotationViewButton" onclick="toggleAnnotationView()">分组列表</button>
        <button class="secondary" id="annotationExportButton" onclick="exportAnnotationWorkbook()" disabled>保存并导出 Excel</button>
      </div>
      <div class="annotation-file-nav">
        <button class="secondary" onclick="moveAnnotationMaterial(-1)">上一个文件</button>
        <select id="annotationMaterialType" onchange="switchAnnotationMaterialType()" disabled></select>
        <button class="secondary" onclick="moveAnnotationMaterial(1)">下一个文件</button>
        <span class="annotation-file-status pending" id="annotationFileStatus">未加载文件</span>
        <span class="status" id="annotationFileStats">文件 0 / 0，规则 0 / 0</span>
      </div>
      <div class="annotation-new-rule hide" id="annotationNewRuleForm">
        <div class="annotation-pane-head" style="padding:0 0 10px;border:0"><span>新增规则</span><span class="small">可指定插入行，留空则追加到末尾</span></div>
        <div class="annotation-new-grid" id="annotationNewFields"></div>
        <div style="margin-top:8px;max-width:260px"><label for="annotationNewPosition">插入到 Excel 第几行（留空追加到末尾）</label><input id="annotationNewPosition" type="number" min="2" placeholder="例如：20"></div>
        <div class="annotation-new-actions"><button class="secondary" onclick="toggleNewRuleForm()">取消</button><button onclick="addAnnotationRule()">加入规则列表</button></div>
      </div>
      <div class="annotation-table-pdf-layout" id="annotationTableView">
        <section class="annotation-pane">
          <div class="annotation-pane-head"><span>规则表</span><span class="status" id="annotationTableCount">0 条</span></div>
          <div class="annotation-pane-body"><input id="annotationTableSearch" placeholder="搜索规则、材料或错误场景" oninput="renderAnnotationTable()"></div>
          <div class="annotation-excel-table-wrap"><table class="annotation-excel-table"><thead id="annotationTableHead"></thead><tbody id="annotationTableBody"><tr><td class="annotation-empty">打开 Excel 后显示规则</td></tr></tbody></table></div>
        </section>
        <section class="annotation-pane annotation-pdf">
          <div class="annotation-pane-head"><span>PDF 对照查看</span><span class="status">仅查看</span></div>
          <div id="annotationPdfEmptyTable" class="annotation-empty">上传 PDF 后，在这里对照卷宗内容进行标注</div>
          <iframe id="annotationPdfFrameTable" class="hide" title="PDF 对照查看"></iframe>
        </section>
      </div>
      <div class="annotation-layout hide" id="annotationLayout">
        <section class="annotation-pane">
          <div class="annotation-pane-head"><span>规则列表</span><span class="status" id="annotationRowCount">0 条</span></div>
          <div class="annotation-pane-body">
            <div class="annotation-list-stats" id="annotationListStats"><div>文件：<b>0 / 0</b></div><div>规则：<b>0 / 0</b></div></div>
            <div class="annotation-filter" id="annotationStatusFilter">
              <button class="active" data-filter="all" onclick="setAnnotationStatusFilter('all')">全部</button>
              <button data-filter="pending" onclick="setAnnotationStatusFilter('pending')">未开始</button>
              <button data-filter="progress" onclick="setAnnotationStatusFilter('progress')">标注中</button>
              <button data-filter="completed" onclick="setAnnotationStatusFilter('completed')">已完成</button>
              <button data-filter="issue" onclick="setAnnotationStatusFilter('issue')">有问题</button>
            </div>
            <input id="annotationSearch" placeholder="搜索规则、材料或错误场景" oninput="renderAnnotationRuleList()">
          </div>
          <div class="annotation-rule-list" id="annotationRuleList"><div class="annotation-empty" style="min-height:180px">打开 Excel 后显示规则</div></div>
        </section>
        <section class="annotation-pane annotation-pdf">
          <div class="annotation-pane-head"><span>PDF 对照查看</span><span class="status">仅查看</span></div>
          <div id="annotationPdfEmpty" class="annotation-empty">上传 PDF 后，在这里对照卷宗内容进行标注</div>
          <iframe id="annotationPdfFrame" class="hide" title="PDF 对照查看"></iframe>
        </section>
        <section class="annotation-pane">
          <div class="annotation-pane-head"><span>当前规则</span><span><button class="secondary" id="annotationEditRuleButton" onclick="toggleAnnotationRuleEdit()">编辑规则字段</button> <span class="status" id="annotationCurrentRow">未选择</span></span></div>
          <div class="annotation-pane-body">
            <div class="annotation-meta" id="annotationMeta"><div class="annotation-empty" style="min-height:180px">选择左侧规则开始标注</div></div>
            <div class="annotation-decision">
              <label>是否触发（M列）</label>
              <div class="annotation-choice"><button id="annotationMYes" onclick="setAnnotationM('是')">是</button><button id="annotationMNo" onclick="setAnnotationM('否')">否</button></div>
              <label>金标准结果（N列）</label>
              <div class="annotation-choice"><button id="annotationNYes" onclick="setAnnotationN('是')">是</button><button id="annotationNNo" onclick="setAnnotationN('否')">否</button><button id="annotationNSlash" onclick="setAnnotationN('/')">/</button></div>
              <textarea id="annotationN" placeholder="触发时填写标准结果；不触发时自动填 /" oninput="updateAnnotationN(this.value)"></textarea>
              <button style="margin-top:10px" onclick="saveAnnotationCurrent()">保存当前标注</button>
            </div>
          </div>
        </section>
      </div>
    </section>
    <section id="identityModule" class="module-panel hide">
      <div class="module-heading">
        <div>
          <h2>主体身份卷</h2>
          <p>当前只开放入口和案件信息获取；规则标注、数据获取、复核工作台、专项统计和比对历史后续接入。</p>
        </div>
      </div>
      <div class="identity-placeholder">主体身份卷功能待接入。请先在上方输入 caseId 获取案件基础信息。</div>
    </section>
    <section id="statsModule" class="module-panel hide">
      <div class="module-heading">
        <div>
          <h2>专项统计</h2>
          <p>集中查看签名、日期、捺印、印章、综合、其他维度，以及模型错误原因汇总。</p>
        </div>
        <button class="secondary" onclick="renderStatsModule(true)">刷新统计</button>
      </div>
      <section class="dimension-panel" id="dimensionPanel"></section>
      <section class="error-reason-panel" id="errorReasonPanel"></section>
    </section>
    <section id="historyModule" class="module-panel hide">
      <div class="module-heading">
        <div>
          <h2>比对历史</h2>
          <p>保存本次比对的文件信息、匹配结果、人工修改和统计快照，可随时重新打开。</p>
        </div>
      </div>
      <div class="history-toolbar">
        <button class="secondary" onclick="loadHistoryList(true)">刷新历史</button>
        <button class="secondary" onclick="selectHistoryPage()">全选当前页</button>
        <button class="secondary" onclick="selectAllHistory()">全选全部</button>
        <button class="secondary" onclick="clearHistorySelection()">清空选择</button>
        <button id="exportHistoryButton" onclick="exportSelectedHistories()" disabled>导出已选（0）</button>
        <span class="status" id="historyStatus"></span>
      </div>
      <div id="historyList" class="history-list"></div>
      <div id="historyPagination" class="history-pagination hide"></div>
      <div id="historyDetail" class="history-detail"></div>
    </section>
      </div>
    </div>
    <datalist id="candidateList"></datalist>
    <div id="pickerOverlay" class="picker-overlay hide" onclick="overlayClick(event)">
      <div class="picker-dialog" role="dialog" aria-modal="true">
        <div class="picker-head">
          <div>
            <div class="picker-title" id="pickerTitle">选择</div>
            <div class="picker-subtitle" id="pickerSubtitle"></div>
          </div>
          <button class="picker-close" onclick="closePicker()">关闭</button>
        </div>
        <div class="picker-search">
          <input id="pickerSearch" placeholder="搜索编号、材料、规则或描述" oninput="onPickerSearch()">
        </div>
        <div id="pickerList" class="picker-list"></div>
        <div id="pickerPagination" class="picker-pagination hide"></div>
      </div>
    </div>
    <div id="historySaveOverlay" class="history-save-overlay hide" onclick="historySaveOverlayClick(event)">
      <div class="history-save-dialog" role="dialog" aria-modal="true" aria-labelledby="historySaveTitle">
        <div class="history-save-head">
          <div>
            <div class="history-save-title" id="historySaveTitle">保存比对记录</div>
            <div class="history-save-subtitle">填写本次比对名称和模型阅卷时间，点击确定后生成历史记录。</div>
          </div>
          <button class="history-save-close" type="button" aria-label="关闭" onclick="closeHistorySaveDialog()">×</button>
        </div>
        <div class="history-save-body">
          <div>
            <label for="historyName">比对名称</label>
            <input id="historyName" placeholder="可不填，默认使用 JSON 文件名 × Excel 文件名">
          </div>
          <div class="history-save-time-grid">
            <div>
              <label for="modelReviewStart">模型阅卷开始时间</label>
              <input id="modelReviewStart" type="datetime-local" step="1" onchange="updateReviewDurationPreview()">
            </div>
            <div>
              <label for="modelReviewEnd">模型阅卷结束时间</label>
              <input id="modelReviewEnd" type="datetime-local" step="1" onchange="updateReviewDurationPreview()">
            </div>
          </div>
          <div class="history-save-duration" id="modelReviewDuration">阅卷时长：未计算</div>
        </div>
        <div class="history-save-actions">
          <button class="secondary" type="button" onclick="closeHistorySaveDialog()">取消</button>
          <button type="button" onclick="saveCurrentHistory()">确定保存</button>
        </div>
      </div>
    </div>
    <div id="toastContainer" class="toast-container" aria-live="polite" aria-atomic="true"></div>
  </main>
  <script>
    const defaults = __DEFAULTS__;
    let analysis = null;
    let candidateById = {};
    let assignments = {};
    let materialGroups = [];
    let materialSelections = {};
    let verdictOverrides = {};
    let categoryOverrides = {};
    let errorReasonOverrides = {};
    let goldDirtyRows = new Set();
    let selectedSourceMaterial = "all";
    let pickerState = null;
    const PICKER_PAGE_SIZE = 5;
    let activeVolume = localStorage.getItem('matchingToolActiveVolume') || 'procedure';
    let activeModule = 'annotation';
    let currentCaseInfo = null;
    let historyItems = [];
    let moduleMenuCollapsed = false;
    let materialsCollapsed = false;
    let completedMaterials = {};
    let confirmedRows = {};
    let materialListPage = 1;
    let historyStatsCollapsed = {};
    let activeHistoryRecord = null;
    let historyPage = 1;
    let historySelectedIds = new Set();
    let annotationWorkbook = null;
    let annotationPdfUrl = '';
    let annotationCurrentIndex = -1;
    let annotationMColumn = 13;
    let annotationNColumn = 14;
    let annotationEditMode = false;
    let annotationMaterialType = '';
    let annotationExpandedMaterials = new Set();
    let annotationStatusFilter = 'all';
    let activeAnalysisJobId = '';
    let modelResultSource = localStorage.getItem('matchingToolModelSource') || 'upload';
    let fetchedModelResult = null;
    let activeBatchJobId = '';
    const MATERIAL_PAGE_SIZE = 10;
    const HISTORY_PAGE_SIZE = 5;
    const errorReasonHeaders = ['错误原因', '来源文件', '材料类型', '对应材料名称', '所属阶段', '阶段时间顺序', '规则名称', '特殊条件', '判断规则', '错误场景', '是否触发', '是否正确', '模型是否触发', '模型是否正确', '具体问题内容'];
    setTheme(localStorage.getItem('matchingToolTheme') || 'review');
    setModuleMenuCollapsed(localStorage.getItem('matchingToolModuleMenuCollapsed') === '1');
    setMaterialsCollapsed(localStorage.getItem('matchingToolMaterialsCollapsed') === '1');
    initCaseInfoControls();
    setVolume(activeVolume, false);
    setAnalysisVisibility(false);
    initModelSourceControls();
    updateFileReadyState();

    function escapeHtml(value) {
      return text(value).replace(/[&<>"']/g, ch => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[ch]));
    }
    function jsString(value) {
      return text(value).replace(/\\/g, '\\\\').replace(/'/g, "\\'").replace(/\n/g, '\\n');
    }
    function pct(v) { return (v * 100).toFixed(2) + '%'; }
    function text(v) { return v === null || v === undefined ? '' : String(v); }
    function showToast(message, type='info', duration=3600) {
      const container = document.getElementById('toastContainer');
      if (!container) return;
      const toast = document.createElement('div');
      const kind = ['success', 'error', 'info'].includes(type) ? type : 'info';
      toast.className = `toast ${kind}`;
      const mark = document.createElement('span');
      mark.className = 'toast-mark';
      mark.textContent = kind === 'success' ? '✓' : kind === 'error' ? '!' : 'i';
      const content = document.createElement('span');
      content.textContent = text(message);
      toast.append(mark, content);
      container.appendChild(toast);
      window.setTimeout(() => {
        toast.classList.add('is-leaving');
        window.setTimeout(() => toast.remove(), 220);
      }, Math.max(1200, Number(duration) || 3600));
    }
    function showSuccess(message) { showToast(message, 'success'); }
    function showFailure(message) { showToast(message, 'error', 5200); }
    function setTheme(theme) {
      const activeTheme = ['review', 'compact'].includes(theme) ? theme : 'review';
      document.body.dataset.theme = activeTheme === 'review' ? '' : activeTheme;
      localStorage.setItem('matchingToolTheme', activeTheme);
      for (const name of ['review','compact']) {
        const btn = document.getElementById(`theme-${name}`);
        if (btn) btn.classList.toggle('active', name === activeTheme);
      }
    }
    function setModuleMenuCollapsed(collapsed) {
      moduleMenuCollapsed = Boolean(collapsed);
      const layout = document.getElementById('appLayout');
      const menu = document.getElementById('moduleMenu');
      const button = document.getElementById('toggleModuleMenu');
      if (layout) layout.classList.toggle('module-menu-collapsed', moduleMenuCollapsed);
      if (menu) menu.classList.toggle('collapsed', moduleMenuCollapsed);
      if (button) {
        button.textContent = moduleMenuCollapsed ? '›' : '‹';
        button.title = moduleMenuCollapsed ? '展开菜单' : '收起菜单';
        button.setAttribute('aria-label', moduleMenuCollapsed ? '展开菜单' : '收起菜单');
        button.setAttribute('aria-expanded', String(!moduleMenuCollapsed));
      }
      localStorage.setItem('matchingToolModuleMenuCollapsed', moduleMenuCollapsed ? '1' : '0');
    }
    function toggleModuleMenu() {
      setModuleMenuCollapsed(!moduleMenuCollapsed);
    }
    function setMaterialsCollapsed(collapsed) {
      materialsCollapsed = Boolean(collapsed);
      const workspace = document.getElementById('reviewWorkspace');
      const button = document.getElementById('toggleMaterials');
      if (workspace) workspace.classList.toggle('materials-collapsed', materialsCollapsed);
      if (button) {
        button.textContent = materialsCollapsed ? '›' : '‹';
        button.title = materialsCollapsed ? '展开材料列表' : '收起材料列表';
        button.setAttribute('aria-label', materialsCollapsed ? '展开材料列表' : '收起材料列表');
        button.setAttribute('aria-expanded', String(!materialsCollapsed));
      }
      localStorage.setItem('matchingToolMaterialsCollapsed', materialsCollapsed ? '1' : '0');
    }
    function toggleMaterials() {
      setMaterialsCollapsed(!materialsCollapsed);
    }
    function initCaseInfoControls() {
      const input = document.getElementById('caseInfoBaseUrl');
      if (input) {
        input.value = localStorage.getItem('matchingToolCaseDetailApi') || defaults.caseDetailApi || '';
        input.addEventListener('input', () => localStorage.setItem('matchingToolCaseDetailApi', input.value));
      }
    }
    function setVolume(volume, keepModule=true) {
      activeVolume = volume === 'identity' ? 'identity' : 'procedure';
      localStorage.setItem('matchingToolActiveVolume', activeVolume);
      document.getElementById('volume-procedure')?.classList.toggle('active', activeVolume === 'procedure');
      document.getElementById('volume-identity')?.classList.toggle('active', activeVolume === 'identity');
      document.getElementById('procedureSubmenu')?.classList.toggle('hide', activeVolume !== 'procedure');
      const title = document.getElementById('casePanelTitle');
      if (title) title.textContent = `${activeVolume === 'procedure' ? '程序卷' : '主体身份卷'} / 案件信息`;
      if (activeVolume === 'identity') {
        setModule('identity');
      } else if (!keepModule || activeModule === 'identity') {
        setModule('annotation');
      } else {
        setModule(activeModule);
      }
      renderCaseInfo();
    }
    function caseText(value) {
      return text(value).trim();
    }
    function detectProcedureTemplate(caseInfo) {
      const haystack = [
        caseInfo.caseNature,
        caseInfo.nature,
        caseInfo.caseKinds,
        caseInfo.caseKindNames,
        caseInfo.caseType,
        caseInfo.caseTypeConfirm,
        caseInfo.sourceType,
        caseInfo.violationType,
        caseInfo.dutyCrime,
        caseInfo.name,
      ].map(value => Array.isArray(value) ? value.join('、') : text(value)).join(' ');
      const isLz = Number(caseInfo.isLz || 0) === 1 || Number(caseInfo.isExtendLz || 0) === 1 || /留置/.test(haystack);
      const isDiscipline = /违纪|职务违法|审查调查|初核/.test(haystack);
      const isTown = /乡镇|镇街|街道/.test(haystack);
      const isDispatched = /派驻/.test(haystack);
      if (isLz) return {name:'自办留置模板', reason:'命中优先级：自办留置。当前依据包含 isLz/isExtendLz 或“留置”关键词。'};
      if (isDiscipline) return {name:'自办违纪模板', reason:'命中优先级：自办违纪。当前依据包含违纪、职务违法、审查调查或初核关键词。'};
      if (isTown) return {name:'乡镇模板', reason:'命中优先级：乡镇。当前依据包含乡镇、镇街或街道关键词。'};
      if (isDispatched) return {name:'派驻模板', reason:'命中优先级：派驻。当前依据包含派驻关键词。'};
      return {name:'待人工选择', reason:'暂未命中已配置规则，后续可补充具体判断依据。'};
    }
    function renderCaseInfo() {
      const infoWrap = document.getElementById('caseInfoContent');
      const templateWrap = document.getElementById('caseTemplateResult');
      const status = document.getElementById('caseInfoStatus');
      if (!infoWrap || !templateWrap) return;
      if (!currentCaseInfo) {
        infoWrap.classList.add('hide');
        templateWrap.classList.add('hide');
        if (status) status.textContent = '未获取案件信息';
        return;
      }
      const fileCount = Array.isArray(currentCaseInfo.fileList) ? currentCaseInfo.fileList.length : 0;
      const fields = [
        ['案件名称', currentCaseInfo.name],
        ['案件性质', currentCaseInfo.caseNature],
        ['案件类型', currentCaseInfo.caseType],
        ['案件确认类型', currentCaseInfo.caseTypeConfirm],
        ['是否留置', Number(currentCaseInfo.isLz || 0) === 1 ? '是' : '否'],
        ['被审查人', currentCaseInfo.reviewName],
        ['职务', currentCaseInfo.reviewPosition],
        ['职级', currentCaseInfo.reviewRank],
        ['身份', currentCaseInfo.civilServant],
        ['特殊身份', currentCaseInfo.specialIdentity],
        ['违纪违法类型', currentCaseInfo.violationType],
        ['材料数量', fileCount],
      ];
      infoWrap.innerHTML = fields.map(([label, value]) => `<div class="case-info-item"><div class="case-info-label">${escapeHtml(label)}</div><div class="case-info-value">${escapeHtml(caseText(value) || '—')}</div></div>`).join('');
      infoWrap.classList.remove('hide');
      if (activeVolume === 'procedure') {
        const detected = detectProcedureTemplate(currentCaseInfo);
        templateWrap.innerHTML = `<b>程序卷推荐模板：${escapeHtml(detected.name)}</b><br>${escapeHtml(detected.reason)}<br><span class="small">当前优先级：自办留置 &gt; 自办违纪 &gt; 乡镇 &gt; 派驻</span>`;
      } else {
        templateWrap.innerHTML = `<b>主体身份卷：</b>案件信息已获取，具体模板和业务功能待接入。`;
      }
      templateWrap.classList.remove('hide');
      if (status) status.textContent = `已获取 caseId=${caseText(currentCaseInfo.id || currentCaseInfo.caseId)}`;
    }
    async function fetchCaseInfo() {
      const caseId = document.getElementById('caseInfoCaseId')?.value.trim();
      const baseUrl = document.getElementById('caseInfoBaseUrl')?.value.trim();
      if (!caseId) { showFailure('请先填写 caseId。'); return; }
      const status = document.getElementById('caseInfoStatus');
      if (status) status.textContent = '正在获取案件信息...';
      try {
        const result = await getJson(`/api/case/detail?caseId=${encodeURIComponent(caseId)}&baseUrl=${encodeURIComponent(baseUrl || defaults.caseDetailApi || '')}`);
        currentCaseInfo = result.data || {};
        currentCaseInfo.caseId = result.caseId || currentCaseInfo.id || caseId;
        renderCaseInfo();
        showSuccess('案件信息获取成功。');
      } catch (err) {
        if (status) status.textContent = err.message;
        showFailure(`案件信息获取失败：${err.message}`);
      }
    }
    function setAnalysisVisibility(ready) {
      const visible = Boolean(ready);
      document.getElementById('cards')?.classList.toggle('hide', !visible);
      document.getElementById('reviewWorkspace')?.classList.toggle('hide', !visible);
      document.getElementById('exportButton')?.toggleAttribute('disabled', !visible);
      document.getElementById('saveHistoryButton')?.toggleAttribute('disabled', !visible);
      document.getElementById('analysisEmpty')?.classList.toggle('hide', visible);
    }
    function initModelSourceControls() {
      const apiDefault = localStorage.getItem('matchingToolApiBaseUrl') || defaults.modelResultApi || '';
      const batchDefault = localStorage.getItem('matchingToolBatchBaseUrl') || apiDefault;
      const apiInput = document.getElementById('apiBaseUrl');
      const batchInput = document.getElementById('batchBaseUrl');
      if (apiInput) {
        apiInput.value = apiDefault;
        apiInput.addEventListener('input', () => localStorage.setItem('matchingToolApiBaseUrl', apiInput.value));
      }
      if (batchInput) {
        batchInput.value = batchDefault;
        batchInput.addEventListener('input', () => localStorage.setItem('matchingToolBatchBaseUrl', batchInput.value));
      }
      setModelSource(modelResultSource, false);
      updateBatchModeUI();
    }
    function setModelSource(source, notify=true) {
      const nextSource = ['upload', 'api', 'batch'].includes(source) ? source : 'upload';
      const changed = nextSource !== modelResultSource;
      modelResultSource = nextSource;
      localStorage.setItem('matchingToolModelSource', modelResultSource);
      if (changed) {
        clearFetchedModelResult();
        clearJsonPreview();
      }
      for (const name of ['upload', 'api', 'batch']) {
        document.getElementById(`source-${name}`)?.classList.toggle('active', modelResultSource === name);
      }
      document.getElementById('uploadSourcePanel')?.classList.toggle('hide', modelResultSource !== 'upload');
      document.getElementById('apiSourcePanel')?.classList.toggle('hide', modelResultSource !== 'api');
      document.getElementById('batchSourcePanel')?.classList.toggle('hide', modelResultSource !== 'batch');
      updateBatchModeUI();
      const button = document.getElementById('analyzeButton');
      if (button) button.textContent = modelResultSource === 'api' ? '分析接口结果' : '分析';
      if (notify && modelResultSource === 'batch') showToast('批量测试会逐个 caseId 拉取接口并生成汇总。', 'info', 2400);
      updateFileReadyState();
    }
    function clearFetchedModelResult() {
      fetchedModelResult = null;
    }
    function clearJsonPreview() {
      const wrap = document.getElementById('jsonPreview');
      const meta = document.getElementById('jsonPreviewMeta');
      const body = document.getElementById('jsonPreviewBody');
      if (wrap) wrap.classList.add('hide');
      if (meta) meta.textContent = '未拉取结果';
      if (body) body.textContent = '';
    }
    function modelResultReady() {
      if (modelResultSource === 'upload') return Boolean(document.getElementById('jsonFile')?.files?.length);
      if (modelResultSource === 'api') return Boolean(fetchedModelResult && fetchedModelResult.content);
      return false;
    }
    function updateFileReadyState() {
      const jsonReady = modelResultReady();
      const excelReady = Boolean(document.getElementById('excelFile')?.files?.length);
      const ready = jsonReady && excelReady;
      const button = document.getElementById('analyzeButton');
      if (button) button.disabled = !ready || Boolean(activeAnalysisJobId);
      refreshWorkbenchDatasetInfo();
      if (!analysis) {
        const needModel = modelResultSource === 'api' ? '请先拉取接口模型结果' : '请先上传 JSON / TXT';
        document.getElementById('status').textContent = ready ? '数据已就绪，可以开始分析' : `${needModel}，并上传 Excel 文件`;
      }
    }
    function resetMaterialListPage() {
      materialListPage = 1;
    }
    function setModule(name) {
      activeModule = name;
      for (const moduleName of ['workbench', 'annotation', 'data', 'stats', 'history', 'identity']) {
        const panel = document.getElementById(`${moduleName}Module`);
        const tab = document.getElementById(`module-${moduleName}`);
        if (panel) panel.classList.toggle('hide', moduleName !== name);
        if (tab) tab.classList.toggle('active', moduleName === name);
      }
      if (name !== 'identity' && activeVolume !== 'procedure') {
        activeVolume = 'procedure';
        localStorage.setItem('matchingToolActiveVolume', activeVolume);
        document.getElementById('volume-procedure')?.classList.add('active');
        document.getElementById('volume-identity')?.classList.remove('active');
        document.getElementById('procedureSubmenu')?.classList.remove('hide');
        const title = document.getElementById('casePanelTitle');
        if (title) title.textContent = '程序卷 / 案件信息';
        renderCaseInfo();
      }
      if (name === 'workbench') refreshWorkbenchDatasetInfo();
      if (name === 'stats') renderStatsModule();
      if (name === 'history') loadHistoryList();
    }

    function annotationCurrentSheet() {
      if (!annotationWorkbook) return null;
      return annotationWorkbook.sheets.find(sheet => sheet.name === annotationWorkbook.currentSheet) || annotationWorkbook.sheets[0];
    }
    function annotationFieldValue(row, index) {
      return text(row?.values?.[index] ?? '');
    }
    function renderAnnotationSheetOptions() {
      const select = document.getElementById('annotationSheet');
      if (!select || !annotationWorkbook) return;
      select.innerHTML = annotationWorkbook.sheets.map(sheet => `<option value="${escapeHtml(sheet.name)}">${escapeHtml(sheet.name)}</option>`).join('');
      select.value = annotationWorkbook.currentSheet;
      select.disabled = annotationWorkbook.sheets.length < 2;
      renderAnnotationColumnOptions();
    }
    function renderAnnotationColumnOptions() {
      const sheet = annotationCurrentSheet();
      if (!sheet) return;
      const maxColumns = Math.max(...annotationWorkbook.sheets.map(item => item.headers.length), 14);
      const options = Array.from({length: maxColumns}, (_, index) => `<option value="${index + 1}">${index + 1}. ${escapeHtml(sheet.headers[index] || `第${index + 1}列`)}</option>`).join('');
      const mSelect = document.getElementById('annotationMColumn');
      const nSelect = document.getElementById('annotationNColumn');
      if (mSelect) { mSelect.innerHTML = options; mSelect.value = '13'; mSelect.disabled = true; }
      if (nSelect) { nSelect.innerHTML = options; nSelect.value = '14'; nSelect.disabled = true; }
    }
    function applyAnnotationColumnMapping() {
      const m = Number(document.getElementById('annotationMColumn')?.value || 13);
      const n = Number(document.getElementById('annotationNColumn')?.value || 14);
      if (m === n) { showFailure('M列和N列必须选择不同的列。'); renderAnnotationColumnOptions(); return; }
      annotationMColumn = m;
      annotationNColumn = n;
      for (const sheet of annotationWorkbook?.sheets || []) {
        for (const row of sheet.rows) {
          row.m = text(row.values[m - 1] || '').trim();
          row.n = text(row.values[n - 1] || '').trim();
        }
      }
      renderAnnotationCurrent();
      renderAnnotationRuleList();
      renderAnnotationTable();
      document.getElementById('annotationStatus').textContent = `已切换列映射：M=${m}列，N=${n}列`;
    }
    function switchAnnotationSheet() {
      if (!annotationWorkbook) return;
      annotationWorkbook.currentSheet = document.getElementById('annotationSheet').value;
      annotationCurrentIndex = -1;
      annotationEditMode = false;
      annotationMaterialType = '';
      annotationExpandedMaterials.clear();
      renderAnnotationMaterialTypes();
      renderAnnotationRuleList();
      renderAnnotationTable();
      renderAnnotationCurrent();
      renderAnnotationNewFields();
    }
    function loadAnnotationPdf(input) {
      const file = input?.files?.[0];
      if (!file) return;
      if (annotationPdfUrl) URL.revokeObjectURL(annotationPdfUrl);
      annotationPdfUrl = URL.createObjectURL(file);
      const frame = document.getElementById('annotationPdfFrame');
      frame.src = annotationPdfUrl;
      frame.classList.remove('hide');
      document.getElementById('annotationPdfEmpty')?.classList.add('hide');
      const tableFrame = document.getElementById('annotationPdfFrameTable');
      if (tableFrame) { tableFrame.src = annotationPdfUrl; tableFrame.classList.remove('hide'); }
      document.getElementById('annotationPdfEmptyTable')?.classList.add('hide');
      document.getElementById('annotationStatus').textContent = `已加载 PDF：${file.name}（仅供对照查看）`;
    }
    async function loadAnnotationWorkbook() {
      const file = document.getElementById('annotationExcelFile')?.files?.[0];
      if (!file) { showFailure('请先上传规则 Excel。'); return; }
      const status = document.getElementById('annotationStatus');
      status.textContent = '正在读取规则 Excel...';
      try {
        const result = await post('/api/rules/load', { excelUpload: { name: file.name, base64: await readBase64File(file) } });
      annotationWorkbook = result;
      annotationWorkbook.newRules = [];
      annotationWorkbook.currentSheet = result.sheets.find(sheet => sheet.name === '全规则表')?.name || result.sheets[0]?.name || '';
      annotationMColumn = 13;
      annotationNColumn = 14;
        annotationMaterialType = '';
        annotationCurrentIndex = -1;
        annotationStatusFilter = 'all';
        annotationExpandedMaterials.clear();
        setAnnotationStatusFilter('all');
        renderAnnotationSheetOptions();
        renderAnnotationMaterialTypes();
        renderAnnotationRuleList();
        renderAnnotationTable();
        renderAnnotationCurrent();
        renderAnnotationNewFields();
      document.getElementById('annotationExportButton').disabled = false;
        status.textContent = `已加载 ${result.fileName}，共 ${result.sheets.reduce((sum, sheet) => sum + sheet.rows.length, 0)} 条规则`;
        showSuccess('规则 Excel 加载成功。');
      } catch (err) {
        status.textContent = err.message;
        showFailure(`加载失败：${err.message}`);
      }
    }
    function annotationRowIsComplete(row) {
      return Boolean(text(row?.m).trim() && text(row?.n).trim());
    }
    function annotationRowHasIssue(row) {
      const m = text(row?.m).trim();
      const n = text(row?.n).trim();
      return (m === '是' && !n) || (!m && n);
    }
    function annotationMaterialStatus(rows) {
      const total = rows.length;
      const completed = rows.filter(annotationRowIsComplete).length;
      if (rows.some(annotationRowHasIssue)) return 'issue';
      if (!completed) return 'pending';
      if (completed === total) return 'completed';
      return 'progress';
    }
    function annotationMaterialSummaries() {
      const sheet = annotationCurrentSheet();
      if (!sheet) return [];
      const groups = new Map();
      sheet.rows.forEach((row, index) => {
        const material = annotationFieldValue(row, 3) || '未分类材料';
        if (!groups.has(material)) groups.set(material, []);
        groups.get(material).push({row, index});
      });
      return [...groups.entries()].map(([material, items]) => {
        const rows = items.map(item => item.row);
        return {material, items, total:rows.length, completed:rows.filter(annotationRowIsComplete).length, status:annotationMaterialStatus(rows)};
      });
    }
    function setAnnotationStatusFilter(filter) {
      annotationStatusFilter = ['all', 'pending', 'progress', 'completed', 'issue'].includes(filter) ? filter : 'all';
      document.querySelectorAll('#annotationStatusFilter button').forEach(button => button.classList.toggle('active', button.dataset.filter === annotationStatusFilter));
      renderAnnotationRuleList();
    }
    function renderAnnotationStats() {
      const summaries = annotationMaterialSummaries();
      const totalFiles = summaries.length;
      const completedFiles = summaries.filter(item => item.status === 'completed').length;
      const totalRules = summaries.reduce((sum, item) => sum + item.total, 0);
      const completedRules = summaries.reduce((sum, item) => sum + item.completed, 0);
      const current = summaries.find(item => item.material === annotationMaterialType);
      const listStats = document.getElementById('annotationListStats');
      if (listStats) listStats.innerHTML = `<div>文件：<b>${completedFiles} / ${totalFiles}</b> 已完成</div><div>规则：<b>${completedRules} / ${totalRules}</b> 已标注</div>`;
      const fileStats = document.getElementById('annotationFileStats');
      if (fileStats) fileStats.textContent = `文件 ${completedFiles} / ${totalFiles}，规则 ${completedRules} / ${totalRules}`;
      const status = document.getElementById('annotationFileStatus');
      if (status) {
        const currentStatus = current?.status || 'pending';
        status.className = `annotation-file-status ${currentStatus}`;
          status.textContent = current ? `${current.completed}/${current.total} 条已标注 · ${currentStatus === 'completed' ? '已完成' : currentStatus === 'progress' ? '标注中' : currentStatus === 'issue' ? '有问题' : '未开始'}` : '未加载文件';
      }
    }
    function renderAnnotationRuleList() {
      const list = document.getElementById('annotationRuleList');
      const count = document.getElementById('annotationRowCount');
      const sheet = annotationCurrentSheet();
      if (!list || !sheet) { if (list) list.innerHTML = '<div class="annotation-empty" style="min-height:180px">打开 Excel 后显示规则</div>'; return; }
      const query = text(document.getElementById('annotationSearch')?.value).trim().toLowerCase();
      const items = sheet.rows.map((row, index) => ({row, index})).filter(({row}) => {
        const material = annotationFieldValue(row, 3) || '未分类材料';
        const summary = annotationMaterialSummaries().find(item => item.material === material);
        if (annotationStatusFilter !== 'all' && summary?.status !== annotationStatusFilter) return false;
        if (!query) return true;
        return row.values.some(value => text(value).toLowerCase().includes(query));
      });
      count.textContent = `${items.length} 条`;
      const groups = new Map();
      items.forEach(item => {
        const material = annotationFieldValue(item.row, 3) || '未分类材料';
        if (!groups.has(material)) groups.set(material, []);
        groups.get(material).push(item);
      });
      list.innerHTML = groups.size ? [...groups.entries()].map(([material, group]) => {
        const summary = annotationMaterialSummaries().find(item => item.material === material);
        const expanded = annotationExpandedMaterials.has(material);
        const content = group.map(({row, index}) => {
          const active = index === annotationCurrentIndex ? ' active' : '';
          const rule = annotationFieldValue(row, 7) || '未命名规则';
          const scene = annotationFieldValue(row, 11) || '未填写错误场景';
          const mark = row.m === '是' ? '●' : row.m === '否' ? '○' : '·';
          return `<button class="annotation-rule-item${active}" onclick="selectAnnotationRow(${index})"><b>${mark} ${escapeHtml(rule)}</b><small>${escapeHtml(scene)} · Excel第${row.rowNumber || '新增'}行</small></button>`;
        }).join('');
        const materialArg = escapeHtml(JSON.stringify(material));
        const statusText = summary?.status === 'completed' ? '已完成' : summary?.status === 'progress' ? '标注中' : summary?.status === 'issue' ? '有问题' : '未开始';
        return `<div class="annotation-material-group"><button class="annotation-material-toggle${expanded ? ' expanded' : ''}" onclick='toggleAnnotationMaterial(${materialArg})'><span class="material-toggle-label">${expanded ? '▾' : '▸'} ${escapeHtml(material)}</span><span class="material-toggle-meta ${summary?.status || 'pending'}">${summary?.completed || 0}/${summary?.total || group.length} · ${statusText}</span></button><div class="annotation-material-items${expanded ? '' : ' hide'}">${content}</div></div>`;
      }).join('') : '<div class="annotation-empty" style="min-height:180px">没有匹配的规则</div>';
      renderAnnotationStats();
    }
    function toggleAnnotationMaterial(material) {
      if (annotationExpandedMaterials.has(material)) annotationExpandedMaterials.delete(material);
      else annotationExpandedMaterials.add(material);
      renderAnnotationRuleList();
    }
    function renderAnnotationTable() {
      const head = document.getElementById('annotationTableHead');
      const body = document.getElementById('annotationTableBody');
      const count = document.getElementById('annotationTableCount');
      const sheet = annotationCurrentSheet();
      if (!head || !body || !sheet) return;
      const query = text(document.getElementById('annotationTableSearch')?.value).trim().toLowerCase();
      const items = sheet.rows.map((row, index) => ({row, index})).filter(({row}) => {
        const sameMaterial = !annotationMaterialType || annotationFieldValue(row, 3) === annotationMaterialType;
        return sameMaterial && (!query || row.values.some(value => text(value).toLowerCase().includes(query)));
      });
      const sourceHeaders = [...sheet.headers.slice(0, 12), sheet.headers[12] || 'M列：是否触发', sheet.headers[13] || 'N列：是否正确'];
      const headers = [...sourceHeaders.slice(0, 12), 'M｜是否触发', 'N｜是否正确'];
      count.textContent = `${annotationMaterialType || '全部材料'} · ${items.length} 条`;
      const widths = [8,16,16,22,28,16,12,34,24,52,24,34,110,180];
      head.innerHTML = `<tr>${headers.map((header, index) => `<th title="${escapeHtml(sourceHeaders[index])}" style="width:${widths[index]}px;max-width:${widths[index]}px;white-space:normal;overflow-wrap:anywhere;word-break:break-all;">${escapeHtml(header)}</th>`).join('')}</tr>`;
      const cellSpan = (itemIndex, cellIndex) => {
        if (cellIndex > 10) return 1;
        const current = annotationFieldValue(items[itemIndex].row, cellIndex);
        if (!current || (itemIndex > 0 && annotationFieldValue(items[itemIndex - 1].row, cellIndex) === current)) return 0;
        let span = 1;
        while (itemIndex + span < items.length && annotationFieldValue(items[itemIndex + span].row, cellIndex) === current) span += 1;
        return span;
      };
      body.innerHTML = items.length ? items.map(({row, index}, itemIndex) => {
        const selected = index === annotationCurrentIndex ? ' selected' : '';
        const cells = row.values.slice(0, 12).map((value, cellIndex) => {
          const span = cellSpan(itemIndex, cellIndex);
          if (!span) return '';
          const rowSpan = span > 1 ? ` rowspan="${span}"` : '';
          return `<td${rowSpan} style="width:${widths[cellIndex]}px"><button class="table-row-button" onclick="selectAnnotationRow(${index})">${escapeHtml(annotationFieldValue(row, cellIndex) || '—')}</button></td>`;
        }).join('');
        return `<tr class="${selected}">${cells}<td style="width:${widths[12]}px"><div class="table-choice"><button class="${row.m === '是' ? 'active' : ''}" onclick="setAnnotationMForRow(${index}, '是')">是</button><button class="${row.m === '否' ? 'active' : ''}" onclick="setAnnotationMForRow(${index}, '否')">否</button></div></td><td style="width:${widths[13]}px"><div class="table-choice"><button class="${row.n === '是' ? 'active' : ''}" onclick="setAnnotationNForRow(${index}, '是')">是</button><button class="${row.n === '否' ? 'active' : ''}" onclick="setAnnotationNForRow(${index}, '否')">否</button><button class="${row.n === '/' ? 'active' : ''}" onclick="setAnnotationNForRow(${index}, '/')">/</button></div><textarea class="table-n-input" onchange="setAnnotationNForRow(${index}, this.value)" onclick="event.stopPropagation()">${escapeHtml(row.n || '')}</textarea></td></tr>`;
      }).join('') : '<tr><td colspan="14" class="annotation-empty" style="min-height:180px">没有匹配的规则</td></tr>';
    }
    function getAnnotationMaterialTypes() {
      const sheet = annotationCurrentSheet();
      if (!sheet) return [];
      return [...new Set(sheet.rows.map(row => annotationFieldValue(row, 3) || '未分类材料'))];
    }
    function renderAnnotationMaterialTypes() {
      const select = document.getElementById('annotationMaterialType');
      if (!select) return;
      const types = getAnnotationMaterialTypes();
      if (!types.includes(annotationMaterialType)) annotationMaterialType = types[0] || '';
      select.innerHTML = types.map(type => `<option value="${escapeHtml(type)}">${escapeHtml(type)}</option>`).join('');
      select.value = annotationMaterialType;
      select.disabled = types.length === 0;
    }
    function switchAnnotationMaterialType() {
      annotationMaterialType = document.getElementById('annotationMaterialType')?.value || '';
      annotationCurrentIndex = -1;
      renderAnnotationRuleList();
      renderAnnotationTable();
      renderAnnotationCurrent();
    }
    function moveAnnotationMaterial(direction) {
      const types = getAnnotationMaterialTypes();
      if (!types.length) return;
      const current = Math.max(0, types.indexOf(annotationMaterialType));
      annotationMaterialType = types[(current + direction + types.length) % types.length];
      renderAnnotationMaterialTypes();
      annotationCurrentIndex = -1;
      renderAnnotationRuleList();
      renderAnnotationTable();
      renderAnnotationCurrent();
    }
    function setAnnotationMForRow(index, value) {
      selectAnnotationRow(index);
      setAnnotationM(value);
    }
    function setAnnotationNForRow(index, value) {
      selectAnnotationRow(index);
      setAnnotationN(value);
    }
    function toggleAnnotationView() {
      const table = document.getElementById('annotationTableView');
      const grouped = document.getElementById('annotationLayout');
      const button = document.getElementById('annotationViewButton');
      const showTable = table?.classList.contains('hide');
      table?.classList.toggle('hide', !showTable);
      grouped?.classList.toggle('hide', showTable);
      if (button) button.textContent = showTable ? '分组列表' : 'Excel表格';
    }
    function selectAnnotationRow(index) {
      annotationCurrentIndex = index;
      renderAnnotationRuleList();
      renderAnnotationTable();
      renderAnnotationCurrent();
    }
    function renderAnnotationCurrent() {
      const sheet = annotationCurrentSheet();
      const row = sheet?.rows?.[annotationCurrentIndex];
      const meta = document.getElementById('annotationMeta');
      if (!row) {
        document.getElementById('annotationCurrentRow').textContent = '未选择';
        meta.innerHTML = '<div class="annotation-empty" style="min-height:180px">选择左侧规则开始标注</div>';
        document.getElementById('annotationN').value = '';
        document.getElementById('annotationMYes').classList.remove('active');
        document.getElementById('annotationMNo').classList.remove('active');
        return;
      }
      document.getElementById('annotationCurrentRow').textContent = `Excel第${row.rowNumber}行`;
      const headers = sheet.headers.slice(0, 12);
      meta.innerHTML = headers.map((header, index) => {
        const value = annotationFieldValue(row, index);
        const multiline = [9, 10, 11].includes(index);
        const control = multiline
          ? `<textarea oninput="updateAnnotationRuleField(${index}, this.value)">${escapeHtml(value)}</textarea>`
          : `<input value="${escapeHtml(value)}" oninput="updateAnnotationRuleField(${index}, this.value)">`;
        return `<div class="annotation-meta-item"><div class="annotation-meta-label">${escapeHtml(header)}</div>${annotationEditMode ? control : `<div class="annotation-meta-value">${escapeHtml(value || '—')}</div>`}</div>`;
      }).join('');
      document.getElementById('annotationEditRuleButton').textContent = annotationEditMode ? '完成编辑' : '编辑规则字段';
      document.getElementById('annotationN').value = row.n || '';
      document.getElementById('annotationMYes').classList.toggle('active', row.m === '是');
      document.getElementById('annotationMNo').classList.toggle('active', row.m === '否');
      document.getElementById('annotationNYes').classList.toggle('active', row.n === '是');
      document.getElementById('annotationNNo').classList.toggle('active', row.n === '否');
      document.getElementById('annotationNSlash').classList.toggle('active', row.n === '/');
    }
    function toggleAnnotationRuleEdit() {
      if (!annotationCurrentSheet()?.rows?.[annotationCurrentIndex]) return;
      annotationEditMode = !annotationEditMode;
      renderAnnotationCurrent();
    }
    function updateAnnotationRuleField(index, value) {
      const row = annotationCurrentSheet()?.rows?.[annotationCurrentIndex];
      if (!row) return;
      row.values[index] = value;
    }
    function setAnnotationM(value) {
      const sheet = annotationCurrentSheet();
      const row = sheet?.rows?.[annotationCurrentIndex];
      if (!row) return;
      row.m = value;
      row.values[annotationMColumn - 1] = value;
      if (value === '否') { row.n = '/'; row.values[annotationNColumn - 1] = '/'; }
      renderAnnotationCurrent();
      renderAnnotationRuleList();
      renderAnnotationTable();
    }
    function setAnnotationN(value) {
      updateAnnotationN(value);
      renderAnnotationRuleList();
      renderAnnotationCurrent();
      renderAnnotationTable();
    }
    function updateAnnotationN(value) {
      const sheet = annotationCurrentSheet();
      const row = sheet?.rows?.[annotationCurrentIndex];
      if (!row) return;
      row.n = value;
      row.values[annotationNColumn - 1] = value;
      document.getElementById('annotationNYes').classList.toggle('active', value === '是');
      document.getElementById('annotationNNo').classList.toggle('active', value === '否');
      document.getElementById('annotationNSlash').classList.toggle('active', value === '/');
      renderAnnotationStats();
    }
    function saveAnnotationCurrent() {
      const sheet = annotationCurrentSheet();
      const row = sheet?.rows?.[annotationCurrentIndex];
      if (!row) { showFailure('请先选择一条规则。'); return; }
      if (row.m === '是' && !text(row.n).trim()) { showFailure('选择“是”后，请填写 N 列金标准结果。'); return; }
      document.getElementById('annotationStatus').textContent = `已保存当前标注：Excel第${row.rowNumber}行`;
      showSuccess('当前标注已保存到页面。');
    }
    function toggleNewRuleForm() {
      document.getElementById('annotationNewRuleForm')?.classList.toggle('hide');
      renderAnnotationNewFields();
    }
    function renderAnnotationNewFields() {
      const wrap = document.getElementById('annotationNewFields');
      const sheet = annotationCurrentSheet();
      if (!wrap || !sheet) return;
      wrap.innerHTML = sheet.headers.slice(0, 12).map((header, index) => {
        const multiline = [9, 10, 11].includes(index);
        const control = multiline ? `<textarea id="annotationNew-${index}" placeholder="${escapeHtml(header)}"></textarea>` : `<input id="annotationNew-${index}" placeholder="${escapeHtml(header)}">`;
        return `<div><label>${escapeHtml(header)}</label>${control}</div>`;
      }).join('');
    }
    function addAnnotationRule() {
      const sheet = annotationCurrentSheet();
      if (!sheet) { showFailure('请先打开规则 Excel。'); return; }
      const values = sheet.headers.slice(0, 12).map((_, index) => document.getElementById(`annotationNew-${index}`)?.value || '');
      if (!values[7] && !values[9] && !values[11]) { showFailure('至少填写规则名称、判断规则或错误场景。'); return; }
      const rowValues = [...values];
      while (rowValues.length < Math.max(14, annotationMColumn, annotationNColumn)) rowValues.push('');
      const row = {rowNumber: null, values: rowValues, m: '', n: ''};
      const requestedRow = Number(document.getElementById('annotationNewPosition')?.value || 0);
      const insertIndex = requestedRow >= 2 ? Math.min(sheet.rows.length, requestedRow - 2) : sheet.rows.length;
      sheet.rows.splice(insertIndex, 0, row);
      annotationWorkbook.newRules.push({sheetName: sheet.name, values: row.values, insertRow: requestedRow || null});
      annotationCurrentIndex = insertIndex;
      document.getElementById('annotationNewPosition').value = '';
      document.getElementById('annotationNewRuleForm').classList.add('hide');
      renderAnnotationRuleList();
      renderAnnotationCurrent();
      document.getElementById('annotationStatus').textContent = '新规则已加入列表，填写 M、N 后再导出。';
      renderAnnotationMaterialTypes();
      renderAnnotationTable();
    }
    async function exportAnnotationWorkbook() {
      if (!annotationWorkbook) { showFailure('请先打开规则 Excel。'); return; }
      const edits = [];
      for (const sheet of annotationWorkbook.sheets) {
        for (const row of sheet.rows) if (row.rowNumber) edits.push({sheetName: sheet.name, rowNumber: row.rowNumber, m: row.m, n: row.n});
      }
      const result = await post('/api/rules/export', {sourcePath: annotationWorkbook.excelPath, edits, newRules: annotationWorkbook.newRules || [], sheetName: annotationWorkbook.currentSheet, mColumn: annotationMColumn, nColumn: annotationNColumn});
      document.getElementById('annotationStatus').textContent = `已导出：${result.outputPath}`;
      showSuccess('规则金标准 Excel 已导出。');
    }
    function candidateLabel(candidate, score) {
      if (!candidate) return '无匹配';
      return `${score === undefined ? '' : '[' + Number(score).toFixed(2) + '] '}#${candidate.id} ${candidate.materialName} / ${candidate.ruleName}`;
    }
    function shortText(value, maxLen=70) {
      const s = text(value);
      return s.length > maxLen ? s.slice(0, maxLen) + '...' : s;
    }
    function materialLabel(group) {
      return `${group.index + 1}. ${group.displayName || group.materialName || '未命名材料'} / ${group.outlineName || '无对应材料名'}（${group.candidates.length}项）`;
    }
    function buildMaterialGroups() {
      materialGroups = [];
      const materialIndexByKey = {};
      for (const c of analysis.candidates) {
        const copyCount = Number(c.materialOccurrenceTotal || 1);
        const copyIndex = Number(c.materialOccurrence || 1);
        const key = `${c.materialName}||${c.outlineName}||${copyCount > 1 ? copyIndex : 0}`;
        if (materialIndexByKey[key] === undefined) {
          materialIndexByKey[key] = materialGroups.length;
          const suffix = copyCount > 1 ? `（第${copyIndex}份）` : '';
          materialGroups.push({
            index: materialGroups.length,
            materialName: c.materialName,
            outlineName: c.outlineName,
            displayName: `${c.materialName || '未命名材料'}${suffix}`,
            copyIndex,
            copyCount,
            candidates: []
          });
        }
        c.materialIndex = materialIndexByKey[key];
        materialGroups[c.materialIndex].candidates.push(c);
      }
    }
    function ruleLabel(candidate, score) {
      const scoreText = score === undefined ? '' : `推荐分${Number(score).toFixed(2)} `;
      return `#${candidate.id} ${scoreText}${candidate.ruleName} | 是否正确=${candidate.isCorrect} | ${shortText(candidate.questionDesc, 58)}`;
    }
    function scoreForRow(row, candidateId) {
      const item = row.options.find(opt => Number(opt.candidateId) === Number(candidateId));
      return item ? item.score : undefined;
    }
    function readTextFile(file) {
      return new Promise((resolve, reject) => {
        const reader = new FileReader();
        reader.onload = () => resolve(String(reader.result || ''));
        reader.onerror = () => reject(reader.error);
        reader.readAsText(file, 'utf-8');
      });
    }
    function readBase64File(file) {
      return new Promise((resolve, reject) => {
        const reader = new FileReader();
        reader.onload = () => {
          const bytes = new Uint8Array(reader.result);
          let binary = '';
          const chunk = 0x8000;
          for (let i = 0; i < bytes.length; i += chunk) {
            binary += String.fromCharCode(...bytes.subarray(i, i + chunk));
          }
          resolve(btoa(binary));
        };
        reader.onerror = () => reject(reader.error);
        reader.readAsArrayBuffer(file);
      });
    }
    async function post(url, body) {
      const res = await fetch(url, {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify(body)});
      const data = await res.json();
      if (!res.ok || data.error) throw new Error(data.error || '请求失败');
      return data;
    }
    async function getJson(url) {
      const res = await fetch(url);
      const data = await res.json();
      if (!res.ok || data.error) throw new Error(data.error || '请求失败');
      return data;
    }
    function renderStatsModule(notify=false) {
      if (!analysis) {
        document.getElementById('dimensionPanel').innerHTML = '<div class="small">请先完成一次分析。</div>';
        document.getElementById('errorReasonPanel').innerHTML = '';
        if (notify) showFailure('刷新统计失败：请先完成分析。');
        return;
      }
      renderDimensionStats();
      renderErrorReasonStats();
      if (notify) showSuccess('统计刷新成功。');
    }
    function setAnalyzeBusy(busy) {
      const button = document.getElementById('analyzeButton');
      if (button) {
        button.disabled = busy || !modelResultReady() || !document.getElementById('excelFile').files[0];
        button.textContent = busy ? '分析中' : (modelResultSource === 'api' ? '分析接口结果' : '分析');
      }
    }
    function setAnalysisProgress(percent, message, visible=true) {
      const wrap = document.getElementById('analysisProgress');
      const bar = document.getElementById('analysisProgressBar');
      const textNode = document.getElementById('analysisProgressText');
      const safePercent = Math.max(0, Math.min(100, Number(percent) || 0));
      if (wrap) wrap.classList.toggle('hide', !visible);
      if (bar) bar.style.width = `${safePercent}%`;
      if (textNode) textNode.textContent = `${Math.round(safePercent)}% · ${message || '正在分析...'}`;
      document.getElementById('status').textContent = message || '正在分析...';
    }
    function finishAnalysis(result) {
      analysis = result;
      setAnalysisVisibility(true);
      candidateById = {};
      assignments = {};
      materialGroups = [];
      materialSelections = {};
      verdictOverrides = {};
      categoryOverrides = {};
      errorReasonOverrides = {};
      goldDirtyRows = new Set();
      completedMaterials = {};
      confirmedRows = {};
      materialListPage = 1;
      selectedSourceMaterial = "all";
      for (const c of analysis.candidates) candidateById[c.id] = c;
      buildMaterialGroups();
      updateGoldEditStatus();
      document.getElementById('candidateList').innerHTML = analysis.candidates.map(c =>
        `<option value="${c.id}">#${c.id} ${c.materialName} / ${c.ruleName}</option>`
      ).join('');
      renderCards();
      renderFormulaPanel();
      renderDimensionStats();
      renderErrorReasonStats();
      renderMaterialList();
      renderRows();
      setModule('workbench');
      setAnalysisProgress(100, `分析完成：${analysis.summary.reviewRows} 行需要复核`, true);
      showSuccess(`分析成功：已生成 ${analysis.candidates.length} 条 JSON 规则匹配结果。`);
      window.setTimeout(() => setAnalysisProgress(100, '', false), 1200);
    }
    async function waitForAnalysisJob(jobId) {
      activeAnalysisJobId = jobId;
      while (activeAnalysisJobId === jobId) {
        const job = await getJson(`/api/analyze/status?id=${encodeURIComponent(jobId)}`);
        setAnalysisProgress(job.percent, job.message, true);
        if (job.status === 'done') {
          finishAnalysis(job.result);
          return;
        }
        if (job.status === 'error') {
          throw new Error(job.error || '分析失败');
        }
        await new Promise(resolve => window.setTimeout(resolve, 400));
      }
    }
    async function fetchModelResult() {
      const baseUrl = document.getElementById('apiBaseUrl')?.value || '';
      const caseId = document.getElementById('apiCaseId')?.value || '';
      document.getElementById('status').textContent = '正在拉取接口模型结果...';
      try {
        const result = await post('/api/model-result/fetch', {baseUrl, caseId});
        fetchedModelResult = result;
        previewJsonResult(result);
        updateFileReadyState();
        showSuccess(`接口拉取成功：caseId=${result.caseId}，共 ${result.count} 条规则。`);
        document.getElementById('status').textContent = `接口结果已就绪：${result.count} 条规则`;
      } catch (err) {
        fetchedModelResult = null;
        updateFileReadyState();
        document.getElementById('status').textContent = err.message;
        showFailure(`接口拉取失败：${err.message}`);
      }
    }
    async function readBatchCaseIds() {
      const file = document.getElementById('batchCaseFile')?.files?.[0];
      const typed = document.getElementById('batchCaseIds')?.value || '';
      const fileText = file ? await readTextFile(file) : '';
      return [typed, fileText].filter(Boolean).join('\n');
    }
    function updateBatchModeUI() {
      const mapped = document.getElementById('batchExcelMode')?.value === 'mapped';
      document.getElementById('batchMappingWrap')?.classList.toggle('hide', !mapped);
      document.getElementById('batchExcelFilesWrap')?.classList.toggle('hide', !mapped);
      document.getElementById('singleExcelWrap')?.classList.toggle('hide', modelResultSource === 'batch' && mapped);
      updateFileReadyState();
    }
    function currentDatasetLabel() {
      const excelFile = document.getElementById('excelFile')?.files?.[0];
      if (modelResultSource === 'api') {
        const apiName = fetchedModelResult ? `接口 caseId=${fetchedModelResult.caseId}（${fetchedModelResult.count || 0} 条规则）` : '接口 JSON 未拉取';
        return `${apiName} / Excel：${excelFile ? excelFile.name : '未上传'}`;
      }
      if (modelResultSource === 'batch') {
        return '批量测试结果在“数据获取”模块中查看；单案复核请点批量表里的“查看明细”。';
      }
      const jsonFile = document.getElementById('jsonFile')?.files?.[0];
      return `手工上传 JSON：${jsonFile ? jsonFile.name : '未上传'} / Excel：${excelFile ? excelFile.name : '未上传'}`;
    }
    function refreshWorkbenchDatasetInfo() {
      const target = document.getElementById('currentDatasetInfo');
      if (target) target.textContent = currentDatasetLabel();
    }
    function useCurrentDataset() {
      if (modelResultSource === 'batch') {
        showToast('批量测试不直接进入复核工作台；请在批量结果表中选择某个 case 查看明细。', 'info', 3600);
        return;
      }
      if (!modelResultReady() || !document.getElementById('excelFile')?.files?.[0]) {
        updateFileReadyState();
        showFailure('请先准备模型 JSON/TXT，并上传金标准 Excel。');
        return;
      }
      refreshWorkbenchDatasetInfo();
      setModule('workbench');
      showSuccess('数据已送入复核工作台，可以开始分析。');
    }
    function previewJsonResult(result) {
      const wrap = document.getElementById('jsonPreview');
      const meta = document.getElementById('jsonPreviewMeta');
      const body = document.getElementById('jsonPreviewBody');
      if (!wrap || !meta || !body) return;
      wrap.classList.remove('hide');
      meta.textContent = `caseId=${result.caseId || ''}；规则数量=${result.count || 0}；URL=${result.url || ''}`;
      body.textContent = JSON.stringify(result.content || [], null, 2);
    }
    async function previewUploadedJson() {
      clearFetchedModelResult();
      const file = document.getElementById('jsonFile')?.files?.[0];
      const wrap = document.getElementById('jsonPreview');
      const meta = document.getElementById('jsonPreviewMeta');
      const body = document.getElementById('jsonPreviewBody');
      if (!file || !wrap || !meta || !body) {
        updateFileReadyState();
        return;
      }
      const content = await readTextFile(file);
      let display = content;
      let countText = '未解析';
      try {
        const parsed = JSON.parse(content);
        const items = Array.isArray(parsed) ? parsed : (Array.isArray(parsed?.data) ? parsed.data : null);
        countText = items ? `${items.length} 条` : 'JSON对象';
        display = JSON.stringify(parsed, null, 2);
      } catch (_) {
        countText = '文本文件';
      }
      wrap.classList.remove('hide');
      meta.textContent = `上传文件=${file.name}；内容=${countText}`;
      body.textContent = display;
      updateFileReadyState();
    }
    async function copyPreviewJson() {
      const content = document.getElementById('jsonPreviewBody')?.textContent || '';
      if (!content) { showFailure('没有可复制的 JSON。'); return; }
      await navigator.clipboard.writeText(content);
      showSuccess('JSON 已复制。');
    }
    function downloadPreviewJson() {
      const content = document.getElementById('jsonPreviewBody')?.textContent || '';
      if (!content) { showFailure('没有可下载的 JSON。'); return; }
      const name = fetchedModelResult?.caseId ? `接口结果-${fetchedModelResult.caseId}.json` : '接口结果.json';
      const blob = new Blob([content], {type:'application/json;charset=utf-8'});
      const link = document.createElement('a');
      link.href = URL.createObjectURL(blob);
      link.download = name;
      link.click();
      URL.revokeObjectURL(link.href);
    }
    async function readBatchMapping() {
      const file = document.getElementById('batchMappingFile')?.files?.[0];
      return file ? await readTextFile(file) : '';
    }
    async function readBatchExcelFiles() {
      const files = [...(document.getElementById('batchExcelFiles')?.files || [])];
      const result = [];
      for (const file of files) {
        result.push({name: file.name, base64: await readBase64File(file)});
      }
      return result;
    }
    function renderBatchResult(job) {
      document.getElementById('batchResult')?.classList.remove('hide');
      const summary = job.summary || {};
      const statusText = `${job.message || ''}（${job.completed || 0}/${job.total || 0}，成功 ${job.success || 0}，失败 ${job.failed || 0}）`;
      document.getElementById('batchStatus').textContent = statusText;
      document.getElementById('batchExportButton').disabled = !job.cases || !job.cases.length;
      document.getElementById('batchSummary').innerHTML = [
        `总case <b>${summary.totalCases || 0}</b>`,
        `成功 <b>${summary.successCases || 0}</b>`,
        `失败 <b>${summary.failedCases || 0}</b>`,
        `JSON规则 <b>${summary.jsonRules || 0}</b>`,
        `Excel检测项 <b>${summary.excelRows || 0}</b>`,
        `漏检 <b>${summary.missedRows || 0}</b>`,
        `触发率 <b>${pct(summary.triggerRate || 0)}</b>`,
        `召回率 <b>${pct(summary.recall || 0)}</b>`,
        `准确率 <b>${pct(summary.accuracy || 0)}</b>`,
      ].join('');
      const rows = job.cases || [];
      document.getElementById('batchRows').innerHTML = rows.map(item => `
        <tr>
          <td>${escapeHtml(item.caseId)}</td>
          <td>${escapeHtml(item.excelFileName || '')}</td>
          <td>${escapeHtml(item.status || '')}<div class="small">${escapeHtml(item.message || '')}</div></td>
          <td>${item.jsonRules || 0}</td>
          <td>${item.excelRows || 0}</td>
          <td>${item.missedRows || 0}</td>
          <td>${pct(item.triggerRate || 0)}</td>
          <td>${pct(item.recall || 0)}</td>
          <td>${pct(item.accuracy || 0)}</td>
          <td>${item.status === 'success' ? `<button class="secondary mini" onclick="loadBatchCase('${jsString(item.caseId)}')">查看明细</button>` : '/'}</td>
        </tr>`).join('') || '<tr><td colspan="10" class="small">暂无批量结果。</td></tr>';
    }
    async function waitForBatchJob(jobId) {
      activeBatchJobId = jobId;
      window.lastBatchJobId = jobId;
      while (activeBatchJobId === jobId) {
        const job = await getJson(`/api/batch/status?id=${encodeURIComponent(jobId)}`);
        renderBatchResult(job);
        if (['done', 'done_with_errors', 'error'].includes(job.status)) {
          activeBatchJobId = '';
          showToast(job.message || '批量测试完成。', job.failed ? 'info' : 'success', 4200);
          return;
        }
        await new Promise(resolve => window.setTimeout(resolve, 900));
      }
    }
    async function startBatchTest() {
      const excelFile = document.getElementById('excelFile')?.files?.[0];
      const excelMode = document.getElementById('batchExcelMode')?.value || 'shared';
      if (excelMode === 'shared' && !excelFile) { showFailure('共用 Excel 模式请先上传 Excel 文件。'); return; }
      try {
        const payload = {
          baseUrl: document.getElementById('batchBaseUrl')?.value || defaults.modelResultApi,
          excelMode,
          caseIds: await readBatchCaseIds(),
          matchThreshold: Number(document.getElementById('matchThreshold').value),
          reviewThreshold: Number(document.getElementById('reviewThreshold').value),
        };
        if (excelMode === 'mapped') {
          payload.mappingContent = await readBatchMapping();
          payload.excelFiles = await readBatchExcelFiles();
        } else {
          payload.excelUpload = { name: excelFile.name, base64: await readBase64File(excelFile) };
        }
        document.getElementById('batchResult')?.classList.remove('hide');
        document.getElementById('batchStatus').textContent = '正在启动批量任务...';
        const started = await post('/api/batch/start', payload);
        await waitForBatchJob(started.jobId);
      } catch (err) {
        activeBatchJobId = '';
        document.getElementById('batchStatus').textContent = err.message;
        showFailure(`批量测试失败：${err.message}`);
      }
    }
    async function loadBatchCase(caseId) {
      if (!activeBatchJobId && !document.getElementById('batchResult')) return;
      try {
        const jobId = activeBatchJobId || window.lastBatchJobId;
        const result = await getJson(`/api/batch/case?id=${encodeURIComponent(jobId)}&caseId=${encodeURIComponent(caseId)}`);
        finishAnalysis(result);
        showSuccess(`已载入 caseId=${caseId} 的分析明细。`);
      } catch (err) {
        showFailure(`载入明细失败：${err.message}`);
      }
    }
    async function exportBatchResult() {
      const jobId = activeBatchJobId || window.lastBatchJobId;
      if (!jobId) { showFailure('没有可导出的批量任务。'); return; }
      try {
        const result = await post('/api/batch/export', {jobId});
        showSuccess(`批量结果已导出：${result.outputPath}`);
      } catch (err) {
        showFailure(`导出批量结果失败：${err.message}`);
      }
    }
    async function analyze() {
      setAnalyzeBusy(true);
      setAnalysisProgress(0, '正在准备分析...', true);
      showToast('正在生成匹配结果，请稍候。', 'info', 2400);
      try {
        const payload = {
          matchThreshold: Number(document.getElementById('matchThreshold').value),
          reviewThreshold: Number(document.getElementById('reviewThreshold').value)
        };
        const jsonFile = document.getElementById('jsonFile').files[0];
        const excelFile = document.getElementById('excelFile').files[0];
        if (!modelResultReady() || !excelFile) {
          updateFileReadyState();
          setAnalysisProgress(0, '', false);
          showFailure('分析失败：请准备模型结果并上传 Excel 文件。');
          return;
        }
        if (modelResultSource === 'api') {
          payload.jsonUpload = {
            name: `接口结果-${fetchedModelResult.caseId || 'case'}.json`,
            content: JSON.stringify(fetchedModelResult.content || [], null, 2)
          };
        } else if (jsonFile) {
          payload.jsonUpload = { name: jsonFile.name, content: await readTextFile(jsonFile) };
        }
        if (excelFile) payload.excelUpload = { name: excelFile.name, base64: await readBase64File(excelFile) };
        setAnalysisProgress(1, '上传文件已读取，正在启动分析任务...', true);
        const started = await post('/api/analyze', payload);
        await waitForAnalysisJob(started.jobId);
      } catch (err) {
        if (!analysis) setAnalysisVisibility(false);
        document.getElementById('status').textContent = err.message;
        setAnalysisProgress(100, err.message, false);
        showFailure(`分析失败：${err.message}`);
      } finally {
        activeAnalysisJobId = '';
        setAnalyzeBusy(false);
      }
    }
    function historyMetric(value) {
      return pct(Number(value || 0));
    }
    function formatReviewDuration(seconds) {
      if (seconds === null || seconds === undefined || seconds === '') return '未计算';
      let remaining = Math.max(0, Number(seconds) || 0);
      const days = Math.floor(remaining / 86400);
      remaining %= 86400;
      const hours = Math.floor(remaining / 3600);
      remaining %= 3600;
      const minutes = Math.floor(remaining / 60);
      const secs = Math.floor(remaining % 60);
      const parts = [];
      if (days) parts.push(`${days}天`);
      if (hours || days) parts.push(`${hours}小时`);
      if (minutes || hours || days) parts.push(`${minutes}分`);
      parts.push(`${secs}秒`);
      return parts.join('');
    }
    function updateReviewDurationPreview() {
      const start = document.getElementById('modelReviewStart')?.value;
      const end = document.getElementById('modelReviewEnd')?.value;
      const target = document.getElementById('modelReviewDuration');
      if (!target) return;
      if (!start && !end) {
        target.textContent = '阅卷时长：未计算';
        return;
      }
      if (!start || !end) {
        target.textContent = '阅卷时长：请补全时间';
        return;
      }
      const seconds = Math.floor((new Date(end).getTime() - new Date(start).getTime()) / 1000);
      target.textContent = seconds < 0 ? '阅卷时长：结束时间有误' : `阅卷时长：${formatReviewDuration(seconds)}`;
    }
    function selectedHistoryCount() {
      return historySelectedIds.size;
    }
    function updateHistorySelectionUI() {
      const button = document.getElementById('exportHistoryButton');
      const count = selectedHistoryCount();
      if (button) {
        button.disabled = count === 0;
        button.textContent = `导出已选（${count}）`;
      }
    }
    function toggleHistorySelection(id, checked) {
      if (checked) historySelectedIds.add(String(id));
      else historySelectedIds.delete(String(id));
      updateHistorySelectionUI();
      renderHistoryList();
    }
    function selectHistoryPage() {
      const totalPages = Math.max(1, Math.ceil(historyItems.length / HISTORY_PAGE_SIZE));
      historyPage = Math.min(Math.max(1, historyPage), totalPages);
      const start = (historyPage - 1) * HISTORY_PAGE_SIZE;
      historyItems.slice(start, start + HISTORY_PAGE_SIZE).forEach(item => historySelectedIds.add(String(item.id)));
      updateHistorySelectionUI();
      renderHistoryList();
    }
    function selectAllHistory() {
      historyItems.forEach(item => historySelectedIds.add(String(item.id)));
      updateHistorySelectionUI();
      renderHistoryList();
    }
    function clearHistorySelection() {
      historySelectedIds.clear();
      updateHistorySelectionUI();
      renderHistoryList();
    }
    function setHistoryPage(page) {
      const totalPages = Math.max(1, Math.ceil(historyItems.length / HISTORY_PAGE_SIZE));
      historyPage = Math.min(Math.max(1, Number(page) || 1), totalPages);
      renderHistoryList();
    }
    function renderHistoryList() {
      const list = document.getElementById('historyList');
      if (!historyItems.length) {
        list.innerHTML = '<div class="history-empty">还没有保存过比对记录。完成分析和人工复核后，可在这里保存当前比对。</div>';
        document.getElementById('historyPagination')?.classList.add('hide');
        updateHistorySelectionUI();
        return;
      }
      const totalPages = Math.max(1, Math.ceil(historyItems.length / HISTORY_PAGE_SIZE));
      historyPage = Math.min(Math.max(1, historyPage), totalPages);
      const pageItems = historyItems.slice((historyPage - 1) * HISTORY_PAGE_SIZE, historyPage * HISTORY_PAGE_SIZE);
      list.innerHTML = pageItems.map(item => `
        <article class="history-item">
          <div class="history-item-head">
            <div class="history-title-wrap">
              <input class="history-select" type="checkbox" ${historySelectedIds.has(String(item.id)) ? 'checked' : ''} aria-label="选择${escapeHtml(item.title)}" onchange="toggleHistorySelection('${item.id}', this.checked)">
              <div class="history-item-title">${escapeHtml(item.title)}</div>
            </div>
            <div class="history-item-time">${escapeHtml(item.createdAt)}</div>
          </div>
          <div class="history-files">JSON：${escapeHtml(item.jsonFileName)}<br>Excel：${escapeHtml(item.excelFileName)}</div>
          <div class="history-review-time">模型阅卷：${escapeHtml(item.modelReviewStartAt || '未填写')}${item.modelReviewEndAt ? ` 至 ${escapeHtml(item.modelReviewEndAt)}` : ''}；阅卷时长：${escapeHtml(item.modelReviewDurationText || '未填写')}</div>
          <div class="history-metrics">
            <span>召回率 <b>${historyMetric(item.recall)}</b></span>
            <span>触发率 <b>${historyMetric(item.triggerRate)}</b></span>
            <span>准确率 <b>${historyMetric(item.accuracy)}</b></span>
            <span>Excel行数 <b>${item.excelRows || 0}</b></span>
            <span>JSON规则 <b>${item.jsonRules || 0}</b></span>
            <span>待复核 <b>${item.reviewRows || 0}</b></span>
            <span>漏检 <b>${item.missedRows || 0}</b></span>
          </div>
          <div class="history-actions">
            <button class="secondary" id="history-toggle-${item.id}" onclick="toggleHistoryStats('${item.id}')">${historyDetailState(item.id)}</button>
            <button onclick="loadHistory('${item.id}')">载入复核</button>
            <button class="secondary" onclick="deleteHistory('${item.id}')">删除记录</button>
          </div>
          ${historyDetailState(item.id) === '收起统计' ? historyDetailHtml(activeHistoryRecord) : ''}
        </article>`).join('');
      const pagination = document.getElementById('historyPagination');
      if (pagination) {
        pagination.classList.toggle('hide', totalPages <= 1);
        pagination.innerHTML = `
          <button class="secondary" onclick="setHistoryPage(${historyPage - 1})" ${historyPage <= 1 ? 'disabled' : ''}>上一页</button>
          <span class="history-page-status">第 ${historyPage} / ${totalPages} 页，共 ${historyItems.length} 条</span>
          <button class="secondary" onclick="setHistoryPage(${historyPage + 1})" ${historyPage >= totalPages ? 'disabled' : ''}>下一页</button>`;
      }
      updateHistorySelectionUI();
    }
    function historyDetailState(id) {
      return activeHistoryRecord && activeHistoryRecord.id === id && !historyStatsCollapsed[id] ? '收起统计' : '展开统计';
    }
    function emptyDimensionStat() {
      return {goldTriggered:0,modelTriggered:0,correct:0,recall:0,accuracy:0,hitRate:0};
    }
    function dimensionStatsNote() {
      return `
        <div class="dimension-panel-note">
          统计范围：按每条检测项的“统计维度”归类，仅统计金标准是否触发=是的检测项。金标准触发数=该维度中金标准是否触发=是的检测项数；模型触发数=其中模型是否触发=是的检测项数；判断正确数=其中模型是否触发=是且AA=是的检测项数。<br>
          召回率=模型触发数/金标准触发数；准确率=判断正确数/模型触发数；综合命中率=判断正确数/金标准触发数。
        </div>`;
    }
    function dimensionColgroup() {
      return `
        <colgroup>
          <col style="width:10%">
          <col style="width:11%">
          <col style="width:11%">
          <col style="width:11%">
          <col style="width:9%">
          <col style="width:9%">
          <col style="width:10%">
          <col style="width:29%">
        </colgroup>`;
    }
    function dimensionStatsRow(category, item, metricFormatter) {
      const data = item || emptyDimensionStat();
      return `<tr>
        <td><b>${escapeHtml(category)}</b></td>
        <td class="num">${data.goldTriggered || 0}</td>
        <td class="num">${data.modelTriggered || 0}</td>
        <td class="num">${data.correct || 0}</td>
        <td class="num rate">${metricFormatter(data.recall || 0)}</td>
        <td class="num rate">${metricFormatter(data.accuracy || 0)}</td>
        <td class="num rate">${metricFormatter(data.hitRate || 0)}</td>
        <td class="formula small">召回率=模型触发数/金标准触发数；准确率=判断正确数/模型触发数；综合命中率=判断正确数/金标准触发数</td>
      </tr>`;
    }
    async function loadHistoryList(notify=false) {
      try {
        historyItems = await getJson('/api/history');
        const validIds = new Set(historyItems.map(item => String(item.id)));
        historySelectedIds = new Set([...historySelectedIds].filter(id => validIds.has(id)));
        if (activeHistoryRecord && !validIds.has(String(activeHistoryRecord.id))) activeHistoryRecord = null;
        const totalPages = Math.max(1, Math.ceil(historyItems.length / HISTORY_PAGE_SIZE));
        historyPage = Math.min(Math.max(1, historyPage), totalPages);
        renderHistoryList();
        document.getElementById('historyStatus').textContent = `共 ${historyItems.length} 条记录；已选择 ${selectedHistoryCount()} 条`;
        if (notify) showSuccess(`历史记录刷新成功，共 ${historyItems.length} 条。`);
      } catch (err) {
        document.getElementById('historyStatus').textContent = err.message;
        if (notify) showFailure(`刷新历史失败：${err.message}`);
      }
    }
    function historyDetailHtml(record) {
      if (!record) return '';
      const summary = record.summary || {};
      const dimensionStats = summary.dimensionStats || {};
      const categories = ['签名', '日期', '捺印', '印章', '综合', '其他'];
      const errorReasonDetails = (record.errorReasonDetails || []).filter(item => String(item['错误原因'] || '').trim());
      const reasonCount = new Set(errorReasonDetails.map(item => item['错误原因'])).size;
      return `
        <div class="history-detail">
        <div class="history-detail-title">历史统计：${escapeHtml(record.title || '未命名比对')} <span class="small">${escapeHtml(record.createdAt || '')}</span></div>
        <section class="dimension-panel">
          ${dimensionStatsNote()}
          <table class="dimension-table">
            ${dimensionColgroup()}
            <thead><tr><th>统计维度</th><th class="num">金标准触发数</th><th class="num">模型触发数</th><th class="num">判断正确数</th><th class="num">召回率</th><th class="num">准确率</th><th class="num">综合命中率</th><th>计算公式</th></tr></thead>
            <tbody>${categories.map(category => { const item = dimensionStats[category] || emptyDimensionStat(); return dimensionStatsRow(category, item, historyMetric); }).join('')}</tbody>
          </table>
        </section>
        <section class="error-reason-panel">
          <div class="error-reason-panel-title">错误原因汇总 <span class="dimension-panel-note">共 ${reasonCount} 类；相同错误原因已合并</span></div>
          <div class="error-reason-table-wrap">
            <table class="error-reason-table">
              <thead><tr>${errorReasonHeaders.map(header => `<th>${header}</th>`).join('')}</tr></thead>
              <tbody>${errorReasonDetails.length ? errorReasonRowsHtml(errorReasonDetails) : `<tr><td colspan="${errorReasonHeaders.length}" class="small">暂无已填写的错误原因。</td></tr>`}</tbody>
            </table>
          </div>
        </section>
        </div>`;
    }
    function renderHistoryDetail(record) {
      activeHistoryRecord = record || null;
      const detail = document.getElementById('historyDetail');
      if (detail) {
        detail.dataset.recordId = record && record.id ? record.id : '';
        detail.classList.add('hide');
        detail.innerHTML = '';
      }
    }
    function toggleHistoryStats(id) {
      if (!activeHistoryRecord || activeHistoryRecord.id !== id) {
        viewHistory(id);
        return;
      }
      historyStatsCollapsed[id] = !historyStatsCollapsed[id];
      renderHistoryList();
    }
    async function viewHistory(id) {
      try {
        const record = await getJson(`/api/history/detail?id=${encodeURIComponent(id)}`);
        historyStatsCollapsed[id] = false;
        renderHistoryDetail(record);
        renderHistoryList();
      } catch (err) {
        document.getElementById('historyStatus').textContent = err.message;
      }
    }
    function openHistorySaveDialog() {
      if (!analysis) {
        showFailure('保存失败：请先完成一次分析。');
        return;
      }
      document.getElementById('historySaveOverlay')?.classList.remove('hide');
      updateReviewDurationPreview();
      window.setTimeout(() => document.getElementById('historyName')?.focus(), 0);
    }
    function closeHistorySaveDialog() {
      document.getElementById('historySaveOverlay')?.classList.add('hide');
    }
    function historySaveOverlayClick(event) {
      if (event.target.id === 'historySaveOverlay') closeHistorySaveDialog();
    }
    async function saveCurrentHistory() {
      if (!analysis) {
        document.getElementById('historyStatus').textContent = '请先完成一次分析。';
        showFailure('保存失败：请先完成一次分析。');
        return;
      }
      if (goldDirtyRows.size) {
        showFailure(`保存失败：还有 ${goldDirtyRows.size} 条金标准修改未保存，请先点击“保存金标准修改”。`);
        return;
      }
      const modelReviewStart = document.getElementById('modelReviewStart').value;
      const modelReviewEnd = document.getElementById('modelReviewEnd').value;
      if (!modelReviewStart || !modelReviewEnd) {
        showFailure('保存失败：请填写模型阅卷开始时间和结束时间。');
        return;
      }
      if (new Date(modelReviewEnd).getTime() < new Date(modelReviewStart).getTime()) {
        showFailure('保存失败：模型阅卷结束时间不能早于开始时间。');
        return;
      }
      const status = document.getElementById('historyStatus');
      status.textContent = '正在保存...';
      try {
        const result = await post('/api/history/save', {
          title: document.getElementById('historyName').value.trim(),
          analysis,
          assignments,
          verdictOverrides,
          categoryOverrides,
          errorReasonOverrides,
          completedMaterials,
          confirmedRows,
          summarySnapshot: currentSummarySnapshot(),
          modelReviewStart,
          modelReviewEnd
        });
        document.getElementById('historyName').value = '';
        document.getElementById('modelReviewStart').value = '';
        document.getElementById('modelReviewEnd').value = '';
        updateReviewDurationPreview();
        historyPage = 1;
        closeHistorySaveDialog();
        status.textContent = `已保存：${result.title}`;
        showSuccess(`保存成功：${result.title}`);
        await loadHistoryList();
      } catch (err) {
        status.textContent = err.message;
        showFailure(`保存失败：${err.message}`);
      }
    }
    async function exportSelectedHistories() {
      const ids = [...historySelectedIds];
      if (!ids.length) {
        showFailure('导出失败：请至少选择一条比对历史。');
        return;
      }
      try {
        showToast(`正在导出 ${ids.length} 条比对历史...`, 'info');
        const result = await post('/api/history/export', {ids});
        document.getElementById('historyStatus').textContent = `已导出 ${result.count} 条历史：${result.outputPath}`;
        showSuccess(`导出成功：已导出 ${result.count} 条比对历史。`);
      } catch (err) {
        document.getElementById('historyStatus').textContent = err.message;
        showFailure(`导出失败：${err.message}`);
      }
    }
    async function deleteHistory(id) {
      if (!confirm('确定删除这条历史比对记录吗？删除后无法恢复。')) return;
      try {
        await post('/api/history/delete', {id});
        historySelectedIds.delete(String(id));
        if (activeHistoryRecord && activeHistoryRecord.id === id) activeHistoryRecord = null;
        updateHistorySelectionUI();
        document.getElementById('historyDetail').innerHTML = '';
        showSuccess('历史记录删除成功。');
        await loadHistoryList();
      } catch (err) {
        document.getElementById('historyStatus').textContent = err.message;
        showFailure(`删除历史失败：${err.message}`);
      }
    }
    function hydrateAnalysisState(sourceAnalysis) {
      analysis = sourceAnalysis;
      candidateById = {};
      materialGroups = [];
      for (const c of analysis.candidates) candidateById[c.id] = c;
      buildMaterialGroups();
      document.getElementById('candidateList').innerHTML = analysis.candidates.map(c => `<option value="${c.id}">#${c.id} ${escapeHtml(c.materialName)} / ${escapeHtml(c.ruleName)}</option>`).join('');
      document.getElementById('matchThreshold').value = analysis.matchThreshold || 0.67;
      document.getElementById('reviewThreshold').value = analysis.reviewThreshold || 0.82;
    }
    async function loadHistory(id) {
      try {
        const record = await getJson(`/api/history/detail?id=${encodeURIComponent(id)}`);
        hydrateAnalysisState(record.analysis);
        assignments = record.assignments || {};
        verdictOverrides = record.verdictOverrides || {};
        categoryOverrides = record.categoryOverrides || {};
        errorReasonOverrides = record.errorReasonOverrides || {};
        goldDirtyRows = new Set();
        completedMaterials = record.completedMaterials || {};
        confirmedRows = record.confirmedRows || {};
        materialListPage = 1;
        materialSelections = {};
        selectedSourceMaterial = 'all';
        setAnalysisVisibility(true);
        updateGoldEditStatus();
        renderCards();
        renderFormulaPanel();
        renderStatsModule();
        renderMaterialList();
        renderRows();
        document.getElementById('status').textContent = `已载入历史：${record.title}`;
        setModule('workbench');
        showSuccess(`历史记录载入成功：${record.title}`);
      } catch (err) {
        document.getElementById('historyStatus').textContent = err.message;
        showFailure(`载入历史失败：${err.message}`);
      }
    }
    function renderCards() {
      const s = currentSummary();
      const cards = [
        { k: 'Excel行数', v: s.excelRows, f: 'Excel 数据行总数' },
        { k: 'JSON规则', v: s.jsonRules, f: 'JSON 展开后的候选规则数' },
        { k: '漏检项', v: s.missedRows, f: `金标准要求触发但模型未触发的检测项 = <code>${s.missedRows}</code> 条` },
        { k: '触发率', v: pct(s.triggerRate), f: `模型触发的金标准检测项 / 金标准要求触发的检测项 = <code>${s.expectedTriggered}/${s.expectedTotal}</code>` },
        { k: '召回率', v: pct(s.recall), f: `错误场景中模型触发的检测项 / 错误场景 = <code>${s.errorRecalled}/${s.errorTotal}</code>` },
        { k: '准确率', v: pct(s.accuracy), f: `模型判断正确的已触发检测项 / 模型已触发的检测项 = <code>${s.modelCorrect}/${s.modelTriggered}</code>` },
      ];
      document.getElementById('cards').innerHTML = cards.map(card => `<div class="card"><div class="k">${card.k}</div><div class="v">${card.v}</div><div class="formula">${card.f}</div></div>`).join('');
    }
    function rowCandidate(row) {
      const key = rowKey(row);
      const assigned = assignments[key] !== undefined ? assignments[key] : row.assignedId;
      return assigned === 'none' || assigned === null || assigned === undefined ? null : candidateById[assigned];
    }
    function rowErrorReason(row) {
      const key = rowKey(row);
      return errorReasonOverrides[key] !== undefined ? errorReasonOverrides[key] : (row.errorReason || '');
    }
    function updateErrorReason(rowNum, value) {
      invalidateMaterialCompletion(rowNum);
      errorReasonOverrides[rowNum] = value;
      renderErrorReasonStats();
    }
    function rowKey(row) {
      return text(row && (row.rowKey !== undefined ? row.rowKey : row.excelRow));
    }
    function rowArg(row) {
      return `'${rowKey(row).replace(/\\/g, '\\\\').replace(/'/g, "\\'").replace(/\n/g, '\\n')}'`;
    }
    function sourceRowNumber(row) {
      return row && row.sourceRowLabel !== undefined
        ? row.sourceRowLabel
        : row && row.sourceExcelRow !== undefined
          ? row.sourceExcelRow
          : row.excelRow;
    }
    function rowAssignedId(row) {
      const key = rowKey(row);
      const assigned = assignments[key] !== undefined ? assignments[key] : row.assignedId;
      return assigned === 'none' || assigned === null || assigned === undefined ? null : Number(assigned);
    }
    const categoryNames = ['签名', '日期', '捺印', '印章', '其他'];
    function rowCategory(row) {
      const key = rowKey(row);
      return categoryNames.includes(categoryOverrides[key]) ? categoryOverrides[key] : (row.category || '其他');
    }
    function rowModelIdentifiesError(row) {
      const candidate = rowCandidate(row);
      return Boolean(candidate && candidate.isCorrect !== true);
    }
    function currentDimensionStats() {
      const rows = analysis ? analysis.rows : [];
      const stats = {};
      for (const category of categoryNames) {
        const goldRows = rows.filter(row => rowCategory(row) === category && row.goldTrigger === '是' && !goldRowPending(row));
        const modelTriggered = goldRows.filter(rowModelTriggered).length;
        const correct = goldRows.filter(row => rowModelTriggered(row) && computedAa(row, rowCandidate(row)) === '是').length;
        stats[category] = {
          goldTriggered: goldRows.length,
          modelTriggered,
          correct,
          recall: goldRows.length ? modelTriggered / goldRows.length : 0,
          accuracy: modelTriggered ? correct / modelTriggered : 0,
          hitRate: goldRows.length ? correct / goldRows.length : 0,
        };
      }
      const combinedGold = ['签名', '日期', '捺印', '印章'].reduce((sum, category) => sum + stats[category].goldTriggered, 0);
      const combinedModel = ['签名', '日期', '捺印', '印章'].reduce((sum, category) => sum + stats[category].modelTriggered, 0);
      const combinedCorrect = ['签名', '日期', '捺印', '印章'].reduce((sum, category) => sum + stats[category].correct, 0);
      stats['综合'] = {
        goldTriggered: combinedGold,
        modelTriggered: combinedModel,
        correct: combinedCorrect,
        recall: combinedGold ? combinedModel / combinedGold : 0,
        accuracy: combinedModel ? combinedCorrect / combinedModel : 0,
        hitRate: combinedGold ? combinedCorrect / combinedGold : 0,
      };
      return stats;
    }
    function renderDimensionStats() {
      if (!analysis) return;
      const stats = currentDimensionStats();
      const categories = ['签名', '日期', '捺印', '印章', '综合', '其他'];
      document.getElementById('dimensionPanel').innerHTML = `
        <div class="dimension-panel-head">
          <div class="dimension-panel-title">统计维度汇总</div>
          ${dimensionStatsNote()}
        </div>
        <table class="dimension-table">
          ${dimensionColgroup()}
          <thead><tr><th>统计维度</th><th class="num">金标准触发数</th><th class="num">模型触发数</th><th class="num">判断正确数</th><th class="num">召回率</th><th class="num">准确率</th><th class="num">综合命中率</th><th>计算公式</th></tr></thead>
          <tbody>${categories.map(category => {
            const item = stats[category];
            return dimensionStatsRow(category, item, pct);
          }).join('')}</tbody>
        </table>`;
    }
    function currentErrorReasonDetails() {
      if (!analysis) return [];
      const details = [];
      for (const row of analysis.rows) {
        const candidate = rowCandidate(row);
        const isUnmatchedError = !candidate && (row.status === 'unmatched' || row.status === 'confirmed_none' || row.goldTrigger === '是');
        const isModelError = candidate && candidate.isCorrect === false;
        const reason = rowErrorReason(row).trim();
        if ((!isUnmatchedError && !isModelError) || !reason) continue;
        details.push({
          '错误原因': reason,
          '来源文件': analysis.excelFileName || '',
          '材料类型': row.materialType || '',
          '对应材料名称': row.materialName || '',
          '所属阶段': row.stage || '',
          '阶段时间顺序': row.stageTimeOrder || '',
          '规则名称': row.ruleName || '',
          '特殊条件': row.specialCondition || '',
          '判断规则': row.judgementRule || '',
          '错误场景': row.errorScene || '',
          '是否触发': row.goldTrigger || '',
          '是否正确': row.goldRaw || '',
          '模型是否触发': rowModelTriggered(row) ? '是' : '否',
          '模型是否正确': computedAa(row, candidate),
          '具体问题内容': (candidate && candidate.questionDesc) || '未匹配到对应的 JSON 检测项'
        });
      }
      const grouped = new Map();
      for (const item of details) {
        if (!grouped.has(item['错误原因'])) grouped.set(item['错误原因'], []);
        grouped.get(item['错误原因']).push(item);
      }
      return Array.from(grouped.values()).flat();
    }
    function errorReasonRowsHtml(items) {
      const body = [];
      for (let index = 0; index < items.length;) {
        let end = index + 1;
        while (end < items.length && items[end]['错误原因'] === items[index]['错误原因']) end += 1;
        const first = items[index];
        const cells = errorReasonHeaders.slice(1).map(header => `<td>${escapeHtml(first[header] ?? '')}</td>`).join('');
        body.push(`<tr><td rowspan="${end - index}">${escapeHtml(first['错误原因'])}</td>${cells}</tr>`);
        for (let rowIndex = index + 1; rowIndex < end; rowIndex += 1) {
          body.push(`<tr>${errorReasonHeaders.slice(1).map(header => `<td>${escapeHtml(items[rowIndex][header] ?? '')}</td>`).join('')}</tr>`);
        }
        index = end;
      }
      return body.join('');
    }
    function renderErrorReasonStats() {
      if (!analysis) return;
      const items = currentErrorReasonDetails();
      const headers = errorReasonHeaders;
      const body = errorReasonRowsHtml(items);
      document.getElementById('errorReasonPanel').innerHTML = `
        <div class="error-reason-panel-title">
          <span>错误原因汇总 <span class="dimension-panel-note">相同错误原因已合并，只显示已填写内容</span></span>
          <button class="secondary" onclick="exportErrorReasons()" ${items.length ? '' : 'disabled'}>导出</button>
        </div>
        <div class="error-reason-table-wrap">
          <table class="error-reason-table">
            <thead><tr>${headers.map(header => `<th>${header}</th>`).join('')}</tr></thead>
            <tbody>${body || `<tr><td colspan="${headers.length}" class="small">暂无已填写的错误原因。</td></tr>`}</tbody>
          </table>
        </div>`;
    }
    function isMissed(row) {
      return row.goldTrigger === '是' && !rowModelTriggered(row);
    }
    function rowModelTriggered(row) {
      const candidate = rowCandidate(row);
      return Boolean(candidate && candidate.modelTriggered !== false);
    }
    function currentSummary() {
      const rows = analysis ? analysis.rows : [];
      const eligibleRows = rows.filter(r => !goldRowPending(r));
      const expectedTotal = eligibleRows.filter(r => r.goldTrigger === '是').length;
      const expectedTriggered = eligibleRows.filter(r => r.goldTrigger === '是' && rowModelTriggered(r)).length;
      const modelTriggered = eligibleRows.filter(r => rowModelTriggered(r)).length;
      const modelCorrect = eligibleRows.filter(r => rowModelTriggered(r) && computedAa(r, rowCandidate(r)) === '是').length;
      const modelError = eligibleRows.filter(r => rowModelTriggered(r) && computedAa(r, rowCandidate(r)) === '否').length;
      const correctRecalled = eligibleRows.filter(r => r.goldTrigger === '是' && rowModelTriggered(r) && computedAa(r, rowCandidate(r)) === '是').length;
      const errorTotal = eligibleRows.filter(r => r.goldTrigger === '是' && r.goldCorrect !== '是').length;
      const errorRecalled = eligibleRows.filter(r => r.goldTrigger === '是' && r.goldCorrect !== '是' && rowModelTriggered(r)).length;
      const missedRows = eligibleRows.filter(isMissed).length;
      return {
        excelRows: rows.length,
        jsonRules: analysis ? analysis.candidates.length : 0,
        expectedTotal,
        expectedTriggered,
        modelTriggered,
        modelCorrect,
        modelError,
        correctRecalled,
        errorTotal,
        errorRecalled,
        missedRows,
        recall: errorTotal ? errorRecalled / errorTotal : 0,
        triggerRate: expectedTotal ? expectedTriggered / expectedTotal : 0,
        accuracy: modelTriggered ? modelCorrect / modelTriggered : 0,
      };
    }
    function currentSummarySnapshot() {
      return {
        ...currentSummary(),
        dimensionStats: currentDimensionStats(),
      };
    }
    function goldRowPending(row) {
      return Boolean(row && row.materialCopyPending && (
        !['是', '否'].includes(row.goldTrigger) || !['是', '否'].includes(row.goldCorrect)
      ));
    }
    function displayMaterialType(row) {
      const value = text(row && row.materialType);
      return goldRowPending(row) ? value : value.replace('·待确认', '');
    }
    function renderFormulaPanel() {
      const statsInfo = document.getElementById('statsInfo');
      const scoreInfo = document.getElementById('scoreInfo');
      if (!statsInfo || !scoreInfo) return;
      if (!analysis) return;
      const matchThreshold = Number(document.getElementById('matchThreshold').value).toFixed(2);
      const reviewThreshold = Number(document.getElementById('reviewThreshold').value).toFixed(2);
      statsInfo.innerHTML = `
        <div class="formula-title">统计口径</div>
        <div class="formula-line">召回率 = <b>错误场景中模型触发的检测项 / 错误场景总数</b>。</div>
        <div class="formula-line">触发率 = <b>模型触发的金标准检测项 / 金标准要求触发的检测项</b>。</div>
        <div class="formula-line">准确率 = <b>模型判断正确的已触发检测项 / 模型已触发的检测项</b>。</div>
        <div class="formula-line">漏检项 = <b>金标准要求触发但模型未触发的检测项</b>。</div>
        <div class="formula-line">错误场景指金标准要求触发、且金标准判定不正确的检测项；材料类型检测按外层 <code>isExit</code> 判断。</div>`;
      scoreInfo.innerHTML = `
        <div class="formula-title">匹配分数</div>
        <div class="formula-line">分数只表示“Excel行”和“JSON规则”的文本相似度，不代表模型回答对错。</div>
        <div class="formula-line"><code>分数 = 0.36 × 材料相似度 + 0.64 × 规则名称相似度 + 少量完全匹配加分</code></div>
        <div class="formula-line">分数 ≥ <code>${matchThreshold}</code> 才自动匹配；低于复核阈值 <code>${reviewThreshold}</code> 会放入待复核。</div>`;
    }
    function rowStatusMatchesFilter(row) {
      const filter = document.getElementById('filter').value;
      if (filter === 'completed') return materialIsCompleted(rowSourceKey(row));
      if (filter === 'incomplete') return !materialIsCompleted(rowSourceKey(row));
      const status = effectiveStatus(row);
      if (filter === 'review' && !['review','unmatched','duplicate_dropped'].includes(status)) return false;
      if (filter === 'missed' && !isMissed(row)) return false;
      if (filter === 'unmatched' && status !== 'unmatched') return false;
      if (filter === 'auto' && status !== 'auto') return false;
      if (filter === 'confirmed' && status !== 'confirmed') return false;
      return true;
    }
    function rowSearchMatches(row) {
      const query = document.getElementById('search').value.trim();
      return !query || (row.materialType + row.materialName + row.ruleName + row.goldRaw).includes(query);
    }
    function rowVisible(row) {
      if (selectedSourceMaterial !== 'all' && rowSourceKey(row) !== selectedSourceMaterial) return false;
      return rowStatusMatchesFilter(row) && rowSearchMatches(row);
    }
    function rowSourceKey(row) {
      return `${row.materialType}||${row.materialName}`;
    }
    function rowSourceParts(row) {
      return {
        type: displayMaterialType(row) || '未填材料类型',
        doc: row.materialName || '未填对应材料名',
      };
    }
    function rowSourceLabel(row) {
      const parts = rowSourceParts(row);
      return `${parts.type} / ${parts.doc}`;
    }
    function effectiveStatus(row) {
      return rowIsConfirmed(row) ? 'confirmed' : row.status;
    }
    function statusLabel(status) {
      return ({
        auto: '自动通过',
        review: '待复核',
        unmatched: '未匹配',
        duplicate_dropped: '重复候选',
        confirmed: '已确认',
        confirmed_none: '已确认（无匹配）'
      })[status] || status;
    }
    function rowIsConfirmed(row) {
      const key = rowKey(row);
      return Object.prototype.hasOwnProperty.call(verdictOverrides, key);
    }
    function goldCorrectFromRaw(value) {
      const raw = text(value).trim();
      if (!raw || ['/', '无', 'none', 'None'].includes(raw)) return '';
      return raw === '是' ? '是' : '否';
    }
    function markGoldDirty(rowNum) {
      const row = getRow(rowNum);
      if (!row) return;
      goldDirtyRows.add(String(rowNum));
      delete confirmedRows[rowNum];
      delete verdictOverrides[rowNum];
      invalidateMaterialCompletion(rowNum);
      updateGoldEditStatus();
    }
    function updateGoldEditStatus() {
      const status = document.getElementById('goldEditStatus');
      const button = document.getElementById('saveGoldButton');
      const count = goldDirtyRows.size;
      if (status) status.textContent = count ? `有 ${count} 条检测项的金标准尚未保存；保存后相关材料需重新确认。` : '金标准未修改';
      if (button) button.disabled = !analysis || count === 0;
    }
    function setGoldTrigger(rowNum, value) {
      const row = getRow(rowNum);
      if (!row || !['', '是', '否'].includes(value)) return;
      row.goldTrigger = value;
      row.goldCorrect = goldCorrectFromRaw(row.goldRaw);
      markGoldDirty(rowNum);
      renderRows();
    }
    function setGoldCorrectChoice(rowNum, value) {
      const row = getRow(rowNum);
      if (!row || !['', '是', '否'].includes(value)) return;
      row.goldCorrect = value;
      if (value === '是') row.goldRaw = '是';
      else if (value === '') row.goldRaw = '';
      else if (!text(row.goldRaw).trim() || text(row.goldRaw).trim() === '是') row.goldRaw = '否';
      markGoldDirty(rowNum);
      renderRows();
    }
    function updateGoldDescription(rowNum, value) {
      const row = getRow(rowNum);
      if (!row) return;
      row.goldRaw = text(value);
      row.goldCorrect = goldCorrectFromRaw(row.goldRaw);
      markGoldDirty(rowNum);
    }
    function restoreGoldStandard(rowNum) {
      const row = getRow(rowNum);
      if (!row) return;
      row.goldTrigger = row.goldTriggerOriginal !== undefined ? row.goldTriggerOriginal : row.goldTrigger;
      row.goldRaw = row.goldRawOriginal !== undefined ? row.goldRawOriginal : row.goldRaw;
      row.goldCorrect = row.goldCorrectOriginal !== undefined ? row.goldCorrectOriginal : goldCorrectFromRaw(row.goldRaw);
      goldDirtyRows.delete(String(rowNum));
      delete confirmedRows[rowNum];
      delete verdictOverrides[rowNum];
      invalidateMaterialCompletion(rowNum);
      updateGoldEditStatus();
      renderRows();
    }
    function saveGoldStandards() {
      if (!analysis) {
        showFailure('保存失败：请先完成分析。');
        return;
      }
      if (!goldDirtyRows.size) {
        showFailure('保存失败：当前没有未保存的金标准修改。');
        return;
      }
      const count = goldDirtyRows.size;
      goldDirtyRows.clear();
      updateGoldEditStatus();
      renderRows();
      showSuccess(`保存成功：已保存 ${count} 条金标准修改，相关材料已恢复为未完成。`);
    }
    function computedAa(row, candidate) {
      const key = rowKey(row);
      if (verdictOverrides[key] !== undefined) return verdictOverrides[key];
      if (!row.goldTrigger || !row.goldCorrect) return '/';
      if (candidate && row.goldCorrect) return Boolean(candidate.isCorrect) === (row.goldCorrect === '是') ? '是' : '否';
      if (candidate && row.goldTrigger === '否') return '否';
      if (!candidate && row.goldTrigger === '是') return '否';
      return '/';
    }
    function aaLabel(value) {
      if (value === '是') return 'AA=是';
      if (value === '否') return 'AA=否';
      return 'AA=/';
    }
    function sourceMaterialStats() {
      const map = new Map();
      const ensure = (key, label, parts) => {
        if (!map.has(key)) map.set(key, {
          key,
          label,
          type: parts ? parts.type : label,
          doc: parts ? parts.doc : '',
          total:0,
          pending:0,
          unmatched:0,
          missed:0,
          confirmed:0
        });
        return map.get(key);
      };
      const all = ensure('all', '全部材料', {type:'全部材料', doc:'所有材料汇总'});
      for (const row of analysis.rows) {
        const key = rowSourceKey(row);
        const parts = rowSourceParts(row);
        const item = ensure(key, rowSourceLabel(row), parts);
        const status = effectiveStatus(row);
        for (const target of [all, item]) {
          target.total += 1;
          if (['review','unmatched','duplicate_dropped'].includes(status)) target.pending += 1;
          if (status === 'unmatched') target.unmatched += 1;
          if (isMissed(row)) target.missed += 1;
          if (status === 'confirmed') target.confirmed += 1;
        }
      }
      return Array.from(map.values());
    }
    function materialRows(materialKey) {
      if (!analysis) return [];
      return analysis.rows.filter(row => materialKey === 'all' || rowSourceKey(row) === materialKey);
    }
    function hasConfirmedAssignment(row) {
      return rowIsConfirmed(row);
    }
    function materialConfirmedCount(materialKey) {
      return materialRows(materialKey).filter(hasConfirmedAssignment).length;
    }
    function materialIsReady(materialKey) {
      const rows = materialRows(materialKey);
      return rows.length > 0 && rows.every(hasConfirmedAssignment);
    }
    function materialIsCompleted(materialKey) {
      return materialKey !== 'all' && completedMaterials[materialKey] === true && materialIsReady(materialKey);
    }
    function materialFileSummary() {
      const groups = sourceMaterialStats().filter(item => item.key !== 'all');
      const matchesSearch = group => analysis.rows.some(row => rowSourceKey(row) === group.key && rowSearchMatches(row));
      const matchesStatus = group => analysis.rows.some(row => rowSourceKey(row) === group.key && rowSearchMatches(row) && rowStatusMatchesFilter(row));
      const searchedGroups = groups.filter(matchesSearch);
      const visibleGroups = searchedGroups.filter(matchesStatus);
      const reviewedFiles = searchedGroups.filter(group => {
        return materialIsCompleted(group.key);
      }).length;
      return {visibleFiles: visibleGroups.length, reviewedFiles, totalFiles: searchedGroups.length};
    }
    function renderMaterialPagination(totalItems) {
      const pagination = document.getElementById('materialPagination');
      if (!pagination) return;
      if (!totalItems) {
        pagination.classList.add('hide');
        pagination.innerHTML = '';
        return;
      }
      const totalPages = Math.max(1, Math.ceil(totalItems / MATERIAL_PAGE_SIZE));
      materialListPage = Math.min(Math.max(Number(materialListPage) || 1, 1), totalPages);
      pagination.classList.remove('hide');
      pagination.innerHTML = `
        <button class="secondary" ${materialListPage <= 1 ? 'disabled' : ''} onclick="changeMaterialListPage(${materialListPage - 1})">上一页</button>
        <span class="material-page-status">第 ${materialListPage} / ${totalPages} 页，共 ${totalItems} 个文件</span>
        <button class="secondary" ${materialListPage >= totalPages ? 'disabled' : ''} onclick="changeMaterialListPage(${materialListPage + 1})">下一页</button>`;
    }
    function changeMaterialListPage(page) {
      materialListPage = Math.max(1, Number(page) || 1);
      renderMaterialList();
    }
    function materialCardHtml(item) {
      const active = item.key === selectedSourceMaterial ? 'active' : '';
      const isAll = item.key === 'all';
      const completed = materialIsCompleted(item.key);
      const readyCount = materialConfirmedCount(item.key);
      const rowCount = materialRows(item.key).length;
      const stateClass = completed ? 'completed' : '';
      const statusTone = isAll ? '' : completed ? 'completed' : item.missed > 0 ? 'has-missed' : item.pending > 0 ? 'pending' : 'incomplete';
      const statusChipTone = isAll ? '' : completed ? 'completed' : item.missed > 0 ? 'missed' : item.pending > 0 ? 'pending' : 'incomplete';
      const statusText = isAll ? '汇总' : completed ? '已完成' : item.missed > 0 ? '漏检' : item.pending > 0 ? '待复核' : '未完成';
      const encodedKey = encodeURIComponent(item.key);
      return `<article class="material-item ${active} ${stateClass} ${statusTone}">
        <button class="material-item-select" onclick="selectSourceMaterial('${encodedKey}')">
          <div class="material-type">材料类型：${escapeHtml(item.type)}</div>
          <div class="material-doc">对应材料名称：${escapeHtml(item.doc)}</div>
          <div class="material-meta"><span>总 ${item.total}</span><span>待复核 ${item.pending}</span><span><strong>漏检 ${item.missed}</strong></span><span>已确认 ${item.confirmed}</span></div>
        </button>
        <div class="material-item-footer">
          <span class="material-file-status ${statusChipTone}"><span class="material-status-chip ${statusChipTone}">${statusText}</span>${isAll ? '' : `已确认 ${readyCount}/${rowCount}`}</span>
          ${isAll ? '' : `<button class="material-submit" onclick="submitMaterial('${encodedKey}')">${completed ? '重新提交' : '提交文件'}</button>`}
        </div>
      </article>`;
    }
    function renderMaterialList() {
      if (!analysis) return;
      const stats = sourceMaterialStats();
      const list = document.getElementById('materialList');
      const fileSummary = materialFileSummary();
      document.getElementById('materialSummary').textContent = `筛选文件：${fileSummary.visibleFiles} 个 · 已完成文件：${fileSummary.reviewedFiles} / ${fileSummary.totalFiles}`;
      const visibleKeys = new Set(analysis.rows.filter(row => rowStatusMatchesFilter(row) && rowSearchMatches(row)).map(rowSourceKey));
      const allItem = stats.find(item => item.key === 'all');
      const groups = stats.filter(item => item.key !== 'all' && visibleKeys.has(item.key));
      if (selectedSourceMaterial !== 'all' && !visibleKeys.has(selectedSourceMaterial)) selectedSourceMaterial = 'all';
      const totalPages = Math.max(1, Math.ceil(groups.length / MATERIAL_PAGE_SIZE));
      materialListPage = Math.min(Math.max(Number(materialListPage) || 1, 1), totalPages);
      const start = (materialListPage - 1) * MATERIAL_PAGE_SIZE;
      const pageGroups = groups.slice(start, start + MATERIAL_PAGE_SIZE);
      list.innerHTML = [allItem, ...pageGroups].filter(Boolean).map(materialCardHtml).join('');
      renderMaterialPagination(groups.length);
      const current = stats.find(item => item.key === selectedSourceMaterial) || stats[0];
      document.getElementById('currentMaterial').textContent = current ? current.label : '全部材料';
      const status = document.getElementById('currentMaterialStatus');
      if (status) {
        const completed = current && materialIsCompleted(current.key);
        status.textContent = current && current.key !== 'all' ? (completed ? '已完成' : '未完成') : '汇总';
        status.classList.toggle('completed', Boolean(completed));
      }
    }
    function selectSourceMaterial(encodedKey) {
      selectedSourceMaterial = decodeURIComponent(encodedKey);
      renderMaterialList();
      renderRows();
    }
    function renderRows() {
      if (!analysis) return;
      renderCards();
      renderFormulaPanel();
      renderDimensionStats();
      renderErrorReasonStats();
      renderMaterialList();
      const tbody = document.getElementById('rows');
      const rows = analysis.rows.filter(rowVisible);
      if (!rows.length) {
        tbody.innerHTML = `<tr><td colspan="8" class="small">当前筛选下没有需要显示的检测项。</td></tr>`;
        return;
      }
      tbody.innerHTML = rows.map(row => {
        const key = rowKey(row);
        const arg = rowArg(row);
        const sourceRow = sourceRowNumber(row);
        const materialTypeLabel = displayMaterialType(row);
        const assigned = assignments[key] !== undefined ? assignments[key] : row.assignedId;
        const candidate = assigned === 'none' || assigned === null ? null : candidateById[assigned];
        const status = effectiveStatus(row);
        const aa = computedAa(row, candidate);
        const verdictClass = aa === '是' ? 'good' : aa === '否' ? 'bad' : '';
        const verdictSource = verdictOverrides[key] !== undefined ? '人工判定' : '自动判定';
        const modelTriggered = rowModelTriggered(row) ? '是' : '否';
        const modelCorrect = aa || '/';
        const modelIsCorrectText = candidate ? `isCorrect：${candidate.isCorrect === null || candidate.isCorrect === undefined ? 'null' : String(candidate.isCorrect)}` : 'isCorrect：未触发';
        const goldTriggerClass = row.goldTrigger === '是' ? 'yes' : 'no';
        const goldCorrectClass = row.goldCorrect === '是' ? 'yes' : row.goldCorrect === '否' ? 'no' : '';
        const modelTriggerClass = modelTriggered === '是' ? 'yes' : 'no';
        const modelCorrectClass = verdictClass;
        const selectedCategory = rowCategory(row);
        const categoryIsManual = categoryOverrides[key] !== undefined;
        const categoryOptions = categoryNames.map(category => `<option value="${category}" ${selectedCategory === category ? 'selected' : ''}>${category}</option>`).join('');
        const errorReason = rowErrorReason(row);
        const goldDescription = row.goldCorrect === '是' ? '' : row.goldRaw || '';
        const goldEditor = `
          <div class="gold-editor">
            <label>是否触发</label>
            <select aria-label="第 ${sourceRow} 行${row.materialCopyCount > 1 ? `第${row.materialCopyIndex}份` : ''}金标准是否触发" onchange="setGoldTrigger(${arg}, this.value)">
              <option value="是" ${row.goldTrigger === '是' ? 'selected' : ''}>是</option>
              <option value="否" ${row.goldTrigger === '否' ? 'selected' : ''}>否</option>
              <option value="" ${!row.goldTrigger ? 'selected' : ''}>待确认</option>
            </select>
            <label>是否正确</label>
            <select aria-label="第 ${sourceRow} 行${row.materialCopyCount > 1 ? `第${row.materialCopyIndex}份` : ''}金标准是否正确" onchange="setGoldCorrectChoice(${arg}, this.value)">
              <option value="是" ${row.goldCorrect === '是' ? 'selected' : ''}>是</option>
              <option value="否" ${row.goldCorrect === '否' ? 'selected' : ''}>否</option>
              <option value="" ${!row.goldCorrect ? 'selected' : ''}>待确认</option>
            </select>
            <textarea aria-label="第 ${sourceRow} 行金标准错误描述" placeholder="不正确时填写具体描述" oninput="updateGoldDescription(${arg}, this.value)">${escapeHtml(goldDescription)}</textarea>
            <div class="gold-editor-note">当前 M：${escapeHtml(row.goldTrigger || '-')}；N：${escapeHtml(row.goldRaw || '-')}</div>
            <button class="gold-editor-reset secondary" onclick="restoreGoldStandard(${arg})">恢复原始</button>
          </div>`;
        const errorReasonField = `
          <div class="error-reason-box">
            <label for="error-reason-${key}">错误原因（可选）</label>
            <textarea id="error-reason-${key}" placeholder="如需说明，请填写该检测项的错误原因" oninput="updateErrorReason(${arg}, this.value)">${escapeHtml(errorReason)}</textarea>
          </div>`;
        const scoreMap = {};
        for (const opt of row.options) scoreMap[opt.candidateId] = opt.score;
        const materialIndex = selectedMaterialIndex(row, candidate);
        const group = materialGroups[materialIndex] || materialGroups[0] || {candidates: []};
        const recommended = row.options.slice(0, 5).map(opt => {
          const c = candidateById[opt.candidateId];
          return `<button class="rec-chip" onclick="setAssignment(${arg}, ${c.id})">[${Number(opt.score).toFixed(2)}] #${c.id} ${escapeHtml(c.materialName)} / ${escapeHtml(c.ruleName)}</button>`;
        }).join('');
        return `<tr class="${status}" data-row="${key}">
          <td>${sourceRow}${row.materialCopyCount > 1 ? `<div class="small">第${row.materialCopyIndex}份</div>` : ''}</td>
          <td><span class="pill ${status}">${statusLabel(status)}</span><div class="small">分数 ${Number(row.score).toFixed(2)}</div></td>
          <td class="rule"><b>${text(materialTypeLabel)}</b><div>${text(row.materialName)}</div><div>${text(row.ruleName)}</div></td>
          <td>${goldEditor}</td>
          <td class="dimension">
            <select class="dimension-select ${categoryIsManual ? 'manual' : ''}" aria-label="第 ${sourceRow} 行统计维度" onchange="setCategory(${arg}, this.value)">${categoryOptions}</select>
            <div class="dimension-note">${categoryIsManual ? '人工标签' : '自动标签，可修改'}</div>
          </td>
          <td class="candidate">
            <div class="choice-card">
              <div class="choice-head">
                <div>
                  <div class="small">JSON材料</div>
                  <div class="choice-title">${escapeHtml(group.displayName || group.materialName || '未选择材料')}</div>
                  <div class="choice-meta">对应材料：${escapeHtml(group.outlineName || '无')}；材料内规则 ${group.candidates.length} 条</div>
                </div>
                <button class="choice-action" onclick="openPicker(${arg}, 'material')">更换材料</button>
              </div>
            </div>
            <div class="choice-card">
              <div class="choice-head">
                <div>
                  <div class="small">材料内规则</div>
                  <div class="choice-title">${candidate ? '#' + candidate.id + ' ' + escapeHtml(candidate.ruleName) : '无匹配'}</div>
                  <div class="choice-meta">${candidate ? '分数 ' + (scoreForRow(row, candidate.id) === undefined ? '非推荐候选' : Number(scoreForRow(row, candidate.id)).toFixed(2)) + '；是否触发=' + modelTriggered + '；' + modelIsCorrectText : '是否触发=否；isCorrect：未触发'}</div>
                </div>
                <button class="choice-action" onclick="openPicker(${arg}, 'rule')">更换规则</button>
              </div>
            </div>
            <div class="detail-line candidate-summary">
              <div class="summary-title">候选摘要</div>
              ${candidate ? '是否触发：' + modelTriggered + '；' + modelIsCorrectText + '；' + shortText(candidate.questionDesc, 95) : '是否触发：否；isCorrect：未触发'}
            </div>
            <div class="recommend small"><b>推荐候选</b><div class="recommend-list">${recommended}</div></div>
            <div class="actions">
              <button class="mini" onclick="confirmCurrent(${arg})">确认匹配</button>
              <button class="mini secondary" onclick="markNone(${arg})">改为无匹配</button>
              <button class="mini secondary" onclick="restoreAuto(${arg})">恢复自动</button>
            </div>
            <div class="manual">
              <input id="manual-${key}" list="candidateList" placeholder="JSON编号">
              <button onclick="applyManual(${arg})">应用编号</button>
            </div>
          </td>
          <td class="verdict">
            <div class="field-stack">
              <div class="field-row"><span class="field-label">是否触发</span><span class="field-value ${modelTriggerClass}">${modelTriggered}</span></div>
              <div class="field-row"><span class="field-label">是否正确</span><span class="field-value ${modelCorrectClass}">${modelCorrect}</span></div>
            </div>
            <div class="field-note">${verdictSource}，当前写入 AA 列：${aa}</div>
            <div class="verdict-actions">
              <button class="${verdictOverrides[key] === '是' ? 'active' : ''}" onclick="setVerdict(${arg}, '是')">AA=是</button>
              <button class="secondary ${verdictOverrides[key] === '否' ? 'active' : ''}" onclick="setVerdict(${arg}, '否')">AA=否</button>
              <button class="secondary" onclick="clearVerdict(${arg})">恢复自动判定</button>
            </div>
            ${errorReasonField}
          </td>
          <td class="desc">
            <div class="model-desc-box">
              <div class="desc-title">模型描述正文</div>
              ${candidate ? escapeHtml(candidate.questionDesc) : '<span class="model-desc-empty">无模型描述</span>'}
            </div>
          </td>
        </tr>`;
      }).join('');
    }
    function getRow(rowNum) {
      return analysis.rows.find(r => rowKey(r) === String(rowNum));
    }
    function invalidateMaterialCompletion(rowNum) {
      const row = getRow(rowNum);
      if (row) delete completedMaterials[rowSourceKey(row)];
    }
    function submitMaterial(encodedKey) {
      if (!analysis) {
        showFailure('提交失败：请先完成分析。');
        return;
      }
      const materialKey = decodeURIComponent(encodedKey);
      if (materialKey === 'all') {
        document.getElementById('status').textContent = '请先选择一个具体材料文件，再提交文件。';
        showFailure('提交失败：请先选择一个具体材料文件。');
        return;
      }
      const rows = materialRows(materialKey);
      const unconfirmed = rows.filter(row => !hasConfirmedAssignment(row));
      if (unconfirmed.length) {
        completedMaterials[materialKey] = false;
        document.getElementById('status').textContent = `该文件还有 ${unconfirmed.length} 条检测项未确认匹配，暂不能提交。`;
        showFailure(`提交失败：还有 ${unconfirmed.length} 条检测项未确认匹配。`);
        renderMaterialList();
        renderRows();
        return;
      }
      completedMaterials[materialKey] = true;
      document.getElementById('status').textContent = '文件已提交，状态为已完成。';
      showSuccess('提交成功：当前文件已完成复核。');
      renderMaterialList();
      renderRows();
    }
    function submitCurrentMaterial() {
      if (!selectedSourceMaterial || selectedSourceMaterial === 'all') {
        document.getElementById('status').textContent = '请先从左侧选择一个具体材料文件。';
        return;
      }
      submitMaterial(encodeURIComponent(selectedSourceMaterial));
    }
    function currentAssigned(rowNum) {
      const row = getRow(rowNum);
      if (!row) return null;
      const key = rowKey(row);
      return assignments[key] !== undefined ? assignments[key] : row.assignedId;
    }
    function openPicker(rowNum, mode) {
      pickerState = { rowNum, mode, page: 1 };
      document.getElementById('pickerSearch').value = '';
      document.getElementById('pickerOverlay').classList.remove('hide');
      renderPickerList();
      setTimeout(() => document.getElementById('pickerSearch').focus(), 0);
    }
    function closePicker() {
      pickerState = null;
      document.getElementById('pickerOverlay').classList.add('hide');
    }
    function overlayClick(event) {
      if (event.target.id === 'pickerOverlay') closePicker();
    }
    function pickerMatches(query, values) {
      if (!query) return true;
      const q = query.toLowerCase();
      return values.some(v => text(v).toLowerCase().includes(q));
    }
    function onPickerSearch() {
      if (pickerState) pickerState.page = 1;
      renderPickerList();
    }
    function renderPickerPagination(totalItems) {
      const pagination = document.getElementById('pickerPagination');
      if (!pagination) return;
      if (!totalItems) {
        pagination.classList.add('hide');
        pagination.innerHTML = '';
        return;
      }
      const totalPages = Math.max(1, Math.ceil(totalItems / PICKER_PAGE_SIZE));
      const page = Math.min(Math.max(Number(pickerState.page) || 1, 1), totalPages);
      pickerState.page = page;
      pagination.classList.remove('hide');
      pagination.innerHTML = `
        <button class="secondary" ${page <= 1 ? 'disabled' : ''} onclick="changePickerPage(${page - 1})">上一页</button>
        <span class="picker-page-status">第 ${page} / ${totalPages} 页，共 ${totalItems} 项</span>
        <button class="secondary" ${page >= totalPages ? 'disabled' : ''} onclick="changePickerPage(${page + 1})">下一页</button>`;
    }
    function changePickerPage(page) {
      if (!pickerState) return;
      pickerState.page = Math.max(1, Number(page) || 1);
      renderPickerList();
    }
    function renderPickerList() {
      if (!pickerState || !analysis) return;
      const row = getRow(pickerState.rowNum);
      const arg = rowArg(row);
      const candidate = rowCandidate(row);
      const query = document.getElementById('pickerSearch').value.trim();
      const title = document.getElementById('pickerTitle');
      const subtitle = document.getElementById('pickerSubtitle');
      const list = document.getElementById('pickerList');
      if (pickerState.mode === 'material') {
        title.textContent = `第 ${sourceRowNumber(row)} 行${row.materialCopyCount > 1 ? `（第${row.materialCopyIndex}份）` : ''}：选择 JSON 材料`;
        subtitle.textContent = `${displayMaterialType(row)} / ${row.materialName} / ${row.ruleName}`;
        const currentIndex = selectedMaterialIndex(row, candidate);
        const items = materialGroups.filter(g => pickerMatches(query, [g.index + 1, g.materialName, g.outlineName]));
        const start = (Math.max(1, Number(pickerState.page) || 1) - 1) * PICKER_PAGE_SIZE;
        const pageItems = items.slice(start, start + PICKER_PAGE_SIZE);
        list.innerHTML = pageItems.map(g => `
          <div class="picker-item ${Number(g.index) === Number(currentIndex) ? 'active' : ''}" role="button" tabindex="0" onclick="chooseMaterial(${arg}, ${g.index})" onkeydown="if (event.key === 'Enter' || event.key === ' ') { event.preventDefault(); chooseMaterial(${arg}, ${g.index}); }">
            <div class="picker-item-top">
              <span class="id-badge">材料 ${g.index + 1}</span>
              <span class="score-badge">${g.candidates.length} 条规则</span>
            </div>
            <div class="picker-rule">${escapeHtml(g.displayName || g.materialName || '未命名材料')}</div>
            <div class="picker-desc">对应材料名称：${escapeHtml(g.outlineName || '无')}</div>
          </div>`).join('') || `<div class="small">没有搜索到材料。</div>`;
        renderPickerPagination(items.length);
      } else {
        title.textContent = `第 ${sourceRowNumber(row)} 行${row.materialCopyCount > 1 ? `（第${row.materialCopyIndex}份）` : ''}：选择材料内规则`;
        subtitle.textContent = `${displayMaterialType(row)} / ${row.materialName} / ${row.ruleName}`;
        const materialIndex = selectedMaterialIndex(row, candidate);
        const group = materialGroups[materialIndex] || { candidates: [] };
        const activeId = rowAssignedId(row);
        const scoreMap = {};
        for (const opt of row.options) scoreMap[opt.candidateId] = opt.score;
        const candidates = group.candidates.filter(c => pickerMatches(query, [c.id, c.ruleName, c.questionDesc, c.materialName, c.outlineName]));
        const noneItem = `
          <div class="picker-item ${activeId === null ? 'active' : ''}" role="button" tabindex="0" onclick="chooseRule(${arg}, 'none')" onkeydown="if (event.key === 'Enter' || event.key === ' ') { event.preventDefault(); chooseRule(${arg}, 'none'); }">
            <div class="picker-item-top"><span class="id-badge">无匹配</span><span class="bool-badge false">Z=否</span></div>
            <div class="picker-rule">JSON 没有触发这个检测项</div>
            <div class="picker-desc">适用于 Excel M列为“是”但模型结果中找不到对应检测项的漏检情况。</div>
          </div>`;
        const ruleItems = candidates.map(c => {
          const score = scoreMap[c.id];
          return `<div class="picker-item ${Number(activeId) === Number(c.id) ? 'active' : ''}" role="button" tabindex="0" onclick="chooseRule(${arg}, ${c.id})" onkeydown="if (event.key === 'Enter' || event.key === ' ') { event.preventDefault(); chooseRule(${arg}, ${c.id}); }">
            <div class="picker-item-top">
              <span class="id-badge">#${c.id}</span>
              <span class="score-badge">${score === undefined ? '非推荐候选' : '分数 ' + Number(score).toFixed(2)}</span>
              <span class="bool-badge ${c.isCorrect ? 'true' : 'false'}">是否正确=${c.isCorrect}</span>
            </div>
            <div class="picker-rule">${escapeHtml(c.ruleName)}</div>
            <div class="picker-desc">${escapeHtml(c.questionDesc || '无描述')}</div>
          </div>`;
        });
        const items = [noneItem, ...ruleItems];
        const start = (Math.max(1, Number(pickerState.page) || 1) - 1) * PICKER_PAGE_SIZE;
        const pageItems = items.slice(start, start + PICKER_PAGE_SIZE);
        list.innerHTML = pageItems.join('') || `<div class="small">该材料下没有搜索到规则。</div>`;
        renderPickerPagination(items.length);
      }
    }
    function chooseMaterial(rowNum, index) {
      invalidateMaterialCompletion(rowNum);
      delete confirmedRows[rowNum];
      materialSelections[rowNum] = Number(index);
      const current = currentAssigned(rowNum);
      const currentCandidate = current === 'none' || current === null || current === undefined ? null : candidateById[current];
      if (currentCandidate && Number(currentCandidate.materialIndex) !== Number(index)) {
        assignments[rowNum] = 'none';
        delete verdictOverrides[rowNum];
      }
      pickerState = { rowNum, mode: 'rule', page: 1 };
      document.getElementById('pickerSearch').value = '';
      renderRows();
      renderPickerList();
    }
    function chooseRule(rowNum, id) {
      setAssignment(rowNum, id);
      closePicker();
    }
    function selectedMaterialIndex(row, candidate) {
      const key = rowKey(row);
      if (materialSelections[key] !== undefined) return Number(materialSelections[key]);
      if (candidate && candidate.materialIndex !== undefined) return Number(candidate.materialIndex);
      if (row.options.length) return Number(candidateById[row.options[0].candidateId].materialIndex);
      return 0;
    }
    function setMaterial(rowNum, value) {
      invalidateMaterialCompletion(rowNum);
      delete confirmedRows[rowNum];
      materialSelections[rowNum] = Number(value);
      const current = currentAssigned(rowNum);
      const currentCandidate = current === 'none' || current === null || current === undefined ? null : candidateById[current];
      if (currentCandidate && Number(currentCandidate.materialIndex) !== Number(value)) {
        assignments[rowNum] = 'none';
        delete verdictOverrides[rowNum];
      }
      renderRows();
    }
    function setAssignment(row, value) {
      invalidateMaterialCompletion(row);
      delete confirmedRows[row];
      assignments[row] = value === 'none' ? 'none' : Number(value);
      const c = candidateById[assignments[row]];
      if (c && c.materialIndex !== undefined) materialSelections[row] = c.materialIndex;
      delete verdictOverrides[row];
      renderRows();
    }
    function confirmCurrent(rowNum) {
      invalidateMaterialCompletion(rowNum);
      const assigned = currentAssigned(rowNum);
      assignments[rowNum] = assigned === 'none' || assigned === null || assigned === undefined ? 'none' : Number(assigned);
      confirmedRows[rowNum] = true;
      renderRows();
      focusNextPendingRow(rowNum);
    }
    function focusNextPendingRow(rowNum) {
      if (!analysis) return;
      const visibleRows = analysis.rows.filter(rowVisible);
      const currentIndex = visibleRows.findIndex(row => rowKey(row) === String(rowNum));
      const pendingRows = visibleRows.filter(row => rowKey(row) !== String(rowNum) && !hasConfirmedAssignment(row));
      const next = visibleRows.slice(Math.max(0, currentIndex + 1)).find(row => !hasConfirmedAssignment(row)) || pendingRows[0];
      if (!next) return;
      window.setTimeout(() => {
        const element = document.querySelector(`#rows tr[data-row="${rowKey(next)}"]`);
        const wrap = document.querySelector('.review-table-wrap');
        if (!element || !wrap) return;
        const top = element.offsetTop - wrap.offsetTop - (wrap.clientHeight - element.offsetHeight) / 2;
        wrap.scrollTo({top: Math.max(0, top), behavior:'smooth'});
        element.classList.add('row-focus');
        window.setTimeout(() => element.classList.remove('row-focus'), 1500);
      }, 40);
    }
    function markNone(rowNum) {
      invalidateMaterialCompletion(rowNum);
      delete confirmedRows[rowNum];
      assignments[rowNum] = 'none';
      delete verdictOverrides[rowNum];
      renderRows();
    }
    function restoreAuto(rowNum) {
      invalidateMaterialCompletion(rowNum);
      delete confirmedRows[rowNum];
      delete assignments[rowNum];
      delete materialSelections[rowNum];
      delete verdictOverrides[rowNum];
      renderRows();
    }
    function applyManual(rowNum) {
      const input = document.getElementById(`manual-${rowNum}`);
      const id = Number(String(input.value).trim());
      if (!Number.isInteger(id) || !candidateById[id]) {
        showFailure('应用失败：请输入有效的 JSON 编号。');
        return;
      }
      invalidateMaterialCompletion(rowNum);
      delete confirmedRows[rowNum];
      assignments[rowNum] = id;
      materialSelections[rowNum] = candidateById[id].materialIndex;
      delete verdictOverrides[rowNum];
      renderRows();
    }
    function setVerdict(rowNum, value) {
      const key = String(rowNum);
      invalidateMaterialCompletion(key);
      verdictOverrides[key] = value;
      confirmedRows[key] = true;
      const row = getRow(key);
      if (row) {
        if (row._statusBeforeVerdict === undefined) row._statusBeforeVerdict = row.status;
        row.status = 'confirmed';
      }
      document.getElementById('status').textContent = `已人工设置第 ${key} 行 AA=${value}`;
      renderRows();
    }
    function clearVerdict(rowNum) {
      const key = String(rowNum);
      invalidateMaterialCompletion(key);
      delete verdictOverrides[key];
      delete confirmedRows[key];
      const row = getRow(key);
      if (row) {
        row.status = row._statusBeforeVerdict !== undefined
          ? row._statusBeforeVerdict
          : row.assignedId === null || row.assignedId === undefined ? 'unmatched' : 'auto';
        delete row._statusBeforeVerdict;
      }
      document.getElementById('status').textContent = `已恢复第 ${key} 行自动 AA 判定`;
      renderRows();
    }
    function setCategory(rowNum, value) {
      if (!categoryNames.includes(value)) return;
      invalidateMaterialCompletion(rowNum);
      categoryOverrides[rowNum] = value;
      renderRows();
    }
    function confirmCurrentMaterial() {
      if (!analysis) {
        showFailure('确认失败：请先完成分析。');
        return;
      }
      let confirmedCount = 0;
      for (const row of analysis.rows) {
        if (selectedSourceMaterial !== 'all' && rowSourceKey(row) !== selectedSourceMaterial) continue;
        const key = rowKey(row);
        invalidateMaterialCompletion(key);
        const assigned = currentAssigned(key);
        assignments[key] = assigned === 'none' || assigned === null || assigned === undefined ? 'none' : Number(assigned);
        confirmedCount += 1;
      }
      renderRows();
      showSuccess(`确认成功：已确认 ${confirmedCount} 条检测项的匹配。`);
    }
    function jumpNextPendingMaterial() {
      if (!analysis) return;
      const stats = sourceMaterialStats().filter(item => item.key !== 'all' && item.pending > 0);
      if (!stats.length) {
        selectedSourceMaterial = 'all';
      } else {
        const currentIndex = stats.findIndex(item => item.key === selectedSourceMaterial);
        selectedSourceMaterial = stats[(currentIndex + 1) % stats.length].key;
      }
      document.getElementById('filter').value = 'review';
      renderMaterialList();
      renderRows();
    }
    async function exportErrorReasons() {
      if (!analysis) {
        showFailure('导出失败：请先完成分析。');
        return;
      }
      if (goldDirtyRows.size) {
        showFailure(`导出失败：还有 ${goldDirtyRows.size} 条金标准修改未保存，请先点击“保存金标准修改”。`);
        return;
      }
      const details = currentErrorReasonDetails();
      if (!details.length) {
        document.getElementById('status').textContent = '暂无已填写的错误原因可导出。';
        showFailure('导出失败：暂无已填写的错误原因。');
        return;
      }
      document.getElementById('status').textContent = '正在导出错误原因明细...';
      showToast('正在导出错误原因明细，请稍候。', 'info', 2400);
      try {
        const result = await post('/api/export-error-reasons', {
          analysis,
          assignments,
          verdictOverrides,
          categoryOverrides,
          errorReasonOverrides,
          outputPath: defaults.errorReasonOutputPath
        });
        document.getElementById('status').textContent = `错误原因明细已导出：${result.outputPath}，共 ${result.count} 条`;
        showSuccess(`导出成功：错误原因明细共 ${result.count} 条。`);
      } catch (err) {
        document.getElementById('status').textContent = err.message;
        showFailure(`导出失败：${err.message}`);
      }
    }
    async function exportWorkbook() {
      if (!analysis) {
        showFailure('导出失败：请先完成分析。');
        return;
      }
      if (goldDirtyRows.size) {
        showFailure(`导出失败：还有 ${goldDirtyRows.size} 条金标准修改未保存，请先点击“保存金标准修改”。`);
        return;
      }
      document.getElementById('status').textContent = '正在导出...';
      showToast('正在导出复核结果，请稍候。', 'info', 2400);
      try {
        const result = await post('/api/export', {
          analysis,
          assignments,
          verdictOverrides,
          categoryOverrides,
          errorReasonOverrides,
          summarySnapshot: currentSummarySnapshot()
        });
        document.getElementById('status').textContent = `已导出：${result.outputPath}`;
        showSuccess('导出成功：复核结果已生成。');
      } catch (err) {
        document.getElementById('status').textContent = err.message;
        showFailure(`导出失败：${err.message}`);
      }
    }
    loadHistoryList();
  </script>
</body>
</html>
"""


def update_analysis_job(job_id, **changes):
    with ANALYSIS_JOBS_LOCK:
        job = ANALYSIS_JOBS.get(job_id)
        if not job:
            return
        job.update(changes)
        job["updatedAt"] = time.time()


def cleanup_analysis_jobs():
    now = time.time()
    with ANALYSIS_JOBS_LOCK:
        expired_ids = [
            job_id
            for job_id, job in ANALYSIS_JOBS.items()
            if now - float(job.get("updatedAt") or job.get("createdAt") or now) > ANALYSIS_JOB_TTL_SECONDS
        ]
        for job_id in expired_ids:
            ANALYSIS_JOBS.pop(job_id, None)


def create_analysis_job(payload):
    cleanup_analysis_jobs()
    job_id = uuid.uuid4().hex
    now = time.time()
    with ANALYSIS_JOBS_LOCK:
        ANALYSIS_JOBS[job_id] = {
            "id": job_id,
            "status": "queued",
            "percent": 0,
            "message": "分析任务已创建，等待开始...",
            "current": None,
            "total": None,
            "result": None,
            "error": "",
            "createdAt": now,
            "updatedAt": now,
        }

    thread = threading.Thread(target=run_analysis_job, args=(job_id, copy.deepcopy(payload)), daemon=True)
    thread.start()
    return job_id


def run_analysis_job(job_id, payload):
    def progress(percent, message, current=None, total=None):
        update_analysis_job(
            job_id,
            status="running",
            percent=max(0, min(100, int(percent))),
            message=message,
            current=current,
            total=total,
        )

    try:
        progress(3, "正在保存上传文件...")
        json_path, excel_path = save_uploaded_inputs(payload)
        result = build_analysis(
            json_path,
            excel_path,
            float(payload.get("matchThreshold", 0.67)),
            float(payload.get("reviewThreshold", 0.82)),
            progress,
        )
        result["jsonFileName"] = (payload.get("jsonUpload") or {}).get("name") or Path(json_path).name
        result["excelFileName"] = (payload.get("excelUpload") or {}).get("name") or Path(excel_path).name
        update_analysis_job(
            job_id,
            status="done",
            percent=100,
            message="分析完成",
            current=None,
            total=None,
            result=result,
            error="",
        )
    except Exception as exc:
        update_analysis_job(
            job_id,
            status="error",
            percent=100,
            message="分析失败",
            current=None,
            total=None,
            result=None,
            error=str(exc),
        )


def analysis_job_snapshot(job_id):
    cleanup_analysis_jobs()
    with ANALYSIS_JOBS_LOCK:
        job = ANALYSIS_JOBS.get(job_id)
        return copy.deepcopy(job) if job else None


def batch_case_ids(value):
    raw = str(value or "")
    ids = []
    seen = set()
    for part in re.split(r"[\s,，;；]+", raw):
        case_id = part.strip()
        if not case_id or case_id in seen:
            continue
        ids.append(case_id)
        seen.add(case_id)
    if not ids:
        raise ValueError("请至少填写一个 caseId。")
    return ids


def save_batch_excel_upload(excel_upload, prefix="batch"):
    excel_upload = excel_upload or {}
    if not excel_upload.get("base64"):
        raise ValueError("请先上传 Excel 文件。")
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    excel_filename = safe_filename(excel_upload.get("name"), "gold_standard.xlsx")
    excel_path = UPLOAD_DIR / f"{stamp}-{prefix}-{excel_filename}"
    excel_path.write_bytes(base64.b64decode(excel_upload["base64"]))
    return excel_path, excel_filename


def save_batch_excel_files(payload):
    files = payload.get("excelFiles") or []
    saved = {}
    duplicates = set()
    for item in files:
        if not isinstance(item, dict) or not item.get("base64"):
            continue
        path, filename = save_batch_excel_upload(item, "batch-map")
        key = str(item.get("name") or filename).strip()
        safe_key = safe_filename(key, filename)
        for name in {key, safe_key, filename}:
            if not name:
                continue
            if name in saved:
                duplicates.add(name)
            saved[name] = {"path": str(path), "fileName": filename}
    if not saved:
        raise ValueError("映射模式请上传对应的 Excel 文件。")
    if duplicates:
        raise ValueError(f"存在重复 Excel 文件名：{', '.join(sorted(duplicates))}")
    return saved


def parse_case_excel_mapping(text):
    lines = [line.strip() for line in str(text or "").splitlines() if line.strip()]
    if not lines:
        raise ValueError("映射模式请上传 case-Excel 映射 CSV。")
    pairs = []
    for index, line in enumerate(lines):
        parts = [part.strip().strip('"') for part in re.split(r"[,，\t]", line) if part.strip()]
        if len(parts) < 2:
            raise ValueError(f"映射表第 {index + 1} 行格式不正确，应为 caseId,excelName。")
        if index == 0 and norm_text(parts[0]).lower() in {"caseid", "案件id"}:
            continue
        pairs.append({"caseId": parts[0], "excelName": parts[1]})
    if not pairs:
        raise ValueError("映射表没有有效数据。")
    return pairs


def resolve_batch_cases(payload):
    mode = payload.get("excelMode") or "shared"
    excel_upload = payload.get("excelUpload") or {}
    if mode == "mapped":
        excel_files = save_batch_excel_files(payload)
        pairs = parse_case_excel_mapping(payload.get("mappingContent") or "")
        cases = []
        missing = []
        seen = set()
        for pair in pairs:
            case_id = pair["caseId"]
            excel_name = pair["excelName"]
            if case_id in seen:
                continue
            seen.add(case_id)
            match = excel_files.get(excel_name) or excel_files.get(safe_filename(excel_name, excel_name)) or excel_files.get(Path(excel_name).name)
            if not match:
                missing.append(f"{case_id}->{excel_name}")
                continue
            cases.append({
                "caseId": case_id,
                "excelPath": match["path"],
                "excelFileName": match["fileName"],
            })
        if missing:
            raise ValueError(f"映射表中有 Excel 未上传：{'; '.join(missing[:8])}{'...' if len(missing) > 8 else ''}")
        return cases, "按映射表选择 Excel"
    ids = batch_case_ids(payload.get("caseIds"))
    excel_path, excel_filename = save_batch_excel_upload(excel_upload, "batch")
    return [
        {"caseId": case_id, "excelPath": str(excel_path), "excelFileName": excel_filename}
        for case_id in ids
    ], excel_filename


def batch_summary_from_rows(rows):
    totals = {
        "totalCases": len(rows),
        "successCases": sum(1 for row in rows if row.get("status") == "success"),
        "failedCases": sum(1 for row in rows if row.get("status") == "error"),
        "jsonRules": sum(int(row.get("jsonRules") or 0) for row in rows),
        "excelRows": sum(int(row.get("excelRows") or 0) for row in rows),
        "reviewRows": sum(int(row.get("reviewRows") or 0) for row in rows),
        "missedRows": sum(int(row.get("missedRows") or 0) for row in rows),
        "modelTriggered": sum(int(row.get("modelTriggered") or 0) for row in rows),
        "modelCorrect": sum(int(row.get("modelCorrect") or 0) for row in rows),
        "expectedTotal": sum(int(row.get("expectedTotal") or 0) for row in rows),
        "expectedTriggered": sum(int(row.get("expectedTriggered") or 0) for row in rows),
        "errorTotal": sum(int(row.get("errorTotal") or 0) for row in rows),
        "errorRecalled": sum(int(row.get("errorRecalled") or 0) for row in rows),
    }
    totals["triggerRate"] = pct(totals["expectedTriggered"], totals["expectedTotal"])
    totals["recall"] = pct(totals["errorRecalled"], totals["errorTotal"])
    totals["accuracy"] = pct(totals["modelCorrect"], totals["modelTriggered"])
    return totals


def cleanup_batch_jobs():
    now = time.time()
    with BATCH_JOBS_LOCK:
        expired = [
            job_id for job_id, job in BATCH_JOBS.items()
            if now - float(job.get("updatedAt") or job.get("createdAt") or now) > BATCH_JOB_TTL_SECONDS
        ]
        for job_id in expired:
            BATCH_JOBS.pop(job_id, None)


def update_batch_job(job_id, **changes):
    with BATCH_JOBS_LOCK:
        job = BATCH_JOBS.get(job_id)
        if not job:
            return
        job.update(changes)
        job["updatedAt"] = time.time()


def db_batch_rows(job_id):
    with db_connect() as conn:
        rows = conn.execute(
            """
            SELECT case_id AS caseId, excel_file_name AS excelFileName, excel_path AS excelPath,
                   status, message, json_rules AS jsonRules, excel_rows AS excelRows,
                   review_rows AS reviewRows, missed_rows AS missedRows, model_triggered AS modelTriggered,
                   model_correct AS modelCorrect, expected_total AS expectedTotal,
                   expected_triggered AS expectedTriggered, error_total AS errorTotal,
                   error_recalled AS errorRecalled, trigger_rate AS triggerRate, recall, accuracy,
                   analysis_path AS analysisPath
            FROM batch_cases
            WHERE job_id=?
            ORDER BY id
            """,
            (job_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def create_batch_job(payload):
    cleanup_batch_jobs()
    init_app_db()
    base_url = str(payload.get("baseUrl") or DEFAULT_MODEL_RESULT_API).strip()
    cases, excel_filename = resolve_batch_cases(payload)
    job_id = uuid.uuid4().hex
    created_at = now_text()
    job = {
        "id": job_id,
        "status": "queued",
        "message": "批量任务已创建，等待开始...",
        "percent": 0,
        "total": len(cases),
        "completed": 0,
        "success": 0,
        "failed": 0,
        "baseUrl": base_url,
        "excelFileName": excel_filename,
        "cases": [{**case, "status": "queued", "message": "等待处理"} for case in cases],
        "summary": batch_summary_from_rows([]),
        "createdAt": time.time(),
        "updatedAt": time.time(),
    }
    with BATCH_JOBS_LOCK:
        BATCH_JOBS[job_id] = job
    with db_connect() as conn:
        conn.execute(
            """
            INSERT OR REPLACE INTO batch_jobs(id, base_url, excel_file_name, total, completed, success, failed, status, message, created_at, updated_at)
            VALUES(?, ?, ?, ?, 0, 0, 0, ?, ?, ?, ?)
            """,
            (job_id, base_url, excel_filename, len(cases), "queued", job["message"], created_at, created_at),
        )
        for case in cases:
            conn.execute(
                """
                INSERT INTO batch_cases(job_id, case_id, excel_file_name, excel_path, status, message, updated_at)
                VALUES(?, ?, ?, ?, ?, ?, ?)
                """,
                (job_id, case["caseId"], case["excelFileName"], case["excelPath"], "queued", "等待处理", created_at),
            )
    thread = threading.Thread(
        target=run_batch_job,
        args=(job_id, base_url, cases, float(payload.get("matchThreshold", 0.67)), float(payload.get("reviewThreshold", 0.82))),
        daemon=True,
    )
    thread.start()
    return job_id


def run_batch_job(job_id, base_url, cases_input, match_threshold, review_threshold):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    success = failed = completed = 0
    cases = []
    total = len(cases_input)
    for index, case_item in enumerate(cases_input, start=1):
        case_id = case_item["caseId"]
        excel_path = case_item["excelPath"]
        excel_file_name = case_item["excelFileName"]
        message = f"正在处理 {index} / {total}：{case_id}"
        update_batch_job(job_id, status="running", message=message, percent=int((index - 1) * 100 / max(total, 1)))
        try:
            fetched = fetch_model_result_from_api(base_url, case_id)
            json_path = UPLOAD_DIR / f"{time.strftime('%Y%m%d-%H%M%S')}-case-{safe_filename(case_id, 'case')}.json"
            json_path.write_text(json.dumps(fetched["content"], ensure_ascii=False), encoding="utf-8")
            analysis = build_analysis(json_path, excel_path, match_threshold, review_threshold)
            analysis["jsonFileName"] = f"接口结果-{case_id}.json"
            analysis["excelFileName"] = excel_file_name
            analysis_path = HISTORY_DIR / f"batch-{job_id}-{safe_filename(case_id, 'case')}.json"
            HISTORY_DIR.mkdir(parents=True, exist_ok=True)
            analysis_path.write_text(json.dumps(analysis, ensure_ascii=False), encoding="utf-8")
            summary = analysis.get("summary", {})
            row = {
                "caseId": case_id,
                "excelFileName": excel_file_name,
                "excelPath": excel_path,
                "status": "success",
                "message": "成功",
                "jsonRules": summary.get("jsonRules", len(fetched["content"])),
                "excelRows": summary.get("excelRows", 0),
                "reviewRows": summary.get("reviewRows", 0),
                "missedRows": summary.get("missedRows", 0),
                "modelTriggered": summary.get("modelTriggered", 0),
                "modelCorrect": summary.get("modelCorrect", 0),
                "expectedTotal": summary.get("expectedTotal", 0),
                "expectedTriggered": summary.get("expectedTriggered", 0),
                "errorTotal": summary.get("errorTotal", 0),
                "errorRecalled": summary.get("errorRecalled", 0),
                "triggerRate": summary.get("triggerRate", 0),
                "recall": summary.get("recall", 0),
                "accuracy": summary.get("accuracy", 0),
                "analysisPath": str(analysis_path),
            }
            success += 1
        except Exception as exc:
            row = {
                "caseId": case_id,
                "excelFileName": excel_file_name,
                "excelPath": excel_path,
                "status": "error",
                "message": str(exc),
                "jsonRules": 0,
                "excelRows": 0,
                "reviewRows": 0,
                "missedRows": 0,
                "modelTriggered": 0,
                "modelCorrect": 0,
                "expectedTotal": 0,
                "expectedTriggered": 0,
                "errorTotal": 0,
                "errorRecalled": 0,
                "triggerRate": 0,
                "recall": 0,
                "accuracy": 0,
                "analysisPath": "",
            }
            failed += 1
        completed += 1
        cases.append(row)
        updated_at = now_text()
        with db_connect() as conn:
            conn.execute(
                """
                UPDATE batch_cases
                   SET status=?, message=?, json_rules=?, excel_rows=?, review_rows=?, missed_rows=?,
                       model_triggered=?, model_correct=?, expected_total=?, expected_triggered=?,
                       error_total=?, error_recalled=?, trigger_rate=?, recall=?, accuracy=?,
                       analysis_path=?, updated_at=?
                 WHERE job_id=? AND case_id=?
                """,
                (
                    row["status"], row["message"], row["jsonRules"], row["excelRows"], row["reviewRows"], row["missedRows"],
                    row["modelTriggered"], row["modelCorrect"], row["expectedTotal"], row["expectedTriggered"],
                    row["errorTotal"], row["errorRecalled"], row["triggerRate"], row["recall"], row["accuracy"],
                    row["analysisPath"], updated_at, job_id, case_id,
                ),
            )
            conn.execute(
                """
                UPDATE batch_jobs
                   SET completed=?, success=?, failed=?, status=?, message=?, updated_at=?
                 WHERE id=?
                """,
                (completed, success, failed, "running", f"已完成 {completed} / {total}", updated_at, job_id),
            )
        update_batch_job(
            job_id,
            cases=cases + [{**item, "status": "queued", "message": "等待处理"} for item in cases_input[index:]],
            summary=batch_summary_from_rows(cases),
            completed=completed,
            success=success,
            failed=failed,
            percent=int(completed * 100 / max(total, 1)),
            message=f"已完成 {completed} / {total}",
        )
    final_status = "done" if failed == 0 else "done_with_errors"
    final_message = f"批量完成：成功 {success}，失败 {failed}"
    update_batch_job(job_id, status=final_status, message=final_message, percent=100)
    with db_connect() as conn:
        conn.execute(
            "UPDATE batch_jobs SET status=?, message=?, updated_at=? WHERE id=?",
            (final_status, final_message, now_text(), job_id),
        )


def batch_job_snapshot(job_id):
    cleanup_batch_jobs()
    with BATCH_JOBS_LOCK:
        job = BATCH_JOBS.get(job_id)
        if job:
            return copy.deepcopy(job)
    rows = db_batch_rows(job_id)
    if not rows:
        return None
    with db_connect() as conn:
        meta = conn.execute("SELECT * FROM batch_jobs WHERE id=?", (job_id,)).fetchone()
    if not meta:
        return None
    return {
        "id": job_id,
        "status": meta["status"],
        "message": meta["message"],
        "percent": int((meta["completed"] or 0) * 100 / max(meta["total"] or 1, 1)),
        "total": meta["total"],
        "completed": meta["completed"],
        "success": meta["success"],
        "failed": meta["failed"],
        "cases": rows,
        "summary": batch_summary_from_rows(rows),
    }


def write_batch_export(job_id):
    snapshot = batch_job_snapshot(job_id)
    if not snapshot:
        raise ValueError("批量任务不存在。")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_DIR / f"批量测试结果_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    summary = snapshot.get("summary") or {}
    rows = [
        ("总case数", summary.get("totalCases", 0)),
        ("成功case数", summary.get("successCases", 0)),
        ("失败case数", summary.get("failedCases", 0)),
        ("JSON规则数", summary.get("jsonRules", 0)),
        ("Excel检测项", summary.get("excelRows", 0)),
        ("漏检项", summary.get("missedRows", 0)),
        ("触发率", pct_text(summary.get("triggerRate", 0))),
        ("召回率", pct_text(summary.get("recall", 0))),
        ("准确率", pct_text(summary.get("accuracy", 0))),
    ]
    for row in rows:
        ws.append(row)
    detail = wb.create_sheet("明细")
    headers = ["caseId", "金标准Excel", "状态", "信息", "JSON规则数", "Excel检测项", "待复核", "漏检", "触发率", "召回率", "准确率"]
    detail.append(headers)
    for item in snapshot.get("cases", []):
        detail.append([
            item.get("caseId", ""),
            item.get("excelFileName", ""),
            item.get("status", ""),
            item.get("message", ""),
            item.get("jsonRules", 0),
            item.get("excelRows", 0),
            item.get("reviewRows", 0),
            item.get("missedRows", 0),
            pct_text(float(item.get("triggerRate") or 0)),
            pct_text(float(item.get("recall") or 0)),
            pct_text(float(item.get("accuracy") or 0)),
        ])
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 20
    wb.save(output_path)
    return output_path


class ToolHandler(BaseHTTPRequestHandler):
    server_version = "MatchingReviewTool/2.0"

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        if parsed.path == "/api/history":
            self.send_json(refreshed_history_index())
            return
        if parsed.path == "/api/history/detail":
            query = urllib.parse.parse_qs(parsed.query)
            record = load_history_record(query.get("id", [""])[0])
            if not record:
                self.send_json({"error": "历史记录不存在"}, 404)
                return
            self.send_json(record)
            return
        if parsed.path == "/api/analyze/status":
            query = urllib.parse.parse_qs(parsed.query)
            job = analysis_job_snapshot(query.get("id", [""])[0])
            if not job:
                self.send_json({"error": "分析任务不存在或已过期"}, 404)
                return
            self.send_json(job)
            return
        if parsed.path == "/api/batch/status":
            query = urllib.parse.parse_qs(parsed.query)
            job = batch_job_snapshot(query.get("id", [""])[0])
            if not job:
                self.send_json({"error": "批量任务不存在或已过期"}, 404)
                return
            self.send_json(job)
            return
        if parsed.path == "/api/batch/case":
            query = urllib.parse.parse_qs(parsed.query)
            job_id = query.get("id", [""])[0]
            case_id = query.get("caseId", [""])[0]
            rows = [row for row in db_batch_rows(job_id) if str(row.get("caseId")) == str(case_id)]
            if not rows or not rows[0].get("analysisPath"):
                self.send_json({"error": "该 case 没有可载入的分析明细"}, 404)
                return
            path = Path(rows[0]["analysisPath"])
            if not path.exists():
                self.send_json({"error": "分析明细文件不存在"}, 404)
                return
            self.send_json(json.loads(path.read_text(encoding="utf-8")))
            return
        if parsed.path == "/api/case/detail":
            query = urllib.parse.parse_qs(parsed.query)
            try:
                result = fetch_case_detail_from_api(
                    query.get("baseUrl", [DEFAULT_CASE_DETAIL_API])[0] or DEFAULT_CASE_DETAIL_API,
                    query.get("caseId", [""])[0],
                )
                self.send_json(result)
            except ValueError as exc:
                self.send_json({"error": str(exc)}, 400)
            return
        if parsed.path not in {"/", "/index.html"}:
            self.send_json({"error": "not found"}, 404)
            return
        defaults = {
            "outputPath": str(DEFAULT_OUTPUT),
            "errorReasonOutputPath": str(DEFAULT_ERROR_REASON_OUTPUT),
            "modelResultApi": DEFAULT_MODEL_RESULT_API,
            "caseDetailApi": DEFAULT_CASE_DETAIL_API,
        }
        html = HTML.replace("__DEFAULTS__", json.dumps(defaults, ensure_ascii=False))
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.end_headers()
        self.wfile.write(html.encode("utf-8"))

    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8") or "{}")
            if self.path == "/api/rules/load":
                excel_upload = payload.get("excelUpload") or {}
                if not excel_upload.get("base64"):
                    raise ValueError("请先上传规则 Excel 文件。")
                UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
                stamp = time.strftime("%Y%m%d-%H%M%S")
                excel_filename = safe_filename(excel_upload.get("name"), "rules.xlsx")
                excel_path = UPLOAD_DIR / f"{stamp}-annotation-{excel_filename}"
                excel_path.write_bytes(base64.b64decode(excel_upload["base64"]))
                self.send_json(load_rule_annotation_workbook(excel_path))
            elif self.path == "/api/rules/export":
                excel_path = safe_uploaded_excel_path(payload.get("sourcePath"))
                output_path = export_rule_annotation_workbook(
                    excel_path,
                    payload.get("sheetName") or "",
                    payload.get("edits") or [],
                    payload.get("newRules") or [],
                    payload.get("mColumn", 13),
                    payload.get("nColumn", 14),
                )
                self.send_json({"outputPath": str(output_path)})
            elif self.path == "/api/analyze":
                job_id = create_analysis_job(payload)
                self.send_json({"jobId": job_id})
            elif self.path == "/api/model-result/fetch":
                result = fetch_model_result_from_api(
                    payload.get("baseUrl") or DEFAULT_MODEL_RESULT_API,
                    payload.get("caseId") or "",
                )
                self.send_json(result)
            elif self.path == "/api/batch/start":
                job_id = create_batch_job(payload)
                self.send_json({"jobId": job_id})
            elif self.path == "/api/batch/export":
                output_path = write_batch_export(payload.get("jobId") or "")
                self.send_json({"outputPath": str(output_path)})
            elif self.path == "/api/export":
                analysis = payload["analysis"]
                reviewed = apply_review_assignments(
                    analysis,
                    payload.get("assignments", {}),
                    payload.get("verdictOverrides", {}),
                    payload.get("categoryOverrides", {}),
                    payload.get("errorReasonOverrides", {}),
                )
                reviewed["summary"] = merge_summary_snapshot(reviewed.get("summary"), payload.get("summarySnapshot"))
                output_path = Path(payload.get("outputPath") or review_output_path(analysis))
                write_reviewed_workbook(analysis["excelPath"], output_path, reviewed)
                self.send_json({"outputPath": str(output_path), "summary": reviewed["summary"]})
            elif self.path == "/api/export-error-reasons":
                analysis = payload["analysis"]
                reviewed = apply_review_assignments(
                    analysis,
                    payload.get("assignments", {}),
                    payload.get("verdictOverrides", {}),
                    payload.get("categoryOverrides", {}),
                    payload.get("errorReasonOverrides", {}),
                )
                output_path = Path(payload.get("outputPath") or DEFAULT_ERROR_REASON_OUTPUT)
                write_error_reason_workbook(reviewed, output_path)
                self.send_json({
                    "outputPath": str(output_path),
                    "count": len(collect_error_reason_details(reviewed)),
                })
            elif self.path == "/api/history/save":
                record = save_history_record(
                    payload["analysis"],
                    payload.get("assignments", {}),
                    payload.get("verdictOverrides", {}),
                    payload.get("categoryOverrides", {}),
                    payload.get("errorReasonOverrides", {}),
                    payload.get("title", ""),
                    payload.get("completedMaterials", {}),
                    payload.get("confirmedRows", {}),
                    payload.get("summarySnapshot"),
                    payload.get("modelReviewStart", ""),
                    payload.get("modelReviewEnd", ""),
                )
                self.send_json(record)
            elif self.path == "/api/history/export":
                ids = payload.get("ids") or []
                records = []
                seen = set()
                for record_id in ids:
                    record_id = history_id(record_id)
                    if not record_id or record_id in seen:
                        continue
                    record = load_history_record(record_id)
                    if record:
                        records.append(record)
                        seen.add(record_id)
                if not records:
                    raise ValueError("没有找到可导出的历史记录。")
                exported_at = time.strftime("%Y-%m-%d %H:%M:%S")
                output_path = OUTPUT_DIR / f"比对历史导出_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
                write_history_export_workbook(records, output_path, exported_at)
                self.send_json({"outputPath": str(output_path), "count": len(records)})
            elif self.path == "/api/history/delete":
                deleted = delete_history_record(payload.get("id", ""))
                if not deleted:
                    self.send_json({"error": "历史记录不存在"}, 404)
                    return
                self.send_json({"deleted": True})
            else:
                self.send_json({"error": "not found"}, 404)
        except ValueError as exc:
            self.send_json({"error": str(exc)}, 400)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def log_message(self, fmt, *args):
        return

    def send_json(self, data, status=200):
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def main():
    parser = argparse.ArgumentParser(description="JSON和Excel匹配复核工具")
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", default=8765, type=int)
    args = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    init_app_db()
    server = ThreadingHTTPServer((args.host, args.port), ToolHandler)
    print(f"匹配复核工具已启动：http://{args.host}:{args.port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()

"""Stage-based offline evaluation for subject identity review."""

import base64
import json
import re
import sqlite3
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = {
    "案件信息": ["case_id", "case_name", "first_violation_date", "complexity", "split", "notes"],
    "文件分类": ["case_id", "file_id", "item_id", "file_name", "expected_file_type", "page", "should_process", "annotation_status"],
    "职位抽取": ["record_id", "case_id", "file_id", "item_id", "file_name", "job", "start_date", "end_date", "job_type", "file_type", "page", "should_display", "before_first_violation", "evidence_text"],
    "任职分组": ["group_id", "record_id", "case_id", "normalized_job", "expected_start", "expected_end", "relation", "should_merge"],
    "身份材料": ["identity_id", "case_id", "identity_type", "identity", "case_input", "report_input", "material_file_id", "material_item_id", "material_name", "page", "should_display"],
    "业务规则": ["rule_id", "case_id", "rule_type", "input_description", "expected_result", "expected_display", "expected_issue", "expected_color_status", "notes"],
}

CHINESE_HEADERS = {
    "案件信息": {"case_id": "案件编号", "case_name": "案件名称", "first_violation_date": "第一次违纪违法时间", "complexity": "案件复杂度", "split": "数据集分组", "notes": "备注"},
    "文件分类": {"case_id": "案件编号", "file_id": "文件编号", "item_id": "材料编号", "file_name": "文件名称", "expected_file_type": "标准文件类型", "page": "页码", "should_process": "是否参与主体身份处理", "annotation_status": "标注状态"},
    "职位抽取": {"record_id": "记录编号", "case_id": "案件编号", "file_id": "文件编号", "item_id": "材料编号", "file_name": "文件名称", "job": "标准职位或身份", "start_date": "开始时间", "end_date": "结束时间", "job_type": "职位类型", "file_type": "文件类型", "page": "页码", "should_display": "是否应展示", "before_first_violation": "是否在第一次违纪违法之前", "evidence_text": "证据原文"},
    "任职分组": {"group_id": "标准分组编号", "record_id": "记录编号", "case_id": "案件编号", "normalized_job": "标准职位名称", "expected_start": "标准开始时间", "expected_end": "标准结束时间", "relation": "材料关系", "should_merge": "是否应合并"},
    "身份材料": {"identity_id": "身份编号", "case_id": "案件编号", "identity_type": "身份类型", "identity": "标准身份名称", "case_input": "案件录入身份", "report_input": "审查调查报告身份", "material_file_id": "关联文件编号", "material_item_id": "关联材料编号", "material_name": "关联材料名称", "page": "页码", "should_display": "是否应展示"},
    "业务规则": {"rule_id": "规则编号", "case_id": "案件编号", "rule_type": "规则类型", "input_description": "测试条件", "expected_result": "预期结果", "expected_display": "是否应展示", "expected_issue": "是否应进入阅卷问题", "expected_color_status": "预期颜色或状态", "notes": "备注"},
}

EXAMPLES = {
    "案件信息": ["C001", "示例案件", "2022-01-01", "复杂", "开发集", "示例行可删除"],
    "文件分类": ["C001", "F001", "I001", "任职文件1", "任职文件", "3", "是", "已确认"],
    "职位抽取": ["R001", "C001", "F001", "I001", "任职文件1", "XX局局长", "2020-01", "2022-06", "职务", "任职文件", "3", "是", "是", "任职文件原文"],
    "任职分组": ["G001", "R001", "C001", "XX局局长", "2020-01", "2022-06", "任职", "是"],
    "身份材料": ["ID001", "C001", "特殊身份", "人大代表", "人大代表", "报告原文", "F002", "I002", "代表证明", "1", "是"],
    "业务规则": ["BR001", "C001", "第一次违纪违法前职位过滤", "存在较早职位", "过滤或标记", "否", "否", "灰色", ""],
}


def text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value).strip()


def norm(value):
    return re.sub(r"[\s\W_]+", "", text(value).lower(), flags=re.UNICODE)


def yes(value):
    return text(value).lower() in {"是", "1", "true", "yes", "应", "需要"}


def similarity(left, right):
    left, right = norm(left), norm(right)
    if not left or not right:
        return 0.0
    if left in right or right in left:
        return 0.65 + 0.35 * min(len(left), len(right)) / max(len(left), len(right))
    return SequenceMatcher(None, left, right).ratio()


def unwrap(payload):
    current = payload
    for _ in range(3):
        if not isinstance(current, dict) or "data" not in current:
            break
        current = current["data"]
    return current if isinstance(current, dict) else {}


def pick(item, *keys):
    if not isinstance(item, dict):
        return ""
    for key in keys:
        if item.get(key) not in (None, ""):
            return text(item[key])
    return ""


def as_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def normalize_file(item, case_id=""):
    return {
        "caseId": pick(item, "case_id", "caseId") or case_id,
        "fileId": pick(item, "file_id", "fileId"), "itemId": pick(item, "item_id", "itemId"),
        "fileName": pick(item, "file_name", "fileName", "item_name", "itemName", "name"),
        "fileType": pick(item, "file_type", "fileType", "type", "material_type", "materialType"),
        "page": pick(item, "page", "page_no", "pageNo"), "raw": item,
    }


def normalize_record(item, default_type="", case_id="", group_id=""):
    return {
        "caseId": pick(item, "case_id", "caseId") or case_id,
        "recordId": pick(item, "record_id", "recordId", "id"),
        "fileId": pick(item, "file_id", "fileId"), "itemId": pick(item, "item_id", "itemId"),
        "fileName": pick(item, "file_name", "fileName", "item_name", "itemName"),
        "job": pick(item, "job", "position", "rank", "name", "title", "职位", "职务"),
        "startDate": pick(item, "start_date", "startDate", "begin_date", "beginDate"),
        "endDate": pick(item, "end_date", "endDate"),
        "jobType": pick(item, "job_type", "jobType", "type") or default_type,
        "fileType": pick(item, "file_type", "fileType", "material_type", "materialType"),
        "page": pick(item, "page", "page_no", "pageNo"),
        "groupId": pick(item, "group_id", "groupId") or group_id,
        "display": item.get("display", item.get("should_display", item.get("shouldDisplay"))), "raw": item,
    }


def flatten_records(value, default_type="", case_id=""):
    records, groups = [], []
    for index, item in enumerate(as_list(value)):
        if not isinstance(item, dict):
            records.append(normalize_record({"job": item}, default_type, case_id))
            continue
        nested = item.get("records") or item.get("items") or item.get("positions") or item.get("jobs")
        group_id = pick(item, "group_id", "groupId") or (f"model-group-{index + 1}" if nested else "")
        if nested:
            group_records = [normalize_record(child if isinstance(child, dict) else {"job": child}, default_type, case_id, group_id) for child in as_list(nested)]
            records.extend(group_records)
            groups.append({"groupId": group_id, "job": pick(item, "job", "position", "name") or (group_records[0]["job"] if group_records else ""),
                           "startDate": pick(item, "start_date", "startDate") or (group_records[0]["startDate"] if group_records else ""),
                           "endDate": pick(item, "end_date", "endDate") or (group_records[-1]["endDate"] if group_records else ""),
                           "recordKeys": [record_key(row) for row in group_records]})
        else:
            record = normalize_record(item, default_type, case_id, group_id)
            records.append(record)
            if group_id:
                groups.append({"groupId": group_id, "job": record["job"], "startDate": record["startDate"], "endDate": record["endDate"], "recordKeys": [record_key(record)]})
    return records, groups


def record_key(row):
    return "|".join(norm(row.get(key)) for key in ("fileId", "itemId", "job", "startDate", "endDate"))


def normalize_identity(item, identity_type, case_id=""):
    materials = []
    for material in as_list((item.get("materials") or item.get("proofDocument")) if isinstance(item, dict) else None):
        if isinstance(material, dict):
            materials.append({"fileId": pick(material, "file_id", "fileId"), "itemId": pick(material, "item_id", "itemId"),
                              "name": pick(material, "material", "name", "file_name", "fileName", "itemName"), "page": pick(material, "page", "pageNo", "pageNum")})
    case_input = pick(item, "caseInput", "case_input", "creationSelection")
    report_input = pick(item, "reportInput", "report_input", "investigateReport")
    return {"caseId": pick(item, "case_id", "caseId") or case_id, "identityType": identity_type,
            "identity": pick(item, "identity", "name", "身份") or case_input or report_input, "caseInput": case_input,
            "reportInput": report_input, "materials": materials, "raw": item}


def flatten_api_position_groups(value, kind, case_id=""):
    """Adapt the actual /subject_identity/identity grouped response."""
    records, groups = [], []
    for index, item in enumerate(as_list(value)):
        if not isinstance(item, dict):
            continue
        group_id = pick(item, "groupId", "group_id") or f"model-{kind}-{index + 1}"
        group_records = []
        if kind == "职级":
            form_job = pick(item, "formRank", "rank", "job")
            investigate_job = pick(item, "investigateRank")
            form_start, form_end = pick(item, "formStartTime"), pick(item, "formEndTime")
            promotion_key, promotion_job = "promotionDocs", "rank"
        else:
            form_job = pick(item, "formPosition", "position", "job")
            investigate_job = pick(item, "investigatePosition")
            form_start, form_end = pick(item, "formStartTime"), pick(item, "formEndTime")
            promotion_key, promotion_job = "appointmentDocs", "position"
        if form_job:
            group_records.append({"caseId": case_id, "recordId": "", "fileId": pick(item, "fileId"), "itemId": pick(item, "itemId"),
                                  "fileName": pick(item, "itemName"), "job": form_job, "startDate": form_start, "endDate": form_end,
                                  "jobType": kind, "fileType": "任职履历", "page": pick(item, "pageNum", "page"), "groupId": group_id, "raw": item})
        if investigate_job:
            group_records.append({"caseId": case_id, "recordId": "", "fileId": "", "itemId": "", "fileName": "审查调查报告",
                                  "job": investigate_job, "startDate": pick(item, "investigateStartTime"), "endDate": "",
                                  "jobType": kind, "fileType": "审查调查报告", "page": "", "groupId": group_id, "raw": item})
        for doc in as_list(item.get(promotion_key)):
            if not isinstance(doc, dict): continue
            group_records.append({"caseId": case_id, "recordId": "", "fileId": pick(doc, "fileId"), "itemId": pick(doc, "itemId"),
                                  "fileName": pick(doc, "itemName"), "job": pick(doc, promotion_job, "job"),
                                  "startDate": pick(doc, "startTime", "time"), "endDate": pick(doc, "endTime"), "jobType": kind,
                                  "fileType": pick(doc, "actionType") or ("晋升文件" if kind == "职级" else "任免职文件"),
                                  "page": pick(doc, "pageNum", "page"), "groupId": group_id, "raw": doc})
        if kind != "职级":
            for doc in as_list(item.get("divisionDocs")):
                if not isinstance(doc, dict): continue
                group_records.append({"caseId": case_id, "recordId": "", "fileId": pick(doc, "fileId"), "itemId": pick(doc, "itemId"),
                                      "fileName": pick(doc, "itemName"), "job": pick(doc, "position", "job"), "startDate": pick(doc, "time", "startTime"),
                                      "endDate": pick(doc, "endTime"), "jobType": "分工", "fileType": "分工文件", "page": pick(doc, "pageNum", "page"),
                                      "groupId": group_id, "raw": doc})
        records.extend(group_records)
        group_job = form_job or investigate_job or (group_records[0]["job"] if group_records else "")
        groups.append({"groupId": group_id, "job": group_job, "startDate": form_start or pick(item, "investigateStartTime"),
                       "endDate": form_end, "recordKeys": [record_key(row) for row in group_records]})
    return records, groups


def parse_output(payload):
    data = unwrap(payload)
    case_id = pick(data, "case_id", "caseId")
    files = [normalize_file(item, case_id) for item in as_list(data.get("fileClassifications") or data.get("classifications")) if isinstance(item, dict)]
    raw_positions, raw_ranks = data.get("positions"), data.get("ranks")
    actual_grouped_positions = any(isinstance(item, dict) and any(key in item for key in ("formPosition", "appointmentDocs", "divisionDocs", "investigatePosition")) for item in as_list(raw_positions))
    actual_grouped_ranks = any(isinstance(item, dict) and any(key in item for key in ("formRank", "promotionDocs", "investigateRank")) for item in as_list(raw_ranks))
    positions, position_groups = flatten_api_position_groups(raw_positions, "职务", case_id) if actual_grouped_positions else flatten_records(raw_positions, "职务", case_id)
    ranks, rank_groups = flatten_api_position_groups(raw_ranks, "职级", case_id) if actual_grouped_ranks else flatten_records(raw_ranks, "职级", case_id)
    explicit_groups = data.get("positionGroups") or data.get("groups")
    group_records, groups = flatten_records(explicit_groups, "", case_id) if explicit_groups else ([], [])
    if group_records:
        existing = {record_key(item) for item in positions + ranks}
        positions.extend(row for row in group_records if record_key(row) not in existing)
    groups = groups or position_groups + rank_groups
    identities = []
    for item in as_list(data.get("specialIdentities")):
        if isinstance(item, dict): identities.append(normalize_identity(item, "特殊身份", case_id))
    for item in as_list(data.get("otherIdentities")):
        if isinstance(item, dict): identities.append(normalize_identity(item, "其他身份", case_id))
    return {"files": files, "records": positions + ranks, "groups": groups, "identities": identities,
            "rules": as_list(data.get("businessRules") or data.get("displayResults")), "raw": data}


def greedy_match(gold, model, score_fn, threshold):
    candidates = []
    for gi, left in enumerate(gold):
        for mi, right in enumerate(model):
            score = score_fn(left, right)
            if score >= threshold:
                candidates.append((score, gi, mi))
    candidates.sort(reverse=True)
    used_g, used_m, result = set(), set(), []
    for score, gi, mi in candidates:
        if gi in used_g or mi in used_m: continue
        used_g.add(gi); used_m.add(mi)
        result.append({"goldIndex": gi, "modelIndex": mi, "score": round(score, 4)})
    return result, used_g, used_m


def metric(total_gold, total_model, matched):
    precision = matched / total_model if total_model else (1.0 if not total_gold else 0.0)
    recall = matched / total_gold if total_gold else 1.0
    return {"goldTotal": total_gold, "modelTotal": total_model, "matched": matched, "precision": precision,
            "recall": recall, "f1": 2 * precision * recall / (precision + recall) if precision + recall else 0.0}


def date_score(left, right):
    a, b = norm(left), norm(right)
    if not a and not b: return 1.0
    if not a or not b: return 0.0
    return similarity(a, b)


def evaluate(gold, payload, threshold=0.58):
    model = parse_output(payload)
    gold_files = [row for row in gold.get("files", []) if yes(row.get("should_process"))]
    def file_score(g, m):
        if norm(g.get("file_id")) and norm(g.get("file_id")) == norm(m.get("fileId")): return 1.0
        return 0.65 * similarity(g.get("file_name"), m.get("fileName")) + 0.35 * similarity(g.get("page"), m.get("page"))
    file_matches, used_gf, used_mf = greedy_match(gold_files, model["files"], file_score, threshold)
    file_type_correct = sum(similarity(gold_files[x["goldIndex"]].get("expected_file_type"), model["files"][x["modelIndex"]].get("fileType")) >= .9 for x in file_matches)

    gold_records = gold.get("records", [])
    def record_score(g, m):
        return (.35 * similarity(g.get("job"), m.get("job")) + .15 * similarity(g.get("job_type"), m.get("jobType")) +
                .15 * date_score(g.get("start_date"), m.get("startDate")) + .15 * date_score(g.get("end_date"), m.get("endDate")) +
                .10 * similarity(g.get("file_type"), m.get("fileType")) + .10 * similarity(g.get("page"), m.get("page")))
    record_matches, used_gr, used_mr = greedy_match(gold_records, model["records"], record_score, threshold)
    field_names = {"job": ("job", "job"), "startDate": ("start_date", "startDate"), "endDate": ("end_date", "endDate"),
                   "jobType": ("job_type", "jobType"), "fileType": ("file_type", "fileType"), "page": ("page", "page")}
    field_stats = {}
    for label, (gkey, mkey) in field_names.items():
        applicable = [match for match in record_matches if norm(gold_records[match["goldIndex"]].get(gkey))]
        correct = sum(similarity(gold_records[x["goldIndex"]].get(gkey), model["records"][x["modelIndex"]].get(mkey)) >= .9 for x in applicable)
        field_stats[label] = {"correct": correct, "total": len(applicable), "accuracy": correct / len(applicable) if applicable else 1.0}

    gold_group_map = {}
    record_by_id = {row.get("record_id"): row for row in gold_records}
    for row in gold.get("groups", []):
        group = gold_group_map.setdefault(row.get("group_id"), {"groupId": row.get("group_id"), "job": row.get("normalized_job"),
                                                                  "startDate": row.get("expected_start"), "endDate": row.get("expected_end"), "recordKeys": []})
        record = record_by_id.get(row.get("record_id"))
        if record: group["recordKeys"].append(record_key({"fileId": record.get("file_id"), "itemId": record.get("item_id"), "job": record.get("job"), "startDate": record.get("start_date"), "endDate": record.get("end_date")}))
    gold_groups = list(gold_group_map.values())
    def group_score(g, m):
        ga, ma = set(g.get("recordKeys") or []), set(m.get("recordKeys") or [])
        overlap = len(ga & ma) / len(ga | ma) if ga or ma else 0
        return .40 * overlap + .25 * similarity(g.get("job"), m.get("job")) + .20 * ((date_score(g.get("startDate"), m.get("startDate")) + date_score(g.get("endDate"), m.get("endDate"))) / 2) + .15 * (1 if overlap else 0)
    group_matches, used_gg, used_mg = greedy_match(gold_groups, model["groups"], group_score, .45)
    split_count = sum(1 for g in gold_groups if sum(group_score(g, m) >= .45 for m in model["groups"]) > 1)
    wrong_merge = sum(1 for m in model["groups"] if sum(group_score(g, m) >= .45 for g in gold_groups) > 1)

    gold_identities = gold.get("identities", [])
    def identity_score(g, m):
        return .65 * similarity(g.get("identity"), m.get("identity")) + .20 * similarity(g.get("identity_type"), m.get("identityType")) + .15 * max(similarity(g.get("case_input"), m.get("caseInput")), similarity(g.get("report_input"), m.get("reportInput")))
    identity_matches, used_gi, used_mi = greedy_match(gold_identities, model["identities"], identity_score, threshold)
    material_total = material_correct = 0
    for match in identity_matches:
        g, m = gold_identities[match["goldIndex"]], model["identities"][match["modelIndex"]]
        if g.get("material_file_id") or g.get("material_item_id"):
            material_total += 1
            material_correct += any((not g.get("material_file_id") or norm(g.get("material_file_id")) == norm(mat.get("fileId"))) and
                                    (not g.get("material_item_id") or norm(g.get("material_item_id")) == norm(mat.get("itemId"))) for mat in m.get("materials", []))

    model_rules = {pick(row, "rule_id", "ruleId", "id"): row for row in model["rules"] if isinstance(row, dict)}
    rule_details = []
    for row in gold.get("rules", []):
        actual = model_rules.get(row.get("rule_id"))
        if not actual:
            verdict = "未执行"
        else:
            expected = row.get("expected_result") or row.get("expected_display") or row.get("expected_issue")
            actual_value = pick(actual, "actual_result", "actualResult", "result", "verdict", "display", "issue")
            verdict = "通过" if similarity(expected, actual_value) >= .9 else "不通过"
        rule_details.append({"gold": row, "actual": actual, "verdict": verdict})
    executed = [row for row in rule_details if row["verdict"] in {"通过", "不通过"}]

    details = {"files": [], "records": [], "groups": [], "identities": [], "rules": rule_details}
    for match in file_matches: details["files"].append({"status": "matched", "gold": gold_files[match["goldIndex"]], "model": model["files"][match["modelIndex"]], **match})
    details["files"] += [{"status": "missed", "gold": row, "model": None} for i, row in enumerate(gold_files) if i not in used_gf]
    details["files"] += [{"status": "false_positive", "gold": None, "model": row} for i, row in enumerate(model["files"]) if i not in used_mf]
    for match in record_matches: details["records"].append({"status": "matched", "gold": gold_records[match["goldIndex"]], "model": model["records"][match["modelIndex"]], **match})
    details["records"] += [{"status": "missed", "gold": row, "model": None} for i, row in enumerate(gold_records) if i not in used_gr]
    details["records"] += [{"status": "false_positive", "gold": None, "model": row} for i, row in enumerate(model["records"]) if i not in used_mr]
    for match in group_matches: details["groups"].append({"status": "matched", "gold": gold_groups[match["goldIndex"]], "model": model["groups"][match["modelIndex"]], **match})
    details["groups"] += [{"status": "missed", "gold": row, "model": None} for i, row in enumerate(gold_groups) if i not in used_gg]
    details["groups"] += [{"status": "false_positive", "gold": None, "model": row} for i, row in enumerate(model["groups"]) if i not in used_mg]
    for match in identity_matches: details["identities"].append({"status": "matched", "gold": gold_identities[match["goldIndex"]], "model": model["identities"][match["modelIndex"]], **match})
    details["identities"] += [{"status": "missed", "gold": row, "model": None} for i, row in enumerate(gold_identities) if i not in used_gi]
    details["identities"] += [{"status": "false_positive", "gold": None, "model": row} for i, row in enumerate(model["identities"]) if i not in used_mi]

    summary = {
        "fileClassification": {**metric(len(gold_files), len(model["files"]), len(file_matches)), "typeAccuracy": file_type_correct / len(file_matches) if file_matches else 1.0},
        "recordExtraction": metric(len(gold_records), len(model["records"]), len(record_matches)), "fields": field_stats,
        "alignment": {**metric(len(gold_groups), len(model["groups"]), len(group_matches)), "groupAccuracy": len(group_matches) / len(gold_groups) if gold_groups else 1.0,
                      "splitRate": split_count / len(gold_groups) if gold_groups else 0.0, "mergeRate": wrong_merge / len(model["groups"]) if model["groups"] else 0.0},
        "identity": metric(len(gold_identities), len(model["identities"]), len(identity_matches)),
        "materialAccuracy": material_correct / material_total if material_total else 1.0,
        "materialCorrect": material_correct, "materialTotal": material_total,
        "businessRules": {"passed": sum(row["verdict"] == "通过" for row in executed), "failed": sum(row["verdict"] == "不通过" for row in executed),
                          "notExecuted": sum(row["verdict"] == "未执行" for row in rule_details),
                          "passRate": sum(row["verdict"] == "通过" for row in executed) / len(executed) if executed else 0.0},
    }
    return {"summary": summary, "details": details, "model": model, "threshold": threshold}


class SubjectIdentityEvaluationService:
    def __init__(self, db_path, output_dir):
        self.db_path, self.output_dir = Path(db_path), Path(output_dir)

    def connect(self):
        conn = sqlite3.connect(self.db_path); conn.row_factory = sqlite3.Row; return conn

    def init_db(self):
        with self.connect() as conn:
            conn.executescript("""
            CREATE TABLE IF NOT EXISTS subject_eval_datasets(id TEXT PRIMARY KEY,name TEXT NOT NULL,version TEXT NOT NULL,description TEXT NOT NULL DEFAULT '',status TEXT NOT NULL DEFAULT 'draft',created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS subject_eval_gold(dataset_id TEXT PRIMARY KEY,content TEXT NOT NULL,updated_at TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS subject_eval_runs(id TEXT PRIMARY KEY,dataset_id TEXT NOT NULL,name TEXT NOT NULL,model_version TEXT NOT NULL DEFAULT '',system_version TEXT NOT NULL DEFAULT '',source TEXT NOT NULL,result TEXT NOT NULL,created_at TEXT NOT NULL);
            CREATE INDEX IF NOT EXISTS idx_subject_runs_dataset ON subject_eval_runs(dataset_id,created_at);
            """)

    def list_datasets(self):
        with self.connect() as conn:
            rows = conn.execute("""SELECT d.*,CASE WHEN g.dataset_id IS NULL THEN 0 ELSE 1 END has_gold,COUNT(r.id) run_count FROM subject_eval_datasets d LEFT JOIN subject_eval_gold g ON g.dataset_id=d.id LEFT JOIN subject_eval_runs r ON r.dataset_id=d.id GROUP BY d.id ORDER BY d.updated_at DESC""").fetchall()
        return [{"id": r["id"], "name": r["name"], "version": r["version"], "description": r["description"], "status": r["status"], "hasGold": bool(r["has_gold"]), "runCount": r["run_count"], "updatedAt": r["updated_at"]} for r in rows]

    def get_dataset(self, dataset_id):
        with self.connect() as conn:
            row = conn.execute("SELECT * FROM subject_eval_datasets WHERE id=?", (dataset_id,)).fetchone()
            if not row: raise ValueError("主体身份评测集不存在。")
            gold_row = conn.execute("SELECT content FROM subject_eval_gold WHERE dataset_id=?", (dataset_id,)).fetchone()
            runs = conn.execute("SELECT id,name,model_version,system_version,created_at,result FROM subject_eval_runs WHERE dataset_id=? ORDER BY created_at DESC", (dataset_id,)).fetchall()
        gold = json.loads(gold_row["content"]) if gold_row else {key: [] for key in ("cases", "files", "records", "groups", "identities", "rules")}
        return {"id": row["id"], "name": row["name"], "version": row["version"], "description": row["description"], "status": row["status"], "gold": gold,
                "runs": [{"id": r["id"], "name": r["name"], "modelVersion": r["model_version"], "systemVersion": r["system_version"], "createdAt": r["created_at"], "summary": json.loads(r["result"])["summary"]} for r in runs]}

    def save_dataset(self, payload):
        dataset_id, name = text(payload.get("id")) or uuid.uuid4().hex, text(payload.get("name"))
        if not name: raise ValueError("请填写评测集名称。")
        now = time.strftime("%Y-%m-%d %H:%M:%S")
        with self.connect() as conn:
            old = conn.execute("SELECT created_at,status FROM subject_eval_datasets WHERE id=?", (dataset_id,)).fetchone()
            if old and old["status"] == "frozen" and payload.get("status") != "frozen": raise ValueError("已冻结评测集不能解除冻结。")
            conn.execute("""INSERT INTO subject_eval_datasets VALUES(?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,version=excluded.version,description=excluded.description,status=excluded.status,updated_at=excluded.updated_at""",
                         (dataset_id, name, text(payload.get("version")) or "v1.0", text(payload.get("description")), text(payload.get("status")) or "draft", old["created_at"] if old else now, now))
        return self.get_dataset(dataset_id)

    def delete_dataset(self, dataset_id):
        with self.connect() as conn:
            row = conn.execute("SELECT name FROM subject_eval_datasets WHERE id=?", (dataset_id,)).fetchone()
            if not row: raise ValueError("主体身份评测集不存在。")
            run_count = conn.execute("SELECT COUNT(*) FROM subject_eval_runs WHERE dataset_id=?", (dataset_id,)).fetchone()[0]
            conn.execute("DELETE FROM subject_eval_runs WHERE dataset_id=?", (dataset_id,)); conn.execute("DELETE FROM subject_eval_gold WHERE dataset_id=?", (dataset_id,)); conn.execute("DELETE FROM subject_eval_datasets WHERE id=?", (dataset_id,))
        return {"deleted": True, "name": row["name"], "runCount": run_count}

    def template(self):
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        for sheet_name, headers in SHEETS.items():
            ws = wb.create_sheet(sheet_name); ws.append([CHINESE_HEADERS[sheet_name][key] for key in headers]); ws.append(EXAMPLES[sheet_name]); ws.freeze_panes = "A2"
            for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="2563EB")
            for idx in range(1, len(headers) + 1): ws.column_dimensions[get_column_letter(idx)].width = 20
        stream = BytesIO(); wb.save(stream); return base64.b64encode(stream.getvalue()).decode("ascii")

    def import_gold(self, dataset_id, encoded):
        dataset = self.get_dataset(dataset_id)
        if dataset["status"] == "frozen": raise ValueError("评测集已冻结，不能修改黄金标准。")
        wb = openpyxl.load_workbook(BytesIO(base64.b64decode(encoded)), data_only=True)
        mapping = {"案件信息": "cases", "文件分类": "files", "职位抽取": "records", "任职分组": "groups", "身份材料": "identities", "业务规则": "rules"}
        gold = {}
        for sheet_name, key in mapping.items():
            if sheet_name not in wb.sheetnames: raise ValueError(f"Excel缺少工作表：{sheet_name}")
            ws = wb[sheet_name]; uploaded_headers = [text(cell.value) for cell in ws[1]]
            header_indexes = {}
            missing = []
            for internal_name in SHEETS[sheet_name]:
                chinese_name = CHINESE_HEADERS[sheet_name][internal_name]
                if chinese_name in uploaded_headers:
                    header_indexes[internal_name] = uploaded_headers.index(chinese_name)
                elif internal_name in uploaded_headers:
                    header_indexes[internal_name] = uploaded_headers.index(internal_name)
                else:
                    missing.append(chinese_name)
            if missing: raise ValueError(f"{sheet_name}缺少列：{'、'.join(missing)}")
            gold[key] = [{internal_name: text(values[index]) for internal_name, index in header_indexes.items()} for values in ws.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in values)]
        now = time.strftime("%Y-%m-%d %H:%M:%S")
        with self.connect() as conn:
            conn.execute("INSERT INTO subject_eval_gold(dataset_id,content,updated_at) VALUES(?,?,?) ON CONFLICT(dataset_id) DO UPDATE SET content=excluded.content,updated_at=excluded.updated_at", (dataset_id, json.dumps(gold, ensure_ascii=False), now))
            conn.execute("UPDATE subject_eval_datasets SET updated_at=? WHERE id=?", (now, dataset_id))
        return self.get_dataset(dataset_id)

    def run(self, payload):
        dataset = self.get_dataset(text(payload.get("datasetId")))
        if not any(dataset["gold"].values()): raise ValueError("请先导入主体身份黄金标准。")
        source = payload.get("source")
        if not isinstance(source, dict): raise ValueError("系统结果必须是JSON对象。")
        result = evaluate(dataset["gold"], source, float(payload.get("threshold") or .58))
        run_id, now = uuid.uuid4().hex, time.strftime("%Y-%m-%d %H:%M:%S")
        name = text(payload.get("name")) or f"主体身份评测 {now}"
        with self.connect() as conn:
            conn.execute("INSERT INTO subject_eval_runs VALUES(?,?,?,?,?,?,?,?)", (run_id, dataset["id"], name, text(payload.get("modelVersion")), text(payload.get("systemVersion")), json.dumps(source, ensure_ascii=False), json.dumps(result, ensure_ascii=False), now))
        return {"id": run_id, "name": name, "createdAt": now, **result}

    def get_run(self, run_id):
        with self.connect() as conn: row = conn.execute("SELECT * FROM subject_eval_runs WHERE id=?", (run_id,)).fetchone()
        if not row: raise ValueError("主体身份评测记录不存在。")
        return {"id": row["id"], "datasetId": row["dataset_id"], "name": row["name"], "modelVersion": row["model_version"], "systemVersion": row["system_version"], "createdAt": row["created_at"], **json.loads(row["result"])}

    def fetch_api(self, base_url, endpoint, case_id, biz_id, timeout=30):
        base_url, case_id, biz_id = text(base_url).rstrip("/"), text(case_id), text(biz_id)
        if not base_url or urllib.parse.urlparse(base_url).scheme not in {"http", "https"}:
            raise ValueError("请填写有效的主体身份接口地址。")
        if not case_id or not biz_id: raise ValueError("caseId 和 bizId 不能为空。")
        if base_url.endswith("/identity") or base_url.endswith("/progress"):
            url = base_url.rsplit("/", 1)[0] + "/" + endpoint
        elif base_url.endswith("/subject_identity"):
            url = base_url + "/" + endpoint
        else:
            url = base_url + "/subject_identity/" + endpoint
        body = json.dumps({"caseId": case_id, "bizId": biz_id}, ensure_ascii=False).encode("utf-8")
        request = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json", "Accept": "application/json"}, method="POST")
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                result = json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            raise ValueError(f"主体身份接口请求失败：HTTP {exc.code}") from exc
        except urllib.error.URLError as exc:
            raise ValueError(f"主体身份接口请求失败：{exc.reason}") from exc
        except json.JSONDecodeError as exc:
            raise ValueError("主体身份接口返回的不是合法JSON。") from exc
        if isinstance(result, dict) and result.get("code") not in (None, 0, "0"):
            raise ValueError(text(result.get("msg") or result.get("message") or "接口返回失败"))
        return {"url": url, "content": result}

    def export_run(self, run_id):
        run = self.get_run(run_id); self.output_dir.mkdir(parents=True, exist_ok=True)
        path = self.output_dir / f"主体身份评测报告_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "综合概览"
        ws.append(["评测名称", run["name"]]); ws.append(["运行时间", run["createdAt"]]); ws.append(["模型版本", run["modelVersion"]]); ws.append(["系统版本", run["systemVersion"]]); ws.append([])
        ws.append(["阶段", "黄金数", "模型数", "匹配数", "精确率", "召回率", "F1"])
        for key, label in (("fileClassification", "文件分类"), ("recordExtraction", "职位抽取"), ("alignment", "任职对齐"), ("identity", "身份识别")):
            m = run["summary"][key]; ws.append([label, m["goldTotal"], m["modelTotal"], m["matched"], m["precision"], m["recall"], m["f1"]])
        fields = wb.create_sheet("字段准确率"); fields.append(["字段", "正确数", "匹配数", "准确率"])
        for key, item in run["summary"]["fields"].items(): fields.append([key, item["correct"], item["total"], item["accuracy"]])
        for section, label in (("files", "文件分类明细"), ("records", "职位抽取明细"), ("groups", "任职分组明细"), ("identities", "身份材料明细")):
            sheet = wb.create_sheet(label); sheet.append(["状态", "匹配分", "黄金标准", "系统结果"])
            for row in run["details"][section]: sheet.append([row.get("status"), row.get("score", ""), json.dumps(row.get("gold"), ensure_ascii=False), json.dumps(row.get("model"), ensure_ascii=False)])
        rules = wb.create_sheet("业务规则"); rules.append(["规则编号", "规则类型", "预期", "实际", "结论"])
        for row in run["details"]["rules"]: rules.append([row["gold"].get("rule_id"), row["gold"].get("rule_type"), row["gold"].get("expected_result"), json.dumps(row.get("actual"), ensure_ascii=False), row["verdict"]])
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
            for idx in range(1, sheet.max_column + 1): sheet.column_dimensions[get_column_letter(idx)].width = 22
        wb.save(path); return path
